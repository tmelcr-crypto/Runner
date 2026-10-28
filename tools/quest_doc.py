#!/usr/bin/env python3
"""The quest pages: a Word file (.docx, opens in Pages) for a quest, made from its data in js/01x-mission-data.js.

  python3 tools/quest_doc.py q02                 write docs/quests/Quest-02-<Title>.docx
  python3 tools/quest_doc.py all                 every quest that has a page part (not Quest 01: its page came first, by hand)

The layout is Quest 01's page (docs/quests/Quest-01-Harbour-Favour.docx): its blocks are copied and filled in - the characters,
the steps, the fails, the dialogue (LIVE lines show during play without pausing), a shot list per cutscene, the rewards, the
start, the notes and the pictures to send. What only the page needs sits in the quest's "page" part; the rest (characters,
steps, lines, fails, rewards) is the same data the game plays. Needs python-docx."""
import copy, json, os, re, sys
import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'js', '01x-mission-data.js')
CHARS = os.path.join(ROOT, 'js', '01y-character-data.js')
BASE = os.path.join(ROOT, 'docs', 'quests', 'Quest-01-Harbour-Favour.docx')
OUT = os.path.join(ROOT, 'docs', 'quests')


def load():                                  # the missions, each with its characters filled in from the character table (js/01y)
    s = open(DATA, encoding='utf-8').read()
    m = re.search(r'/\*MISSIONS-JSON\*/(.*?)/\*END-MISSIONS-JSON\*/', s, re.S)
    c = open(CHARS, encoding='utf-8').read()
    ch = {x['id']: x for x in json.loads(re.search(r'/\*CHARACTERS-JSON\*/(.*?)/\*END-CHARACTERS-JSON\*/', c, re.S).group(1))['characters']}
    ms = json.loads(m.group(1))['missions']
    for q in ms: q['chars'] = [ch[i] if isinstance(i, str) else i for i in q['chars']]
    return ms


def set_text(p, text):                      # a paragraph's words, keeping its first run's look
    runs = p.runs
    if not runs: p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]: r.text = ''


def set_cell(cell, text):                   # a cell's words (several lines: one paragraph each, the first one's look)
    lines = str(text).split('\n'); ps = cell.paragraphs
    set_text(ps[0], lines[0])
    for p in ps[1:]: p._element.getparent().remove(p._element)
    for ln in lines[1:]:
        q = copy.deepcopy(ps[0]._element); ps[0]._element.getparent().append(q)
        set_text(docx.text.paragraph.Paragraph(q, ps[0]._parent), ln)


def fill_rows(table, rows):                 # keep the header row, one copy of row 1 per item
    proto = table.rows[1]._tr
    for r in list(table.rows)[1:]: table._tbl.remove(r._tr)
    for vals in rows:
        tr = copy.deepcopy(proto); table._tbl.append(tr)
        row = docx.table._Row(tr, table)
        for c, v in zip(row.cells, vals): set_cell(c, v)


def build(m):
    P = m['page']; d = docx.Document(BASE); T = d.tables; body = d.element.body
    paras = d.paragraphs
    who = {c['id']: c for c in m['chars']}
    # the title
    for p in paras[:6]:
        if p.text.strip() == 'QUEST 01': set_text(p, 'QUEST ' + m['no'])
        elif p.text.strip() == 'HARBOUR FAVOUR': set_text(p, m['title'])
        elif p.text.startswith('The test mission'): set_text(p, P.get('sub', 'Written from your answers - change anything you like.'))
    for sec in d.sections:                    # the footer: this quest's number and name (the page number after it stays)
        for fp in sec.footer.paragraphs:
            for r in fp.runs:
                if 'Quest 01' in r.text: r.text = 'Block Runner · Quest %s · %s   ·   ' % (m['no'], m['title'].title())
    g = P['glance']
    for r, v in zip(T[0].rows, [m['no'] + ' - ' + g.get('kind', 'story'), g['given'], g['where'], g['unlocked'], g.get('time', 'Any'), g['length'], g['limit']]):
        set_cell(r.cells[1], v)
    # the characters: Quest 01's first block (its table, the mumble line and the blank lines between) once per character
    els = list(body.iterchildren()); i2, i4 = els.index(T[2]._tbl), els.index(T[4]._tbl)
    more = next(e for e in els[i4:] if e.tag.endswith('}p') and ''.join(e.itertext()).startswith('More characters'))
    block, old = els[i2:i4], els[i2:els.index(more)]
    for i, c in enumerate(m['chars']):
        pc = P['chars'][c['id']]; cp = [copy.deepcopy(e) for e in block]
        for e in cp: more.addprevious(e)
        bt, mt = [docx.table.Table(e, d) for e in cp if e.tag.endswith('}tbl')]
        cell0, cell1 = bt.rows[0].cells
        set_text(cell0.paragraphs[0], 'CHARACTER %d' % (i + 1))
        pic = c['id'] + '-calm.png'                # the drawn pixel-art portrait (tools/portraits.py), in the portrait box
        pc0 = cell0.tables[0].rows[0].cells[0]; set_cell(pc0, 'PORTRAIT'); png = os.path.join(ROOT, 'img', 'chars', pic)
        if os.path.exists(png):
            from PIL import Image
            big = os.path.join(OUT, '_' + pic); Image.open(png).resize((192, 300), Image.NEAREST).save(big)
            pc0.add_paragraph().add_run().add_picture(big, width=docx.shared.Cm(1.8)); os.remove(big)
        else: pc0.add_paragraph('img/chars/' + pic)
        for r, v in zip(cell1.tables[0].rows, [c['name'], c.get('letter') if pc.get('gives') else '-', pc['role'], pc['looks'], pc['bubble'], pic]):
            set_cell(r.cells[1], v)
        set_cell(mt.rows[0].cells[0], 'MUMBLE  (chosen by me)  ' + pc['mumble'])
    for e in old: body.remove(e)
    T = d.tables
    # the story, the steps, the fails
    k = 2 + 2 * len(m['chars'])                   # the story's table now
    set_cell(T[k].rows[0].cells[0], m['story'])
    fill_rows(T[k + 1], [[str(i + 1), s['word'], P['steps'][i][0], s['show'], P['steps'][i][1]] for i, s in enumerate(m['steps'])])
    fill_rows(T[k + 3], P['fails'])
    # the dialogue
    for p in d.paragraphs:
        if p.text.startswith('The game pauses for every line'):
            set_text(p, 'In a cutscene the game pauses for every line: the portrait and the bubble, the words appearing letter by letter with the '
                        'speaker\'s mumble; tap for the next. Lines marked LIVE show during play without pausing - a shout in a fight, '
                        'a phone call while you drive - and go by themselves.')
    fill_rows(T[k + 5], [[str(i + 1), who[l['who']]['name'].split()[0], l['text'], l.get('when', ''), l.get('mood', '').capitalize()] for i, l in enumerate(m['lines'])])
    # the cutscenes: Quest 01's first one (its title, its 'Plays' line and its table) once per cutscene
    els = list(body.iterchildren()); t13 = els.index(T[k + 7]._tbl)
    first = els[t13 - 2:t13 + 1]
    end = next(e for e in els[t13:] if e.tag.endswith('}p') and ''.join(e.itertext()).startswith('Every cutscene can be skipped'))
    old = els[t13 - 2:els.index(end)]
    for cu in P['cuts']:
        cp = [copy.deepcopy(e) for e in first]
        for e in cp: end.addprevious(e)
        set_text(docx.text.paragraph.Paragraph(cp[0], d), cu['title'])
        set_text(docx.text.paragraph.Paragraph(cp[1], d), 'Plays: ' + cu['plays'])
        fill_rows(docx.table.Table(cp[2], d), [[str(i + 1)] + [str(v) for v in sh] for i, sh in enumerate(cu['shots'])])
    for e in old: body.remove(e)
    T = d.tables
    k2 = k + 7 + len(P['cuts'])                   # the rewards' table now
    r = m['reward']
    for row, v in zip(T[k2].rows, ['€{:,}'.format(r['cash']), r.get('vehicle') or '-', r.get('clothes') or '-', r.get('next') or '-', r.get('place') or '-', r.get('contact') or '-']):
        set_cell(row.cells[1], v)
    st = P['start']
    set_cell(T[k2 + 1].rows[0].cells[1], ('☒' if st['how'] == 'marker' else '☐') + ' A marker on the map   - the giver\'s letter in a circle; walk in and the opening cutscene plays\n'
             + ('☒' if st['how'] == 'phone' else '☐') + ' A phone call   - the phone rings, the caller\'s portrait and bubble, then a marker')
    set_cell(T[k2 + 1].rows[1].cells[1], st.get('marker', '-'))
    set_cell(T[k2 + 1].rows[2].cells[1], st.get('phone', '-'))
    set_cell(T[k2 + 3].rows[0].cells[1], P.get('failing', '-'))
    n = P['notes']
    for row, v in zip(T[k2 + 4].rows, [n['places'], n['people'], n['music'], n['new'], n.get('other', '-')]): set_cell(row.cells[1], v)
    pics = T[k2 + 5]; proto = pics.rows[0]._tr
    for rr in list(pics.rows): pics._tbl.remove(rr._tr)
    for c in m['chars']:
        tr = copy.deepcopy(proto); pics._tbl.append(tr); row = docx.table._Row(tr, pics)
        set_cell(row.cells[0], c['id'] + '-calm / -angry / -smile.png'); set_cell(row.cells[1], c['name'] + ' - drawn (pixel art, 48 x 75, docs/Characters.docx); a picture of your own replaces them')
    name = 'Quest-%s-%s.docx' % (m['no'], '-'.join(w.capitalize() for w in re.sub(r'[^A-Za-z0-9 ]', '', m['title']).split()))
    path = os.path.join(OUT, name); d.save(path); return path


if __name__ == '__main__':
    if len(sys.argv) < 2: print(__doc__); sys.exit(1)
    ms = load(); want = sys.argv[1]
    for m in ms:
        if 'page' in m and (want == 'all' or m['id'] == want): print('wrote', os.path.relpath(build(m), ROOT))
