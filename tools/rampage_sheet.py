"""Rampage table <-> spreadsheet (Apple Numbers or Excel).

  python3 tools/rampage_sheet.py export rampages.numbers            write js/01h-rampage-data.js as a Numbers file (or .xlsx)
  python3 tools/rampage_sheet.py import rampages.numbers [--dry-run] read an edited file (.numbers or .xlsx) back, check it, rewrite the table

One row per rampage, one column per setting (found by its heading, so columns may be moved). The game places the rampages itself:
the first row nearest to where you start, the rest spread over the city. A row without an ID is left out.
Needs openpyxl for .xlsx and numbers-parser for .numbers."""
import json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
JS = os.path.join(HERE, '..', 'js', '01h-rampage-data.js'); WJS = os.path.join(HERE, '..', 'js', '01f-weapon-data.js')
A, B = '/*RAMPAGE-JSON*/', '/*END-RAMPAGE-JSON*/'
TARGETS = ['people', 'cars']
MAX_N = 40


def weapons():                              # the weapons in the game: id -> name
    s = open(WJS, encoding='utf-8').read(); t = json.loads(s[s.index('/*WEAPON-JSON*/') + 15:s.index('/*END-WEAPON-JSON*/')])
    return {w['id']: w['name'] for w in t['weapons'] if w.get('status') == 'in game'}


W = weapons()
# a column: (key, heading, kind, choices or (min, max), what it means, how to fill)
COLS = [
    ('id', 'ID', 'text', None, 'Internal key of the rampage (saved games remember which ones you passed by it). Never shown.', 'r01, r02 ... unique'),
    ('name', 'Name', 'text', None, 'Shown on the rampage screen and at the bottom while it runs.', 'CAPITALS look best'),
    ('weapon', 'Weapon', 'choice', list(W), 'The weapon you get for it, locked in your hand with endless spare ammo (reloads still take their time). '
     'Pass it the first time and you keep it with its basic load. Weapon IDs: ' + ', '.join('%s = %s' % kv for kv in W.items()) + '.', 'pick from the list'),
    ('target', 'Target', 'choice', TARGETS, 'people: kill anyone on foot (police officers count too). cars: wreck vehicles - one counts when it catches fire.', 'pick from the list'),
    ('count', 'How many', 'int', (1, 200), 'How many people to kill, or vehicles to wreck.', 'count'),
    ('time', 'Time (s)', 'int', (10, 900), 'How long you have. The clock stops while the game waits (pause, map).', 'seconds'),
    ('reward', 'Reward (€)', 'int', (0, 1000000), 'Cash for passing it the first time (your score is your cash). Played again: no reward.', 'euros'),
    ('notes', 'Notes', 'long', None, 'Anything else - not used by the game.', 'free text'),
]
KEYS = [c[0] for c in COLS]; COL = {c[0]: c for c in COLS}
MATCH = [('id', 'id'), ('name', 'name'), ('weapon', 'weapon'), ('target', 'target'), ('count', 'how many'), ('time', 'time'), ('reward', 'reward'), ('notes', 'notes')]
TIPS = [('Yellow cells', 'What you type in. One row per rampage. Weapon and Target are lists to pick from.', 'tap a cell and type'),
        ('Order', 'The first row is the rampage nearest to where you start (make it an easy one); the game spreads the others over the city.', ''),
        ('New rampage', 'Fill the next empty row and give it a new ID. Up to %d.' % MAX_N, ''),
        ('Remove one', 'Delete its row (or clear its ID).', ''),
        ('How hard', 'A bot that aims perfectly but walks badly passed the table\'s rampages with time to spare; people play slower. As a guide, give at least 3 s per person with guns, 5 s with fists, 10-20 s per vehicle with guns, 8-12 s with explosives.', ''),
        ('Sending it back', 'Upload the file in the chat as it is (.numbers or .xlsx). Every row with an ID becomes a rampage.', '')]


def read_table():
    s = open(JS, encoding='utf-8').read(); i, j = s.index(A) + len(A), s.index(B)
    return s, i, j, json.loads(s[i:j])


def dump(rs):
    return '{\n"rampages": [\n' + ',\n'.join('  ' + json.dumps({k: r.get(k) for k in KEYS}, ensure_ascii=False) for r in rs) + '\n]\n}'


fmt = lambda v: ('%d' % v) if float(v).is_integer() else ('%g' % v)
choices = lambda key: COL[key][3] if COL[key][2] == 'choice' else None


def guide_rows(): return [(c[1], c[4], c[5]) for c in COLS]


# ---------- export ----------
def export_xlsx(path, t):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.datavalidation import DataValidation
    F = 'Arial'; f_in = PatternFill('solid', fgColor='FFF7CC'); f_head = PatternFill('solid', fgColor='2B2442'); f_lab = PatternFill('solid', fgColor='E4E0F2'); f_calc = PatternFill('solid', fgColor='EDEDED')
    thin = Side(style='thin', color='C9C9C9'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    rs = t['rampages']; NR = max(MAX_N, len(rs)); wb = Workbook(); sh = wb.active; sh.title = 'Rampages'
    widths = {'id': 8, 'name': 26, 'weapon': 14, 'target': 11, 'count': 11, 'time': 10, 'reward': 13, 'notes': 46}
    heads = [c[1] for c in COLS] + ['Calc: seconds each', 'Calc: Check']
    for j, h in enumerate(heads, start=1):
        c = sh.cell(row=1, column=j, value=h); c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = f_head; c.border = box; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sh.column_dimensions[L(j)].width = widths.get(KEYS[j - 1] if j <= len(KEYS) else '', 16)
    for n in range(NR):
        r = rs[n] if n < len(rs) else {}; i = n + 2
        for j, (key, lab, kind, *_r) in enumerate(COLS, start=1):
            c = sh.cell(row=i, column=j, value=r.get(key)); c.border = box; c.font = Font(name=F, size=10, color='0000FF', bold=key == 'id'); c.fill = f_lab if key == 'id' else f_in
            c.alignment = Alignment(horizontal='left' if kind in ('long', 'text') else 'center', vertical='center', wrap_text=kind == 'long')
        X = lambda key: '%s%d' % (L(KEYS.index(key) + 1), i)
        c = sh.cell(row=i, column=len(KEYS) + 1, value='=IF(OR(%s="",%s=""),"",ROUND(%s/%s,1))' % (X('count'), X('time'), X('time'), X('count')))
        c.fill = f_calc; c.border = box; c.font = Font(name=F, size=10); c.alignment = Alignment(horizontal='center')
        chk = ('=IF({i}="","",IF(COUNTIF($A$2:$A${last},{i})>1,"Duplicate ID",IF(OR({n}="",{w}="",{t}="",{c}="",{s}=""),"Fill every column",'
               'IF(AND({t}="people",{s}/{c}<3),"Very tight",IF(AND({t}="cars",{s}/{c}<8),"Very tight","OK")))))').format(
            i=X('id'), last=NR + 1, n=X('name'), w=X('weapon'), t=X('target'), c=X('count'), s=X('time'))
        c = sh.cell(row=i, column=len(KEYS) + 2, value=chk); c.fill = f_calc; c.border = box; c.font = Font(name=F, size=10); c.alignment = Alignment(horizontal='center')
    for j, key in enumerate(KEYS, start=1):
        span = '%s2:%s%d' % (L(j), L(j), NR + 1); ch = choices(key); kind = COL[key][2]
        if ch: dv = DataValidation(type='list', formula1='"%s"' % ','.join(ch), allow_blank=True, showErrorMessage=True, error='Pick from the list.')
        elif kind == 'int': lo, hi = COL[key][3]; dv = DataValidation(type='whole', operator='between', formula1=fmt(lo), formula2=fmt(hi), allow_blank=True, showErrorMessage=True, error='Enter a whole number from %s to %s.' % (fmt(lo), fmt(hi)))
        else: continue
        sh.add_data_validation(dv); dv.add(span)
    sh.freeze_panes = 'C2'
    g = wb.create_sheet('How to fill'); g.column_dimensions['A'].width = 22; g.column_dimensions['B'].width = 80; g.column_dimensions['C'].width = 24
    for j, h in enumerate(['Item', 'What it means', 'Unit / how to fill'], start=1):
        c = g.cell(row=1, column=j, value=h); c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = f_head; c.border = box
    for i, row in enumerate(TIPS + [('Grey columns', 'Seconds each = time / how many. Check says OK, or what to look at.', '')] + guide_rows(), start=2):
        for j, v in enumerate(row, start=1):
            c = g.cell(row=i, column=j, value=v); c.font = Font(name=F, size=10, bold=j == 1); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = box
            if j == 1: c.fill = f_lab
    g.freeze_panes = 'A2'; wb.save(path)


def export_numbers(path, t):
    from numbers_parser import Alignment, Document, RGB
    sys.path.insert(0, HERE); from settings_sheet import exact_numbers; exact_numbers()
    rs = t['rampages']; NR = max(MAX_N, len(rs))
    doc = Document(sheet_name='Rampages', table_name='Rampages', num_header_rows=1, num_header_cols=1, num_rows=NR + 1, num_cols=len(COLS)); tb = doc.sheets[0].tables[0]
    S = {}
    def style(name, **kw):
        if name not in S: S[name] = doc.add_style(name=name, font_name='Helvetica Neue', font_size=kw.pop('size', 11.0), alignment=Alignment(kw.pop('h', 'left'), kw.pop('v', 'middle')), **kw)
        return S[name]
    head = style('BR head', bg_color=RGB(43, 36, 66), font_color=RGB(255, 255, 255), bold=True, h='center'); lab = style('BR label', bg_color=RGB(228, 224, 242), bold=True)
    inp = style('BR value', bg_color=RGB(255, 247, 204), font_color=RGB(0, 0, 255), h='center'); txt_in = style('BR text value', bg_color=RGB(255, 247, 204), font_color=RGB(0, 0, 255))
    txt = style('BR text')
    for j, w in enumerate([60, 190, 110, 90, 80, 80, 100, 300]): tb.col_width(j, w)
    for j, c in enumerate(COLS): tb.write(0, j, c[1], style=head)
    for n in range(NR):
        r = rs[n] if n < len(rs) else {}; i = n + 1
        for j, (key, _l, kind, *_r) in enumerate(COLS):
            v = r.get(key); st = lab if key == 'id' else txt_in if kind in ('text', 'long') else inp; ch = choices(key)
            if v is None or v == '':
                if ch: tb.write(i, j, '', style=st)
                else: tb.set_cell_style(i, j, st)
            else: tb.write(i, j, v, style=st)
            if ch: tb.set_cell_formatting(i, j, 'popup', popup_values=ch, allow_none=True)
        tb.row_height(i, 26)
    rows = TIPS + [('Checks', 'This Numbers file has no formulas: every row is checked when the file comes back, and anything to fix is listed for you.', '')] + guide_rows()
    doc.add_sheet('How to fill', 'How to fill', num_rows=len(rows) + 1, num_cols=3)
    g = doc.sheets[-1].tables[0]; g.num_header_rows = 1; g.num_header_cols = 1; g.col_width(0, 160); g.col_width(1, 480); g.col_width(2, 160)
    for j, h in enumerate(['Item', 'What it means', 'Unit / how to fill']): g.write(0, j, h, style=head)
    for i, row in enumerate(rows, start=1):
        g.write(i, 0, row[0], style=lab); g.write(i, 1, row[1], style=txt); g.write(i, 2, row[2] or '', style=txt); g.row_height(i, 60 if len(row[1]) > 120 else 40)
    doc.save(path)


# ---------- import ----------
def read_rows(path):
    if path.lower().endswith('.numbers'):
        from numbers_parser import Document
        sheets = {sh.name: [tuple(r) for r in sh.tables[0].rows(values_only=True)] for sh in Document(path).sheets}
    else:
        from openpyxl import load_workbook
        sheets = {ws.title: list(ws.iter_rows(values_only=True)) for ws in load_workbook(path, data_only=True).worksheets}
    rows = next((r for n, r in sheets.items() if n.lower().startswith('rampage')), None) or next(iter(sheets.values()))
    col = {}
    for j, h in enumerate(rows[0]):
        h = str(h or '').strip().lower()
        for key, start in MATCH:
            if (h == start if key == 'id' else h.startswith(start)) and key not in col: col[key] = j; break
    if 'id' not in col: raise SystemExit('No column headed ID in the first row.')
    out = []
    for row in rows[1:]:
        v = row[col['id']] if col['id'] < len(row) else None
        if v in (None, '') or str(v).strip() == '': continue
        out.append({key: (row[j] if j < len(row) else None) for key, j in col.items()})
    return out, set(col)


def num(v):
    if isinstance(v, bool) or v is None: return None
    if isinstance(v, str):
        try: v = float(v.strip().replace(',', '').replace('$', '').replace('€', ''))
        except ValueError: return None
    v = round(float(v), 6); return int(v) if v.is_integer() else v


def build(rows, present, old):
    olds = {r['id']: r for r in old}; out, errors, warns, seen = [], [], [], set()
    for k, row in enumerate(rows):
        rid = re.sub(r'[^a-z0-9]+', '_', str(row['id']).strip().lower()).strip('_'); where = 'Row %d (%s)' % (k + 2, rid)
        if rid in seen: errors.append('%s: the ID is used twice' % where); continue
        seen.add(rid); r = {key: None for key in KEYS}; r.update(olds.get(rid, {})); r['id'] = rid
        for key in present:
            if key == 'id': continue
            raw = row.get(key); kind = COL[key][2]; lab = COL[key][1]
            if isinstance(raw, str): raw = raw.strip()
            if kind in ('text', 'long'): r[key] = '' if raw is None else str(raw)
            elif raw in (None, ''): r[key] = None
            elif kind == 'choice':
                s = str(raw).lower()
                if s in COL[key][3]: r[key] = s
                else: errors.append('%s / %s: "%s" is not one of %s' % (where, lab, raw, ', '.join(COL[key][3])))
            else:
                n = num(raw); lo, hi = COL[key][3]
                if n is None: errors.append('%s / %s: "%s" is not a number' % (where, lab, raw))
                elif n < lo or n > hi: errors.append('%s / %s: %s is outside %s to %s' % (where, lab, fmt(n), fmt(lo), fmt(hi)))
                elif not float(n).is_integer(): errors.append('%s / %s: %s must be a whole number' % (where, lab, fmt(n)))
                else: r[key] = int(n)
        for key in ('name', 'weapon', 'target', 'count', 'time'):
            if r.get(key) in (None, ''): errors.append('%s: %s is missing' % (where, COL[key][1]))
        if r.get('reward') is None: r['reward'] = 0
        if r.get('notes') is None: r['notes'] = ''
        if r.get('target') and r.get('count') and r.get('time'):
            each = r['time'] / r['count']
            if (r['target'] == 'people' and each < 3) or (r['target'] == 'cars' and each < 8):
                warns.append('%s: %.1f s per %s is very tight' % (where, each, 'person' if r['target'] == 'people' else 'vehicle'))
        out.append({key: r.get(key) for key in KEYS})
    if len(out) > MAX_N: errors.append('At most %d rampages (there are %d)' % (MAX_N, len(out)))
    if not out: errors.append('No rampages left - every row needs an ID')
    gone = [k for k in olds if k not in seen]
    if gone: warns.append('Removed (their rows are gone): ' + ', '.join(gone))
    return out, errors, warns


def changes(old, new):
    o = {r['id']: r for r in old}; out = []
    for r in new:
        if r['id'] not in o: out.append('new rampage %s (%s)' % (r['id'], r['name'])); continue
        for k in KEYS:
            if r.get(k) != o[r['id']].get(k): out.append('%s / %s: %s -> %s' % (r['id'], k, o[r['id']].get(k), r.get(k)))
    if [r['id'] for r in new] != [r['id'] for r in old if r['id'] in {x['id'] for x in new}]: out.append('the order changed')
    return out


def do_import(path, dry):
    s, i, j, t = read_table(); rows, present = read_rows(path)
    rs, errors, warns = build(rows, present, t['rampages'])
    for w in warns: print('note: ' + w)
    if errors: print('Not imported - fix these first:'); [print('  ' + e) for e in errors]; sys.exit(1)
    ch = changes(t['rampages'], rs); print('%d change%s' % (len(ch), '' if len(ch) == 1 else 's') + (':' if ch else '')); [print('  ' + c) for c in ch]
    if not dry and ch: open(JS, 'w', encoding='utf-8').write(s[:i] + dump(rs) + s[j:]); print('written', os.path.relpath(JS))


if __name__ == '__main__':
    if len(sys.argv) < 3 or sys.argv[1] not in ('export', 'import'): print(__doc__); sys.exit(2)
    if sys.argv[1] == 'export':
        t = read_table()[3]; (export_numbers if sys.argv[2].lower().endswith('.numbers') else export_xlsx)(sys.argv[2], t); print('saved', sys.argv[2])
    else: do_import(sys.argv[2], '--dry-run' in sys.argv[3:])
