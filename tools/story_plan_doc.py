#!/usr/bin/env python3
"""The story plan as a Word file (.docx, opens in Pages), made from docs/STORY-PLAN.md.

  python3 tools/story_plan_doc.py              write docs/Story-Plan.docx

The plan is written in docs/STORY-PLAN.md; this turns its headings, paragraphs, bullets and tables into a landscape Word page
(**bold** stays bold, `code` becomes plain). Needs python-docx."""
import os, re
import docx
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'docs', 'STORY-PLAN.md')
OUT = os.path.join(ROOT, 'docs', 'Story-Plan.docx')


def runs(p, text, size=None, bold=False):     # a line's words into a paragraph: **bold** parts bold, `code` plain
    for i, part in enumerate(re.split(r'\*\*(.+?)\*\*', text.replace('`', ''))):
        if not part: continue
        r = p.add_run(part); r.bold = bold or i % 2 == 1
        if size: r.font.size = Pt(size)


def shade(cell, colour):                       # a cell's background
    pr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), colour); pr.append(sh)


def table(d, rows, width=26.5):               # a markdown table: the first row is the header; wordier columns get more room
    head = rows[0]; body = [r for r in rows[2:]]
    t = d.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'; t.autofit = False
    long = [max(len(h), 4 + sum(len(r[k]) for r in body) / max(1, len(body))) ** 0.8 for k, h in enumerate(head)]
    cm = [width * w / sum(long) for w in long]
    for c, h in zip(t.rows[0].cells, head):
        c.paragraphs[0].text = ''; runs(c.paragraphs[0], h, 9, True); shade(c, 'D9E2F3')
    for r in body:
        cells = t.add_row().cells
        for c, v in zip(cells, r):
            c.paragraphs[0].text = ''; runs(c.paragraphs[0], v, 9)
    for col, w in zip(t.columns, cm): col.width = Cm(w)
    for row in t.rows:
        for c, w in zip(row.cells, cm): c.width = Cm(w)
    d.add_paragraph()


def build():
    d = docx.Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, side, Cm(1.6))
    st = d.styles['Normal']; st.font.name = 'Helvetica'; st.font.size = Pt(10.5)
    lines = open(SRC, encoding='utf-8').read().split('\n')
    i = 0; para = []

    def flush():
        if para:
            p = d.add_paragraph(); runs(p, ' '.join(para)); para.clear()

    while i < len(lines):
        ln = lines[i].rstrip()
        if ln.startswith('|'):
            flush(); rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]); i += 1
            table(d, rows); continue
        m = re.match(r'(#{1,3}) (.*)', ln)
        if m:
            flush(); h = d.add_heading(level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0)
            runs(h, m.group(2))
            for r in h.runs: r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)
        elif re.match(r'\s*(- |\d+\. )', ln):
            flush(); num = re.match(r'\s*\d+\. ', ln)
            text = re.sub(r'^\s*(- |\d+\. )', '', ln)
            while i + 1 < len(lines) and lines[i + 1].startswith('  ') and not re.match(r'\s*(- |\d+\. )', lines[i + 1]):
                i += 1; text += ' ' + lines[i].strip()
            p = d.add_paragraph(style='List Number' if num else 'List Bullet'); runs(p, text)
        elif not ln.strip():
            flush()
        else:
            para.append(ln.strip())
        i += 1
    flush()
    d.save(OUT)
    print('wrote', os.path.relpath(OUT, ROOT))


if __name__ == '__main__':
    build()
