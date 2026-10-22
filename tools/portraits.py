#!/usr/bin/env python3
"""The characters' pixel-art portraits, drawn from their recipes in js/01y-character-data.js (the character table).

  python3 tools/portraits.py                           draw every portrait: img/chars/<id>-calm.png, -angry.png, -smile.png (48 x 75)
  python3 tools/portraits.py sheet out.png [N] [ids]   a contact sheet (all, or some ids: danny,joe), N times bigger (default 4), on the game's ground
  python3 tools/portraits.py doc                       the Word sheet docs/Characters.docx: each portrait and how the character looks

The style is the owner's reference picture (October 2026): head and shoulders, 48 x 75, a 3/4 view with the ear on the left and the
face turned to the right, lit from the upper right, transparent around the figure (the game draws the night-dark ground and a glow in
the character's bubble colour behind it). The head, the reference's own hair and the shirt are hand-made parts taken from that picture
(tools/portrait_parts.py: letters for colour ramps); every character recolours them from its recipe - skin, hair and eye colours, the
top's two colours - and adds its own parts on top: a hair style (drawn in clumps and strands like the reference's), facial hair, a hat,
glasses, the top and extras (wrinkles, an earring, a chain, a scar, a gold tooth, sweat, tired eyes, a toothpick, hoop earrings,
lipstick, a tattoo, pearls, a cigarette holder, a pencil behind the ear, a bow tie, grey temples, jowls, gold braid, a jeweller's loupe,
a cigar, a rag over the shoulder). Women (told by their hair style - pixie, bob, bun, pulled, wave - or lipstick) get a finer face:
thin brows, lashes, rosy lips, a blush and a softer chin. The three faces - calm, angry, smiling - differ in the brows, eyes and mouth.
Needs Pillow (and python-docx for the Word sheet)."""
import json, os, re, sys
from math import atan2, cos, sin, hypot, pi
from PIL import Image
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import portrait_parts as PARTS

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'js', '01y-character-data.js')
OUT = os.path.join(ROOT, 'img', 'chars')
W, H = 48, 75
FACES = ('calm', 'angry', 'smile')
HAIR = ['short', 'slick', 'buzz', 'combed', 'pixie', 'curly', 'sides', 'bald', 'flattop', 'bob', 'bun', 'pulled', 'wave', 'spiky']
FACIAL = ['none', 'stubble', 'beard', 'moustache', 'pencil', 'goatee']
HATS = ['none', 'captain', 'hardhat', 'police', 'flatcap', 'chauffeur', 'visor']
GLASSES = ['none', 'reading', 'shades', 'square', 'roundshades']
TOPS = ['openshirt', 'denim', 'flowershirt', 'jacketvest', 'darkshirt', 'waistcoat', 'suit', 'hivis', 'uniform', 'jumpsuit',
        'cardigan', 'leather', 'blouse', 'blazer', 'tanktop', 'labcoat', 'dress', 'buttoned', 'seersucker', 'checkjacket', 'windbreaker']
EXTRAS = ['wrinkles', 'earring', 'chain', 'scar', 'goldtooth', 'sweat', 'tired', 'toothpick', 'hoops', 'lipstick', 'tattoo', 'pearls',
          'holder', 'pencil', 'bowtie', 'temples', 'jowls', 'braid', 'loupe', 'cigar', 'rag']
AGES = ['young', 'mid', 'old']
BUILDS = ['slim', 'normal', 'broad']
SMILE_EYES = ['happy', 'open']
CALM_MOUTH = ['line', 'smirk', 'frown']
FEM_HAIR = ('pixie', 'bob', 'bun', 'pulled', 'wave')


def load():
    s = open(DATA, encoding='utf-8').read()
    return json.loads(re.search(r'/\*CHARACTERS-JSON\*/(.*?)/\*END-CHARACTERS-JSON\*/', s, re.S).group(1))


def rgb(h): h = h.lstrip('#'); return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))
def mix(a, b, t): return tuple(max(0, min(255, round(x + (y - x) * t))) for x, y in zip(a, b))
def dark(c, k): return tuple(max(0, min(255, round(x * k))) for x in c)
def lite(c, t): return mix(c, (255, 255, 255), t)
def mul(c, k): return tuple(max(0, min(255, round(x * f))) for x, f in zip(c, k))


def part(p):                                 # a hand-made part: {(x, y): letter}
    top, text = p; lines = text.split('\n')
    if lines and lines[0] == '': lines = lines[1:]
    return {(x, top + i): ch for i, line in enumerate(lines) for x, ch in enumerate(line) if ch != '.'}


HEAD, HAIR_SHORT, BODY = part(PARTS.HEAD_M), part(PARTS.HAIR_SHORT), part(PARTS.BODY)
SKULL = (24.0, 21.5, 14.2, 15.6)              # the skull under the hair: centre and radii
def in_skull(x, y, grow=0.0):
    cx, cy, rx, ry = SKULL
    return ((x + 0.5 - cx) / (rx + grow)) ** 2 + ((y + 0.5 - cy) / (ry + grow)) ** 2 <= 1


# ---- colour ramps, in the reference's own proportions
def skin_ramp(base):
    b = mix(base, (255, 214, 176), 0.12)                       # a little warmer and lighter, like the reference's lit skin
    return {'O': mul(b, (0.30, 0.19, 0.22)), '0': mul(b, (0.47, 0.37, 0.41)), '1': mul(b, (0.63, 0.52, 0.57)), '2': mul(b, (0.75, 0.71, 0.77)),
            '3': mul(b, (0.86, 0.82, 0.87)), '4': b, '5': mix(b, (255, 238, 222), 0.4)}


def hair_ramp(h):
    lum = h[0] * 0.3 + h[1] * 0.59 + h[2] * 0.11
    if lum < 70:                              # dark hair: near-black with a cool blue-grey sheen (the reference)
        return {'a': mix(dark(h, 0.3), (6, 4, 12), 0.6), 'e': dark(h, 0.62), 'b': h, 'c': mix(h, (92, 98, 124), 0.26), 'd': mix(h, (136, 146, 178), 0.48)}
    return {'a': mix(dark(h, 0.36), (24, 18, 30), 0.4), 'e': dark(h, 0.72), 'b': h, 'c': lite(h, 0.2), 'd': lite(h, 0.48)}


def cloth_ramp(c, letters='ADEFG'):
    r = {'A': mul(c, (0.22, 0.21, 0.30)), 'D': mul(c, (0.54, 0.51, 0.64)), 'E': mul(c, (0.82, 0.80, 0.86)), 'F': c, 'G': mix(c, (255, 255, 255), 0.2)}
    if max(c) > 225 and min(c) > 200:         # white cloth: grey-blue shades, not muddy
        r.update({'A': (70, 72, 90), 'D': (150, 154, 172), 'E': (196, 198, 212)})
    return dict(zip(letters, (r['A'], r['D'], r['E'], r['F'], r['G'])))


class Pic:                                    # the picture as it is built: (x, y) -> colour; each part paints over the last
    def __init__(self): self.p = {}
    def put(self, x, y, c):
        if 0 <= x < W and 0 <= y < H and c is not None: self.p[(x, y)] = c
    def stamp(self, px, pal, dx=0, dy=0, only=None):
        for (x, y), ch in px.items():
            if ch in pal and (only is None or only(x + dx, y + dy)): self.put(x + dx, y + dy, pal[ch])
    def get(self, x, y): return self.p.get((x, y))


def poly_px(pts):
    ys = [p[1] for p in pts]; out = set()
    for y in range(int(min(ys)) - 1, int(max(ys)) + 2):
        xs = []
        for (x1, y1), (x2, y2) in zip(pts, pts[1:] + pts[:1]):
            if (y1 <= y + 0.5 < y2) or (y2 <= y + 0.5 < y1): xs.append(x1 + (y + 0.5 - y1) * (x2 - x1) / (y2 - y1))
        xs.sort()
        for a, b in zip(xs[::2], xs[1::2]):
            for x in range(int(round(a)), int(round(b))): out.add((x, y))
    return out


def ellipse_px(cx, cy, rx, ry):
    return {(x, y) for y in range(int(cy - ry) - 1, int(cy + ry) + 2) for x in range(int(cx - rx) - 1, int(cx + rx) + 2)
            if ((x + 0.5 - cx) / rx) ** 2 + ((y + 0.5 - cy) / ry) ** 2 <= 1}


def edge(px, x, y): return any((x + dx, y + dy) not in px for dx, dy in ((1, 0), (-1, 0), (0, 1), (0, -1)))


# ---- the face: the reference's features are the smiling face; patches give the calm and angry ones, the women's finer face
def face_patch(g, face, ch, fem, extras):
    def setrow(y, x0, s):
        for i, c in enumerate(s):
            if c != ' ': g[(x0 + i, y)] = c
    brow_px = [(x, y) for (x, y), c in g.items() if c == 'B']
    if face == 'angry':                       # the brows down to the nose, a crease between them, the lids lower; the mouth open, its corners down
        for (x, y) in brow_px: g[(x, y)] = g.get((x, y + 3), '3') if g.get((x, y + 3), '3') in '012345' else '3'
        for x, y in [(15, 25), (16, 25), (17, 25), (16, 26), (17, 26), (18, 26), (19, 26), (18, 27), (19, 27), (20, 27), (21, 27), (22, 27), (21, 28), (22, 28), (23, 28), (24, 28), (25, 28),
                     (37, 25), (38, 25), (35, 26), (36, 26), (37, 26), (33, 27), (34, 27), (35, 27), (31, 28), (32, 28), (33, 28), (34, 28)]:
            g[(x, y)] = 'B'
        g[(28, 27)] = '2'; g[(28, 28)] = '1'
        setrow(30, 18, 'LL'); setrow(30, 32, 'L')
        setrow(40, 22, '444MMMMMMMM4'); setrow(41, 22, '4Mnttttttnn0'); setrow(42, 22, 'M4nnnnnnnn4M'); setrow(43, 23, '4MMMMMMMM3')
    elif face == 'calm':                      # the smile straightened (the corner no longer curls up)
        cm = ch.get('calmMouth', 'line')
        setrow(40, 22, '344'); setrow(41, 24, '00000000'); g[(32, 41)] = '2'
        if cm == 'smirk': setrow(40, 31, '00'); g[(32, 41)] = '0'
        elif cm == 'frown': g[(23, 42)] = '0'; g[(32, 42)] = '0'; setrow(41, 24, '30000000')
    else:                                     # smiling: the reference's own mouth
        if ch.get('smileEyes', 'happy') == 'happy':              # the eyes narrowed by the cheeks
            setrow(31, 18, '33333'); setrow(31, 31, '3333'); setrow(32, 18, '44444'); setrow(32, 31, '44444')
        toothy = 'goldtooth' in extras or ch['id'] in ('tito', 'nico', 'rocco', 'sol')
        if toothy: setrow(41, 25, 'tttttt'); setrow(42, 25, '000000')
        if 'goldtooth' in extras: g[(28, 41)] = 'g'
    if fem:                                   # thin arched brows, lashes at the outer corners, a lighter, smaller nose, full lips, a blush
        brows = [(x, y) for (x, y), c in g.items() if c == 'B']
        for (x, y) in brows:                  # one pixel thick: only the top of each column stays
            if any((x, yy) in brows for yy in range(y - 2, y)): g[(x, y)] = '3'
        for (x, y) in ((24, 25), (23, 25)) if face != 'angry' else ():
            if g.get((x, y)) == 'B': g[(x, y)] = '3'
        for x, y in ((17, 29), (16, 28), (15, 28), (37, 29), (38, 28), (17, 30), (37, 30)): g[(x, y)] = 'L'
        for (x, y), c in list(g.items()):
            if 26 <= x <= 33 and 33 <= y <= 38 and c in '012': g[(x, y)] = {'0': '1', '1': '2', '2': '3'}[c]
        for x in (26, 27, 32, 33):
            if g.get((x, 38)) in ('1', '2'): g[(x, 38)] = '3'
        for (x, y) in [(x, y) for x in range(18, 22) for y in (35, 36)] + [(x, y) for x in range(33, 36) for y in (34, 35)]:
            if g.get((x, y)) in ('3', '4'): g[(x, y)] = 'k'
        if face != 'angry':
            for x in range(24, 32):
                if g.get((x, 41)) in ('0', 'M'): g[(x, 41)] = 'M'
            for x in range(25, 31):
                if g.get((x, 40)) not in ('t',): g[(x, 40)] = 'm'
                if g.get((x, 42)) not in ('t', '0'): g[(x, 42)] = 'm'
            g[(24, 40)] = 'M' if face == 'smile' else g[(24, 40)]
    return g


def frow(left, right, inner):                 # a hand-drawn row of the women's face: the outline at both ends
    assert len(inner) == right - left - 1
    return {left: 'O', right: 'O', **{left + 1 + i: c for i, c in enumerate(inner)}}
FEM_JAW = {                                   # a narrower, rounder chin (rows 41-48)
    41: frow(13, 36, '233' + '4444444' + 'M' + '0000000' + '2' + '44' + '3'),
    42: frow(13, 36, '233' + '44444444' + '444444' + '4444' + '3'),
    43: frow(14, 35, '234' + '44444444' + '2222' + '4444' + '3'),
    44: frow(14, 34, '123' + '4' * 11 + '333' + '32'),
    45: frow(15, 33, '12' + '3' + '4' * 9 + '333' + '21'),
    46: frow(16, 32, '12' + '3' + '4' * 7 + '33' + '221'),
    47: frow(17, 31, '11' + '22' + '33333' + '222' + '1'),
    48: {x: 'O' for x in range(18, 31)},
}


def head_letters(ch, face, fem, extras):
    g = dict(HEAD)
    if fem:
        for y, cols in FEM_JAW.items():
            for x in range(10, 40): g.pop((x, y), None)
            for x, c in cols.items(): g[(x, y)] = c
    old = ch['age'] == 'old' or 'wrinkles' in extras
    if old:                                   # lines: the forehead, the eyes' corners, under the eyes, from the nose to the mouth
        for x in range(21, 34):
            if x % 3: g[(x, 19 + (x // 5) % 2)] = '2'
        for x, y in ((15, 31), (16, 32), (38, 30), (37, 32), (19, 33), (20, 33), (21, 33), (32, 33), (33, 33), (34, 33), (23, 39), (22, 40), (22, 41), (34, 39), (35, 40)):
            if g.get((x, y)) in ('2', '3', '4', '5'): g[(x, y)] = '1' if g[(x, y)] in '23' else '2'
    if 'tired' in extras:
        for x in range(18, 23): g[(x, 33)] = '2'
        for x in range(31, 36): g[(x, 33)] = '2'
    if 'jowls' in extras:
        for x in range(17, 33):
            if g.get((x, 47)) in ('2', '3', '4'): g[(x, 47)] = '1'
    face_patch(g, face, ch, fem, extras)
    if 'scar' in extras:                      # down the near cheek under the eye
        for x, y in ((14, 33), (15, 34), (15, 35), (16, 36), (16, 37)): g[(x, y)] = 's'
    return g


def stubble(pic, fc, face):
    """a shadow of beard over the jaw, the chin and the upper lip: speckled, denser along the jaw (the reference's)"""
    mouth = {(x, y) for x in range(23, 33) for y in (40, 41, 42)} | ({(x, y) for x in range(22, 34) for y in range(40, 44)} if face == 'angry' else set())
    for y in range(38, 50):
        for x in range(10, 38):
            c = pic.get(x, y)
            if c is None or HEAD.get((x, y)) in (None, 'O') or (x, y) in mouth: continue
            jaw = y >= 44 or x <= 15 or (y >= 40 and (x <= 20 or x >= 33))
            lip = y == 39 and 23 <= x <= 32
            if not (jaw or lip): continue
            dense = y >= 46 or x <= 14 or (y >= 44 and (x <= 17 or x >= 31))
            pic.put(x, y, mix(c, fc, (0.42 if (x + y) % 2 else 0.34) if dense else (0.3 if (x + y) % 3 else 0.22)))   # a soft shadow, denser along the jaw


# ---- hair: masks for each style, shaded in clumps and strands from a crown, the outline dark
def hair_mask(hs):
    """front: over the head; back: behind the head and neck; crown: where the strands start; flow: how they run"""
    front, back = set(), set()
    sk = lambda grow: {(x, y) for y in range(0, 50) for x in range(0, W) if in_skull(x, y, grow)}
    def hairline(y_of_x, grow, side_l=28, side_r=25):           # the skull's cap down to a hairline over the forehead; the sides to the ear
        out = set()
        for (x, y) in sk(grow):
            if y <= y_of_x(x): out.add((x, y))
            elif x <= 14 and y <= side_l: out.add((x, y))
            elif x >= 37 and y <= side_r: out.add((x, y))
        return out
    crown, flow = (20, 6), 'fall'
    if hs == 'slick':
        front = hairline(lambda x: 17 + (0 if x > 24 else 1), 1.2); flow, crown = 'back', (34, 14)
    elif hs == 'combed':
        front = hairline(lambda x: 17, 1.6) | {(x, y) for (x, y) in ellipse_px(27, 14, 13, 8.5) if y <= 18}; flow, crown = 'part', (17, 7)
    elif hs == 'buzz':
        front = hairline(lambda x: 17 + (1.5 if x < 17 or x > 33 else 0), 0.3, 26, 23); flow = 'buzz'
    elif hs == 'pixie':
        front = hairline(lambda x: 17 + max(0, min(7, (x - 16) * 0.45)), 1.6, 32, 26); flow, crown = 'part', (15, 8)
        front |= {(x, y) for y in range(33, 38) for x in range(10, 14 - (y - 33) // 2)}           # a lock in front of the ear
        front |= {p for p in poly_px([(3, 22), (7, 11), (16, 6), (14, 44), (8, 45), (4, 36)]) if p not in HEAD}   # the nape, past the jaw
    elif hs == 'curly':
        for k in range(18):
            a = pi * (0.95 + k / 17 * 1.1)
            front |= ellipse_px(24 + cos(a) * 15.5, 21 + sin(a) * 16, 4.2, 3.8)
        front |= hairline(lambda x: 18 + (1 if (x // 3) % 2 else 0), 2.5, 32, 28)
        front = {(x, y) for (x, y) in front if y <= 34 and not (14 < x < 37 and y > 20)}
        flow = 'curl'
    elif hs == 'sides':
        front = {(x, y) for (x, y) in sk(0.8) if (x <= 14 and 16 <= y <= 31) or (x >= 36 and 17 <= y <= 27)}
    elif hs == 'flattop':
        front = poly_px([(10, 7), (11.5, 3), (37, 3), (38.5, 7), (39, 18), (9.5, 18)]) | hairline(lambda x: 17, 0.5, 26, 23)
        flow = 'up'
    elif hs == 'bob':                         # a blunt fringe at the brows, straight to the jaw, behind the neck
        front = hairline(lambda x: 25, 1.6) | {(x, y) for y in range(18, 48) for x in range(5, 14) if x >= 5 + max(0, (y - 42) // 2)} | {(x, y) for y in range(19, 47) for x in range(37, 42)}
        back = poly_px([(5, 21), (8, 10), (16, 5), (32, 5), (40, 10), (43, 21), (43, 48), (5, 48)])
    elif hs in ('bun', 'pulled'):
        front = hairline(lambda x: 16.5, 0.8, 27, 24); flow, crown = 'back', (14, 13)
        if hs == 'bun':                       # a messy bun on top, loose curls at the temples
            front |= ellipse_px(22, 5.5, 7.5, 5.0)
            front |= {(x, y) for y in range(25, 38) for x in (10, 11) if (y // 2 + x) % 3} | {(12, y) for y in range(25, 31)}
            front |= {(x, y) for y in range(21, 30) for x in (37, 38) if (y + x) % 3}
            front |= {p for p in poly_px([(3, 22), (7, 11), (16, 6), (14, 41), (9, 43), (5, 38), (3, 30)]) if p not in HEAD}   # curls fallen loose behind the ear
        else: back = ellipse_px(8, 16, 5.0, 5.4)
    elif hs == 'wave':                        # a set 50s wave: full round the face to the jaw, rolls over the brow
        front = hairline(lambda x: 16 + 2.5 * sin((x - 8) / 4.5), 2.6) | {(x, y) for y in range(19, 43) for x in range(4, 14)} | {(x, y) for y in range(17, 39) for x in range(37, 43)}
        back = poly_px([(3, 23), (6, 9), (16, 3), (32, 3), (42, 9), (45, 23), (44, 43), (4, 43)])
        flow = 'wave'
    elif hs == 'spiky':
        front = hairline(lambda x: 17 + (1 if x % 5 == 0 else 0), 0.6, 27, 24)
        for i, (x0, h) in enumerate(((11, 6), (15, 9), (20, 11), (25, 11), (30, 9), (35, 7), (39, 4))):
            front |= poly_px([(x0 - 2.6, 10), (x0 + 0.5 + i * 0.3, 8 - h), (x0 + 2.6, 10)])
        flow = 'up'
    return front, back, crown, flow


def flow_dir(flow, crown, x, y):              # which way the strands run at a point
    cx, cy = crown
    if flow == 'back': a = atan2(-0.25 + (y - 14) * 0.05, -1.0)              # from the brow back over the head, falling at the back
    elif flow == 'up': return (0.15, -1.0)
    elif flow == 'wave': return (1.0, 0.45 * cos(x / 3.0))
    else: a = atan2(y - cy + 0.01, x - cx)
    return (cos(a), sin(a))


def shade_hair(px, crown, flow):
    """the hair's letters, like the reference's: broad tone bands (lit upper right c, body b, shade e), long strand lines traced along
    the flow, a sheen streak (d) beside the strands where the light falls, the outline (a)"""
    out = {}
    lit = lambda x, y: (x - 24) * 0.55 - (y - 10) * 0.85
    for (x, y) in px:
        l = lit(x, y); out[(x, y)] = 'c' if l > 7 else 'b' if l > -10 else 'e'
    if flow == 'curl':                                              # tiled curls, every other row shifted
        sprite = ['.cc.', 'cbbe', 'cbbe', '.ee.']
        for (x, y) in px:
            row = y // 4; sx = (x + (2 if row % 2 else 0)) % 4; ch = sprite[y % 4][sx]
            if ch == '.': ch = 'e'
            if ch == 'c' and lit(x, y) > 0: ch = 'd'
            out[(x, y)] = ch
    elif flow == 'buzz':                                            # cropped close: fine speckle, lit on top
        for (x, y) in px:
            if (x * 3 + y * 5) % 7 == 0: out[(x, y)] = 'c' if lit(x, y) > -6 else 'e'
            elif lit(x, y) > 8 and (x + y) % 3 == 0: out[(x, y)] = 'd'
    else:
        xs = [x for x, y in px]; ys = [y for x, y in px]
        roots = [(gx + ((gx * 7 + gy * 13) % 3) - 1, gy + ((gx * 5 + gy * 11) % 3) - 1)
                 for gy in range(min(ys), max(ys) + 1, 5) for gx in range(min(xs) + (gy // 5 % 2) * 2, max(xs) + 1, 4)]
        for (rx, ry) in roots:
            if (rx, ry) not in px: continue
            for sign in (1, -1):                                    # trace the strand both ways from its root
                x, y = float(rx), float(ry)
                for step in range(9):
                    fx, fy = flow_dir(flow, crown, x, y); x += fx * sign; y += fy * sign
                    q = (int(round(x)), int(round(y)))
                    if q not in px: break
                    l = lit(*q)
                    out[q] = 'e' if l < 4 else 'b'
                    if l > -2 and 2 <= step <= 6 and sign == 1:     # a sheen streak beside the strand, on the light's side
                        n = (q[0] + (1 if fy > 0 else 0), q[1] - (1 if abs(fx) > 0.5 else 0))
                        if n in px and out.get(n) != 'e': out[n] = 'd' if l > 3 else 'c'
    beside = lambda x, y: (x, y) in px or (y >= 30 and (x, y) in HEAD)          # hair down past the jaw lies on the head: no outline there
    for (x, y) in px:
        if any(not beside(x + dx, y + dy) for dx, dy in ((1, 0), (-1, 0), (0, 1), (0, -1))): out[(x, y)] = 'a'
    return out


# ---- the top: the reference's shirt recoloured, then each top's own collar, fronts and pattern
def draw_top(pic, ch, sk, extras):
    t = ch['top']; top = rgb(ch['topColour']); top2 = rgb(ch['topColour2']) if ch.get('topColour2') else dark(top, 0.8)
    C = cloth_ramp(top); C2 = cloth_ramp(top2)
    shirt = BODY
    cloth = {p for p, l in shirt.items() if l in 'ADEFG'}
    under = {p for p, l in shirt.items() if l in 'UV'}
    V = poly_px([(19.5, 49), (31.5, 49), (27, 66), (24, 66)])
    if t in ('tanktop', 'dress'):             # bare shoulders: the shirt's shape in skin, then the garment
        for (x, y), l in shirt.items():
            pic.put(x, y, sk[{'A': 'O', 'D': '1', 'E': '2', 'F': '3', 'G': '4', 'U': '3', 'V': '2'}[l]])
        if t == 'tanktop':
            vest = poly_px([(13, 52), (17, 50), (21, 60), (30, 60), (34, 50), (38, 52), (40, 75), (11, 75)])
            for (x, y) in vest: pic.put(x, y, C[{'A': 'E', 'D': 'E', 'E': 'E', 'F': 'F', 'G': 'G', 'U': 'F', 'V': 'E'}.get(shirt.get((x, y)), 'F')])
            for (x, y) in vest:
                if edge(vest, x, y) and y < 74: pic.put(x, y, C['D'])
        else:
            gown = poly_px([(0, 61), (14, 59), (24, 62), (36, 59), (48, 61), (48, 75), (0, 75)])
            for (x, y) in gown: pic.put(x, y, C[{'A': 'D', 'D': 'E', 'E': 'E', 'F': 'F', 'G': 'G'}.get(shirt.get((x, y)), 'F')])
            for x in range(48):
                if (x, 60) in gown: pic.put(x, 60, C['D'])
            for x0 in (13, 35):
                for y in range(50, 61): pic.put(x0, y, C['F']); pic.put(x0 + 1, y, C['D'])
            for x in (6, 15, 30, 41):
                for y in range(64, 75): pic.put(x, y, C2['F'])
        draw_top_extras(pic, t, extras); return
    for (x, y) in cloth: pic.put(x, y, C[shirt[(x, y)]])
    for (x, y) in under: pic.put(x, y, (236, 236, 238) if shirt[(x, y)] == 'U' else (176, 178, 190))
    closed = t in ('buttoned', 'uniform', 'jumpsuit', 'windbreaker', 'hivis')
    jacket = t in ('suit', 'seersucker', 'checkjacket', 'jacketvest', 'blazer', 'labcoat', 'leather', 'cardigan', 'waistcoat')
    if t == 'flowershirt':
        for (cx, cy) in ((4, 62), (11, 70), (17, 58), (36, 60), (42, 68), (33, 72), (6, 72), (40, 54), (21, 73), (2, 54)):
            for dx, dy in ((0, 0), (1, 0), (-1, 0), (0, 1), (0, -1), (1, 1), (-1, -1)):
                if (cx + dx, cy + dy) in cloth: pic.put(cx + dx, cy + dy, (255, 170, 50) if (dx, dy) == (0, 0) else C2['F'] if (dx + dy) % 2 == 0 else C2['G'])
    if t == 'checkjacket':
        for (x, y) in cloth:
            if x % 4 == 0 or y % 4 == 0: pic.put(x, y, C2['F'] if (x % 4 == 0) != (y % 4 == 0) else C2['D'])
    if t == 'seersucker':
        for (x, y) in cloth:
            if x % 2 == 0: pic.put(x, y, lite(pic.get(x, y), 0.3))
    if t == 'cardigan':
        for (x, y) in cloth:
            if (x + (y // 2) * 2) % 4 == 0 and y % 2: pic.put(x, y, C['E'])
    if closed:                                # the collar shut: cloth over the open neck, a placket down the middle
        for (x, y) in V | under:
            if y >= 51: pic.put(x, y, C['F'] if x < 26 else C['E'])
        for y in range(51, 75): pic.put(25, y, C['D']); pic.put(26, y, C['E'])
        if t in ('buttoned', 'uniform'):
            for y in range(53, 75, 4): pic.put(25, y, (226, 226, 232))
        if t == 'buttoned':                   # a number on the pocket
            for (x, y) in poly_px([(7, 61), (14, 61), (14, 66), (7, 66)]): pic.put(x, y, C2['F'] if y > 61 else C2['D'])
            for x in (8, 9, 11, 12): pic.put(x, 63, (232, 232, 238))
        if t == 'uniform':                    # a tie, a badge, epaulettes (gold braid for the chief)
            for (x, y) in poly_px([(24, 51), (27, 51), (28, 68), (25.5, 71), (23, 68)]): pic.put(x, y, C2['F'] if x <= 25 else C2['E'])
            for (x, y) in poly_px([(8, 61), (13, 61), (13, 66), (10.5, 68), (8, 66)]): pic.put(x, y, (236, 196, 80) if (x + y) % 3 else (255, 232, 140))
            ep = (220, 180, 70) if 'braid' in extras else C['D']
            for (x, y) in poly_px([(3, 55), (13, 52), (13, 54.5), (3, 57.5)]) | poly_px([(36, 52), (46, 55), (46, 57.5), (36, 54.5)]): pic.put(x, y, ep if 'braid' not in extras or (x + y) % 2 else lite(ep, 0.35))
        if t == 'jumpsuit':                   # the zip, a name patch, grease
            for y in range(52, 75): pic.put(25, y, (176, 178, 190) if y % 2 else C['D'])
            for (x, y) in poly_px([(6, 61), (17, 61), (17, 66), (6, 66)]): pic.put(x, y, (238, 234, 222))
            for x in range(7, 16): pic.put(x, 63, (40, 56, 110) if x % 2 else (238, 234, 222))
            for x, y in ((10, 70), (11, 71), (36, 66), (37, 67), (30, 72)): pic.put(x, y, mix(top, (40, 36, 30), 0.5))
        if t == 'windbreaker':                # a high zipped collar, a band across the chest
            for (x, y) in cloth:
                if 60 <= y <= 64: pic.put(x, y, C2[shirt[(x, y)]])
            for (x, y) in poly_px([(17, 46), (34, 46), (35, 53), (16, 53)]): pic.put(x, y, C['G'] if x < 26 else C['E'])
            for y in range(46, 75): pic.put(25, y, (214, 216, 226) if y % 2 else (150, 152, 166))
        if t == 'hivis':                      # a navy shirt under the orange vest, silver bands
            vest = poly_px([(14, 52), (21, 54), (24, 75), (5, 75), (6, 57)]) | poly_px([(37, 52), (30, 54), (27, 75), (44, 75), (44, 57)])
            for (x, y) in cloth - vest: pic.put(x, y, C2[shirt[(x, y)]])
            for (x, y) in vest: pic.put(x, y, C[shirt[(x, y)]] if shirt.get((x, y)) in C else C['F'])
            for y in (63, 64, 70, 71):
                for x in range(W):
                    if (x, y) in vest: pic.put(x, y, (228, 232, 238) if (x + y) % 3 else (252, 252, 255))
    if jacket:                                # a jacket: the shirt (second colour) in the open neck, lapels down to the first button
        inner = C2 if t != 'checkjacket' else cloth_ramp((238, 234, 222))
        depth = {'leather': 72, 'labcoat': 70, 'blazer': 70, 'jacketvest': 72, 'cardigan': 66, 'waistcoat': 64}.get(t, 66)
        vee = poly_px([(18, 49), (33, 49), (27, depth), (24, depth)])
        for (x, y) in vee: pic.put(x, y, inner['F'] if x < 26 else inner['E'])
        if t in ('suit', 'seersucker', 'checkjacket', 'cardigan', 'waistcoat'):     # the shirt open at the neck: skin, its collar points
            for (x, y) in poly_px([(21, 49), (30, 49), (26.5, 57), (25, 57)]):
                l = HEAD.get((x, y)); pic.put(x, y, sk[l] if l in sk else sk['2'])
            for (x, y) in poly_px([(20, 49), (23, 49), (25.5, 56), (21, 54)]) | poly_px([(31, 49), (28, 49), (25.8, 56), (30, 54)]): pic.put(x, y, inner['G'] if x < 26 else inner['F'])
        elif t in ('leather', 'labcoat', 'blazer'):                                  # a T-shirt or top: a round neck
            for (x, y) in poly_px([(20, 49), (31, 49), (30, 53), (21, 53)]):
                l = HEAD.get((x, y)); pic.put(x, y, sk[l] if l in sk else sk['2'])
            for x in range(20, 32): pic.put(x, 53, inner['E']); pic.put(x, 54, inner['F'])
        elif t == 'jacketvest':                                                      # a white vest, a scoop of skin above it
            for (x, y) in vee: pic.put(x, y, (240, 240, 236) if x < 26 else (210, 212, 220))
            for (x, y) in ellipse_px(25.5, 49, 5, 5):
                l = HEAD.get((x, y)); pic.put(x, y, sk[l] if l in sk else sk['2'])
        L = poly_px([(16, 48), (20, 49), (25.5, depth), (21, depth - 6), (13, 53)]); R = poly_px([(35, 48), (31, 49), (25.5, depth), (30, depth - 6), (38, 53)])
        for (x, y) in L: pic.put(x, y, C['G'] if not edge(L, x, y) else C['D'])
        for (x, y) in R: pic.put(x, y, C['E'] if not edge(R, x, y) else C['A'])
        for y in range(depth + 2, 75, 5): pic.put(25, y, C['A'])
        if t == 'waistcoat':                  # white shirt sleeves at the sides, the waistcoat on the chest
            for (x, y) in cloth:
                if x <= 9 or x >= 41: pic.put(x, y, C2[shirt[(x, y)]])
        if t == 'leather':                    # shine on the leather, the open zip
            for x, y in ((4, 57), (5, 57), (6, 56), (3, 62), (3, 63), (40, 58), (41, 59), (9, 52), (10, 52)): pic.put(x, y, C['G'])
            for y in range(64, 75): pic.put(20, y, (150, 150, 162) if y % 3 == 0 else C['D'])
        if t == 'labcoat':                    # pens in the breast pocket
            for (x, y) in poly_px([(33, 62), (40, 62), (40, 68), (33, 68)]): pic.put(x, y, C['E'] if y == 62 else C['F'])
            for x, y, c in ((34, 61, (60, 90, 200)), (34, 60, (60, 90, 200)), (36, 61, (220, 40, 40)), (36, 60, (220, 40, 40))): pic.put(x, y, c)
        if t == 'blazer':                     # the padded shoulders
            for x in range(1, 47):
                if (x, 54) in cloth: pic.put(x, 54, C['G'] if x < 30 else C['F'])
        if t == 'cardigan':
            for y in (68, 71, 74): pic.put(25, y, (226, 206, 160))
    if t == 'blouse':
        for x, y in ((9, 62), (10, 63), (38, 62), (39, 63), (6, 70), (42, 70)): pic.put(x, y, C['E'])
    draw_top_extras(pic, t, extras)


def draw_top_extras(pic, t, extras):
    if 'chain' in extras:
        ring = poly_px([(19, 50), (32, 50), (30, 57), (25.5, 59), (21, 57)]) - poly_px([(20, 50), (31, 50), (29.4, 55.8), (25.5, 57.6), (21.6, 55.8)])
        for (x, y) in ring: pic.put(x, y, (255, 224, 110) if x % 2 else (214, 164, 54))
    if 'pearls' in extras:
        for i, (x, y) in enumerate(((17, 51), (18, 52), (19, 53), (20, 54), (21, 55), (22, 55), (23, 56), (24, 56), (25, 56), (26, 56), (27, 56), (28, 55), (29, 55), (30, 54), (31, 53), (32, 52), (33, 51))):
            pic.put(x, y, (252, 250, 246) if i % 2 else (196, 192, 212))
    if 'bowtie' in extras or t == 'waistcoat':
        bt = (24, 22, 30) if t == 'waistcoat' else (150, 30, 44)
        for (x, y) in poly_px([(20, 49), (25.5, 51.5), (31, 49), (31, 55), (25.5, 52.5), (20, 55)]): pic.put(x, y, bt if x not in (25, 26) else lite(bt, 0.3))
    if 'tattoo' in extras:                    # on the near shoulder
        for x, y in ((5, 58), (6, 57), (7, 58), (6, 59), (5, 60), (7, 60), (6, 61), (8, 59), (4, 59)): pic.put(x, y, (40, 92, 104))
    if 'rag' in extras:
        for (x, y) in poly_px([(3, 55), (12, 51), (14, 55), (9, 66), (4, 67)]): pic.put(x, y, (204, 70, 60) if (x + y) % 4 else (150, 48, 42))


def draw_facial(pic, ch, face, fc):
    f = ch['facial']; fr = hair_ramp(fc)
    if f in ('moustache', 'goatee', 'beard'):                  # over the upper lip
        mo = poly_px([(21, 39), (33, 39), (34, 42), (31, 41), (23, 41), (20, 42)])
        for (x, y) in mo: pic.put(x, y, fr['a'] if edge(mo, x, y) and y >= 41 else fr['c'] if (x + y) % 3 == 0 else fr['b'])
    if f == 'pencil':
        for x in range(23, 32): pic.put(x, 40, fr['a'] if x in (23, 31) else fr['b'])
    if f == 'goatee':
        gt = poly_px([(22, 44), (31, 44), (30.5, 49.5), (26.5, 51.5), (22.5, 49.5)])
        for (x, y) in gt:
            if face == 'angry' and y <= 44: continue
            pic.put(x, y, fr['a'] if edge(gt, x, y) else fr['c'] if (x * 2 + y) % 5 == 0 else fr['b'])
    if f == 'beard':                          # a full short beard on the jaw and chin, the mouth clear
        bd = poly_px([(9.5, 31), (12.5, 31), (15, 39), (21, 43), (31, 43), (35, 38), (37, 31), (38.5, 32), (36, 44), (31, 50), (26, 52), (19, 50), (14, 46)])
        for (x, y) in bd:
            if 22 <= x <= 32 and 41 <= y <= 44: continue
            k = (x * 3 + y * 5) % 7
            pic.put(x, y, fr['a'] if edge(bd, x, y) else fr['e'] if k < 2 else fr['d'] if (k == 3 and x > 26 and y < 45) else fr['c'] if k == 4 else fr['b'])


def draw_glasses(pic, ch):
    gl = ch.get('glasses', 'none')
    if gl == 'none': return
    gc = rgb(ch['glassesColour']) if ch.get('glassesColour') else (30, 26, 34); g_lt = lite(gc, 0.5)
    if gl == 'shades':
        for (x, y) in poly_px([(16, 28), (25, 28), (24.5, 33.5), (17, 33.5)]) | poly_px([(29.5, 28), (38.5, 28), (38, 33.5), (30.5, 33.5)]): pic.put(x, y, gc)
        for x, y in ((18, 29), (19, 29), (18, 30), (31, 29), (32, 29), (31, 30)): pic.put(x, y, (156, 176, 214))
        for x, y in ((22, 32), (36, 32)): pic.put(x, y, (96, 116, 156))
        for x in range(25, 30): pic.put(x, 29, gc)
        for x in range(11, 16): pic.put(x, 29, gc)
    elif gl == 'roundshades':
        for (x, y) in ellipse_px(20.5, 31, 4.8, 4.2) | ellipse_px(33.5, 31, 4.4, 4.2): pic.put(x, y, gc)
        for (x, y) in ellipse_px(20.5, 31, 3.6, 3.0) | ellipse_px(33.5, 31, 3.2, 3.0): pic.put(x, y, mix(gc, (40, 30, 50), 0.5))
        for x, y in ((19, 29), (20, 28), (32, 29), (33, 28)): pic.put(x, y, (196, 206, 236))
        for x in range(25, 30): pic.put(x, 30, gc)
        for x in range(11, 16): pic.put(x, 30, gc)
    elif gl == 'square':
        for x0, x1 in ((15, 25), (29, 39)):
            for x in range(x0, x1 + 1): pic.put(x, 27, gc); pic.put(x, 34, gc)
            for y in range(27, 35): pic.put(x0, y, gc); pic.put(x1, y, gc)
            pic.put(x0 + 1, 28, g_lt); pic.put(x0 + 2, 28, g_lt)
        for x in range(26, 29): pic.put(x, 29, gc)
        for x in range(10, 15): pic.put(x, 29, gc)
    elif gl == 'reading':                     # half-moon frames low on the nose, a cord to the collar
        for x0, x1 in ((16, 24), (30, 38)):
            for x in range(x0, x1 + 1): pic.put(x, 34, gc)
            for y in (32, 33): pic.put(x0, y, gc); pic.put(x1, y, gc)
            for x in range(x0 + 1, x1): pic.put(x, 33, mix(pic.get(x, 33) or (200, 160, 130), (214, 234, 250), 0.35))
        for x in range(25, 30): pic.put(x, 33, gc)
        for x in range(11, 16): pic.put(x, 32, gc)
        for y in range(33, 53): pic.put(11 - (y - 33) // 7, y, gc if y % 2 else dark(gc, 0.7))


def draw_hat(pic, ch):
    hat = ch.get('hat', 'none')
    if hat == 'none': return
    hc = rgb(ch['hatColour']) if ch.get('hatColour') else (40, 40, 50); hc2 = rgb(ch['hatColour2']) if ch.get('hatColour2') else dark(hc, 0.6)
    T = cloth_ramp(hc)
    def crown(px):
        for (x, y) in px:
            c = T['F']
            if x <= 14: c = T['E']
            if x <= 11: c = T['D']
            if (x, y - 1) not in px and x > 14: c = T['G']
            if edge(px, x, y): c = T['A']
            pic.put(x, y, c)
    visor = lambda: [pic.put(x, y, (12, 12, 18) if y > 20 else (50, 52, 68)) for (x, y) in poly_px([(13, 20), (41, 20), (39, 24), (16, 24)])]
    if hat == 'captain':
        crown(poly_px([(8, 20), (8.5, 11), (14, 5), (34, 5), (40, 10), (41, 20)]))
        for x in range(15, 34): pic.put(x, 6, lite(hc, 0.5))
        for x in range(9, 41): pic.put(x, 17, (196, 160, 70) if x % 2 else (240, 204, 110)); pic.put(x, 18, (24, 24, 32)); pic.put(x, 19, (16, 16, 22))
        visor()
        for x, y in ((25, 9), (25, 10), (25, 11), (25, 12), (25, 13), (24, 10), (26, 10), (23, 13), (27, 13), (24, 14), (25, 14), (26, 14)): pic.put(x, y, (236, 190, 70))
        pic.put(25, 8, (255, 230, 140))
    elif hat in ('police', 'chauffeur'):
        crown(poly_px([(6.5, 5), (41.5, 5), (38.5, 19), (9.5, 19)]))
        for x in range(10, 39): pic.put(x, 18, (18, 18, 24)); pic.put(x, 19, hc2 if hat == 'police' else (18, 18, 24))
        if hat == 'police':
            for x in range(10, 39, 2): pic.put(x, 19, (230, 230, 236))
            for x, y in ((23, 11), (24, 11), (25, 11), (23, 12), (24, 12), (25, 12), (24, 13), (24, 10)): pic.put(x, y, (236, 196, 80) if y < 13 else (200, 150, 50))
        else: pic.put(36, 18, (236, 190, 70)); pic.put(12, 18, (236, 190, 70))
        visor()
    elif hat == 'flatcap':
        px = poly_px([(7, 21), (8, 12), (13, 6.5), (32, 5.5), (39, 8.5), (42, 17), (42, 22)]); crown(px)
        for (x, y) in px:
            if (x * 2 + y * 3) % 6 == 0 and not edge(px, x, y): pic.put(x, y, hc2)
        for (x, y) in poly_px([(14, 21), (43, 20), (41, 24), (17, 24)]): pic.put(x, y, T['E'] if y < 22 else T['D'])
        for x in range(15, 31): pic.put(x, 7, T['G'])
    elif hat == 'hardhat':
        crown({(x, y) for (x, y) in ellipse_px(24.5, 19, 16.5, 15.5) if y <= 19})
        for y in range(5, 19): pic.put(24, y, lite(hc, 0.5)); pic.put(25, y, T['G']); pic.put(26, y, T['E'])
        C2 = cloth_ramp(hc2)
        for (x, y) in poly_px([(5, 19), (45, 19), (46, 22), (4, 22)]): pic.put(x, y, C2['F'] if y == 19 else C2['E'] if y < 21 else C2['D'])
    elif hat == 'visor':                      # a green eyeshade: a band round the head, a see-through peak
        for x in range(9, 40): pic.put(x, 20, hc2); pic.put(x, 21, dark(hc2, 0.8))
        for (x, y) in poly_px([(13, 22), (42, 22), (43, 27), (15, 27)]):
            under = pic.get(x, y) or (200, 160, 130)
            pic.put(x, y, mix(under, hc, 0.6 if (x + y) % 2 else 0.75))
        for x in range(14, 42): pic.put(x, 22, lite(hc, 0.35))


def draw_mouth_props(pic, extras):
    if 'cigar' in extras:
        for x, y in ((32, 42), (33, 42), (34, 42), (35, 41), (36, 41), (37, 41), (38, 41)): pic.put(x, y, (124, 80, 44)); pic.put(x, y + 1, (86, 54, 30))
        pic.put(39, 41, (255, 120, 50)); pic.put(39, 42, (200, 70, 40))
        for x, y in ((40, 38), (41, 36), (40, 34), (41, 32)): pic.put(x, y, (170, 170, 182))
    if 'toothpick' in extras:
        for i, (x, y) in enumerate(((32, 42), (33, 42), (34, 41), (35, 41), (36, 40), (37, 40))): pic.put(x, y, (228, 198, 142) if i < 5 else (190, 160, 110))
    if 'holder' in extras:
        for x, y in ((32, 42), (33, 42), (34, 41), (35, 41), (36, 40), (37, 40), (38, 39), (39, 39), (40, 38)): pic.put(x, y, (24, 22, 28))
        pic.put(41, 38, (240, 236, 228)); pic.put(42, 37, (240, 236, 228)); pic.put(43, 37, (255, 120, 60))
        for x, y in ((44, 34), (43, 32), (44, 30), (43, 28)): pic.put(x, y, (170, 170, 186))


def draw(ch, face):
    extras = set(ch.get('extras', [])); hs = ch['hairStyle']; fem = hs in FEM_HAIR or 'lipstick' in extras
    sk = skin_ramp(rgb(ch['skin']))
    hair = rgb(ch['hair']) if ch.get('hair') else (40, 30, 30); hr = hair_ramp(hair)
    fcol = rgb(ch.get('facialColour') or ch['hair'])
    eyes = rgb(ch['eyes'])
    pal = dict(sk)
    brow = hr['a'] if sum(hair) < 200 else dark(hair, 0.62)
    if hs == 'bald': brow = dark(fcol, 0.7)
    lip_line = mix(sk['0'], (96, 20, 26), 0.35); lip = mix(sk['3'], (206, 96, 96), 0.38)
    if fem: lip_line, lip = mix(sk['0'], (150, 30, 56), 0.55), mix(sk['4'], (214, 92, 112), 0.5)     # rosy lips without lipstick too
    if 'lipstick' in extras: lip_line, lip = (140, 14, 40), (206, 34, 64)
    pal.update({'B': brow, 'L': (16, 10, 16), 'W': (236, 232, 228), 'I': mix(eyes, (255, 255, 255), 0.25), 'J': dark(eyes, 0.45),
                'M': lip_line, 'm': lip, 'k': mix(sk['4'], (236, 128, 138), 0.2), 't': (244, 240, 230), 'n': (84, 18, 28), 'g': (255, 210, 70), 's': mix(sk['4'], (236, 180, 176), 0.6)})
    pic = Pic()
    front, back, crown, flow = hair_mask(hs) if hs not in ('short', 'bald') else (set(), set(), None, None)
    if back:                                  # hair behind the head and neck, in shade
        bh = shade_hair(back, (24, 10), 'fall' if hs != 'wave' else 'wave')
        pic.stamp(bh, {'a': hr['a'], 'e': hr['a'], 'b': hr['e'], 'c': hr['b'], 'd': hr['c']})
    pic.stamp(HEAD, sk)                       # the neck under the collar
    draw_top(pic, ch, sk, extras)
    g = head_letters(ch, face, fem, extras)
    pic.stamp({p: l for p, l in g.items() if p[1] < 50}, pal)
    if ch['facial'] == 'stubble' and not fem: stubble(pic, fcol, face)
    draw_facial(pic, ch, face, fcol)
    if 'earring' in extras: pic.put(7, 37, (255, 214, 90)); pic.put(7, 38, (196, 146, 46))
    if hs == 'short':                         # the reference's own hair (under a hat: only what shows at the sides)
        hat = ch.get('hat', 'none') != 'none'
        pic.stamp(HAIR_SHORT, hr, only=(lambda x, y: 18 <= y <= 30 and (x <= 14 or x >= 37)) if hat else None)
    elif hs == 'bald':
        for x, y in ((28, 9), (29, 9), (30, 10), (31, 10), (27, 10), (30, 11)): pic.put(x, y, sk['5'])
    elif front: pic.stamp(shade_hair(front, crown, flow), hr)
    if 'temples' in extras:
        for (x, y), c in list(pic.p.items()):
            if 10 <= x <= 15 and 16 <= y <= 29 and c in (hr['b'], hr['c'], hr['d'], hr['e']): pic.put(x, y, (170, 170, 178) if c != hr['e'] else (120, 120, 130))
    if 'pencil' in extras:                    # behind the near ear
        for i, (x, y) in enumerate(((3, 26), (4, 27), (5, 28), (6, 29), (7, 30), (8, 31), (9, 32))): pic.put(x, y, (240, 130, 150) if i == 0 else (246, 200, 40) if i < 6 else (60, 50, 40))
    draw_glasses(pic, ch)
    draw_hat(pic, ch)
    draw_mouth_props(pic, extras)
    if 'hoops' in extras:
        for (x, y) in ellipse_px(7.5, 41.5, 2.6, 3.4) - ellipse_px(7.5, 41.5, 1.4, 2.2): pic.put(x, y, (255, 206, 70) if y < 42 else (210, 160, 40))
    if 'sweat' in extras and face != 'smile':
        for x, y, c in ((36, 21, (200, 236, 255)), (36, 22, (150, 205, 245)), (13, 24, (200, 236, 255)), (13, 25, (150, 205, 245))): pic.put(x, y, c)
    if 'loupe' in extras:
        for (x, y) in ellipse_px(33, 21, 3.2, 3.0): pic.put(x, y, (34, 34, 40))
        for (x, y) in ellipse_px(33, 21, 2.0, 1.8): pic.put(x, y, (170, 205, 230) if x < 33 else (120, 155, 190))
    return finish(pic)


def finish(pic):                              # the outline round the figure, very dark, tinted by what it edges
    ink = (12, 8, 18)
    img = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    for (x, y), c in pic.p.items(): img.putpixel((x, y), tuple(c) + (255,))
    for y in range(H):
        for x in range(W):
            if (x, y) in pic.p: continue
            nb = [pic.p[q] for q in ((x + 1, y), (x - 1, y), (x, y + 1), (x, y - 1)) if q in pic.p]
            if nb: img.putpixel((x, y), mix(dark(min(nb, key=sum), 0.25), ink, 0.6) + (255,))
    return img


def draw_all():
    t = load(); os.makedirs(OUT, exist_ok=True); n = 0
    for ch in t['characters']:
        check(ch)
        for f in FACES: draw(ch, f).save(os.path.join(OUT, '%s-%s.png' % (ch['id'], f))); n += 1
    print('drew %d portraits in %s' % (n, os.path.relpath(OUT, ROOT)))


def check(ch):                               # the recipe's words must be ones the tool draws
    for key, allowed in (('hairStyle', HAIR), ('facial', FACIAL), ('hat', HATS), ('glasses', GLASSES), ('top', TOPS), ('age', AGES), ('build', BUILDS),
                         ('smileEyes', SMILE_EYES), ('calmMouth', CALM_MOUTH)):
        if ch.get(key) not in allowed: sys.exit('%s: %s "%s" is not one of %s' % (ch['id'], key, ch.get(key), ', '.join(allowed)))
    for e in ch.get('extras', []):
        if e not in EXTRAS: sys.exit('%s: extra "%s" is not one of %s' % (ch['id'], e, ', '.join(EXTRAS)))


def on_ground(ch, img):                      # a portrait as the game shows it: on the night-dark ground with a glow in the bubble colour
    bub = rgb(ch['bubble']); g = Image.new('RGBA', (W, H))
    for y in range(H):
        for x in range(W):
            d = ((x - W / 2) ** 2 + ((y - H * 0.4) * 0.8) ** 2) ** 0.5
            g.putpixel((x, y), mix((18, 15, 30), bub, max(0.0, 1 - d / (H * 0.5)) * 0.45) + (255,))
    g.alpha_composite(img); return g.convert('RGB')


def sheet(path, k=4, ids=None):
    t = load(); chs = [c for c in t['characters'] if not ids or c['id'] in ids]; pad = 4
    img = Image.new('RGB', (len(FACES) * (W * k + pad) + pad, len(chs) * (H * k + pad) + pad), (30, 28, 40))
    for r, ch in enumerate(chs):
        check(ch)
        for c, f in enumerate(FACES):
            img.paste(on_ground(ch, draw(ch, f)).resize((W * k, H * k), Image.NEAREST), (pad + c * (W * k + pad), pad + r * (H * k + pad)))
    img.save(path); print('wrote', path)


def doc():
    import docx
    from docx.shared import Pt, Cm
    t = load(); d = docx.Document()
    st = d.styles['Normal']; st.font.name = 'Helvetica'; st.font.size = Pt(10)
    d.add_heading('BLOCK RUNNER - THE CHARACTERS', 0)
    d.add_paragraph('How every character looks, written down so they stay the same in every mission. The portraits are pixel art, 48 x 75, '
                    'in the style of the reference picture: head and shoulders in a 3/4 view, the ear on the left, lit from the upper right, '
                    'transparent around the figure (the game shows them on a night-dark ground with a glow in the character\'s bubble colour, as '
                    'here). Each has three faces: calm, angry and smiling - the game picks one from the mood written on each line '
                    '(angry: ' + ', '.join(t['faces']['angry']) + '; smiling: ' + ', '.join(t['faces']['smile']) + '; anything else: calm).')
    d.add_paragraph('They are drawn by tools/portraits.py from the recipes in the character table (js/01y-character-data.js; to change one, '
                    'edit the table - or the exported spreadsheet - and run the tool again). A picture of your own in img/chars/ with the same '
                    'name replaces a drawn one.')
    tmp = os.path.join(OUT, '_doc'); os.makedirs(tmp, exist_ok=True)
    for ch in t['characters']:
        d.add_heading('%s  (%s)' % (ch['name'], ch['letter']), 1)
        row = d.add_paragraph()
        for f in FACES:
            p = os.path.join(tmp, '%s-%s.png' % (ch['id'], f)); on_ground(ch, draw(ch, f)).resize((W * 4, H * 4), Image.NEAREST).save(p)
            row.add_run().add_picture(p, width=Cm(2.6)); row.add_run('  ')
        d.add_paragraph('calm - angry - smiling').runs[0].font.size = Pt(8)
        tb = d.add_table(rows=0, cols=2); tb.style = 'Light Grid Accent 1'
        def line(k, v):
            r = tb.add_row().cells; r[0].text = k; r[1].text = v
        line('Who', ch['role']); line('Looks', ch['look'])
        line('Bubble colour', ch['bubble']); line('In the city', 'shirt %s, trousers %s' % (ch['shirt'], ch['pants']))
        line('Portrait recipe', 'skin %s; hair %s (%s); facial hair: %s; eyes %s; hat: %s; glasses: %s; top: %s (%s / %s); extras: %s; %s, %s build' % (
            ch['skin'], ch['hair'], ch['hairStyle'], ch['facial'], ch['eyes'], ch['hat'], ch['glasses'], ch['top'], ch['topColour'], ch['topColour2'],
            ', '.join(ch['extras']) or 'none', ch['age'], ch['build']))
        line('Mumble', '%s wave, %d-%d Hz, %d syllables a second, pauses %s s' % (ch['mumble']['wave'], ch['mumble']['low'], ch['mumble']['high'], ch['mumble']['rate'], ch['mumble']['gap']))
        for r in tb.rows: r.cells[0].width = Cm(3.5); r.cells[1].width = Cm(13)
    path = os.path.join(ROOT, 'docs', 'Characters.docx'); d.save(path)
    for f in os.listdir(tmp): os.remove(os.path.join(tmp, f))
    os.rmdir(tmp); print('wrote', os.path.relpath(path, ROOT))


if __name__ == '__main__':
    a = sys.argv[1:]
    if not a: draw_all()
    elif a[0] == 'sheet': sheet(a[1], int(a[2]) if len(a) > 2 else 4, a[3].split(',') if len(a) > 3 else None)
    elif a[0] == 'doc': doc()
    else: print(__doc__)
