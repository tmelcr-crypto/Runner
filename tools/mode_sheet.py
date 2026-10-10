"""Game modes table <-> spreadsheet (Apple Numbers or Excel): which features FREE ROAM and STORY have.

  python3 tools/mode_sheet.py export modes.numbers            write the features of js/01j-mode-data.js as a Numbers file (or .xlsx)
  python3 tools/mode_sheet.py import modes.numbers [--dry-run] read an edited file back, check it, write the yes / no values

One row per feature, a yes / no column per mode. Rows are found by the key in the last column; only the yes / no cells are written
back (a new feature needs building first, so a new row is left out). A second sheet, Settings, holds the numbers that differ by mode
(a value per mode, checked against the range allowed). Needs openpyxl for .xlsx and numbers-parser for .numbers."""
import json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__)); JS = os.path.join(HERE, '..', 'js', '01j-mode-data.js')
A, B = '/*MODE-JSON*/', '/*END-MODE-JSON*/'
TIPS = [('Yellow cells', 'yes: the feature is in that mode; no: it is left out. Pick from the list.', 'yes or no'),
        ('New features', 'Every new feature gets a row here, with your answer to: free roam, story or both?', ''),
        ('Story', 'Story mode is built later; until then it shows as COMING SOON in the game.', ''),
        ('Settings sheet', 'Numbers that differ by mode, e.g. how many hidden pickups lie out (% of the weapon and item tables). Type a number in the range shown.', 'a number'),
        ('Keys', 'The grey last column tells the game which feature a row is. Do not change it.', ''),
        ('Sending it back', 'Upload the file in the chat as it is (.numbers or .xlsx).', '')]


def read_table():
    s = open(JS, encoding='utf-8').read(); i, j = s.index(A) + len(A), s.index(B); return s, i, j, json.loads(s[i:j])


def heads(t): return ['Feature'] + [m['name'].title() for m in t['modes']] + ['Notes', 'Key (do not change)']
def sheads(t): return ['Setting'] + [m['name'].title() for m in t['modes']] + ['Unit', 'Allowed', 'Notes', 'Key (do not change)']
fmt = lambda v: ('%d' % v) if float(v).is_integer() else ('%g' % v)


def export_xlsx(dest, t):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    F = 'Arial'; f_in = PatternFill('solid', fgColor='FFF7CC'); f_head = PatternFill('solid', fgColor='2B2442'); f_lab = PatternFill('solid', fgColor='E4E0F2'); f_key = PatternFill('solid', fgColor='EDEDED')
    thin = Side(style='thin', color='C9C9C9'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook(); sh = wb.active; sh.title = 'Modes'; M = t['modes']; H = heads(t)
    for j, h in enumerate(H, start=1):
        c = sh.cell(row=1, column=j, value=h); c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = f_head; c.border = box; c.alignment = Alignment(horizontal='center')
    dv = DataValidation(type='list', formula1='"yes,no"', allow_blank=False, showErrorMessage=True, error='yes or no'); sh.add_data_validation(dv)
    for i, f in enumerate(t['features'], start=2):
        vals = [f['name']] + ['yes' if f.get(m['id']) else 'no' for m in M] + [f.get('note', ''), f['id']]
        for j, v in enumerate(vals, start=1):
            c = sh.cell(row=i, column=j, value=v); c.border = box; mode = 1 < j <= 1 + len(M)
            c.font = Font(name=F, size=10, bold=j == 1, color='0000FF' if mode else '595959' if j == len(H) else '000000')
            c.fill = f_lab if j == 1 else f_in if mode else f_key if j == len(H) else PatternFill(); c.alignment = Alignment(horizontal='center' if mode else 'left', vertical='center', wrap_text=j == len(H) - 1)
            if mode: dv.add(c.coordinate)
    for col, w in zip('ABCDEFG', [44] + [14] * len(M) + [60, 18]): sh.column_dimensions[col].width = w
    sh.freeze_panes = 'B2'
    st = wb.create_sheet('Settings'); SH = sheads(t)
    for j, h in enumerate(SH, start=1):
        c = st.cell(row=1, column=j, value=h); c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = f_head; c.border = box; c.alignment = Alignment(horizontal='center')
    for i, r in enumerate(t.get('settings', []), start=2):
        v = DataValidation(type='decimal', operator='between', formula1=fmt(r['min']), formula2=fmt(r['max']), allow_blank=False, showErrorMessage=True, error='A number from %s to %s' % (fmt(r['min']), fmt(r['max']))); st.add_data_validation(v)
        vals = [r['name']] + [r.get(m['id']) for m in M] + [r['unit'], '%s to %s' % (fmt(r['min']), fmt(r['max'])), r.get('note', ''), r['id']]
        for j, x in enumerate(vals, start=1):
            c = st.cell(row=i, column=j, value=x); c.border = box; mode = 1 < j <= 1 + len(M)
            c.font = Font(name=F, size=10, bold=j == 1, color='0000FF' if mode else '595959' if j == len(SH) else '000000')
            c.fill = f_lab if j == 1 else f_in if mode else f_key if j == len(SH) else PatternFill(); c.alignment = Alignment(horizontal='center' if mode else 'left', vertical='center', wrap_text=j == len(SH) - 1)
            if mode: v.add(c.coordinate)
    for col, w in zip('ABCDEFGHI', [44] + [14] * len(M) + [8, 12, 60, 18]): st.column_dimensions[col].width = w
    g = wb.create_sheet('How to fill'); g.column_dimensions['A'].width = 20; g.column_dimensions['B'].width = 80
    for i, row in enumerate([('Item', 'What it means', 'How to fill')] + TIPS, start=1):
        for j, v in enumerate(row, start=1): c = g.cell(row=i, column=j, value=v); c.font = Font(name=F, size=10, bold=i == 1 or j == 1); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = box
    wb.save(dest)


def export_numbers(dest, t):
    from numbers_parser import Alignment, Document, RGB
    sys.path.insert(0, HERE); from settings_sheet import exact_numbers; exact_numbers()
    M = t['modes']; H = heads(t); fs = t['features']
    doc = Document(sheet_name='Modes', table_name='Modes', num_header_rows=1, num_header_cols=1, num_rows=len(fs) + 1, num_cols=len(H)); tb = doc.sheets[0].tables[0]
    S = {}
    def style(name, **kw):
        if name not in S: S[name] = doc.add_style(name=name, font_name='Helvetica Neue', font_size=kw.pop('size', 11.0), alignment=Alignment(kw.pop('h', 'left'), kw.pop('v', 'middle')), **kw)
        return S[name]
    head = style('BR head', bg_color=RGB(43, 36, 66), font_color=RGB(255, 255, 255), bold=True, h='center'); lab = style('BR label', bg_color=RGB(228, 224, 242), bold=True)
    val = style('BR value', bg_color=RGB(255, 247, 204), font_color=RGB(0, 0, 255), h='center'); note = style('BR note', size=10.0)
    key = style('BR key', bg_color=RGB(237, 237, 237), font_color=RGB(89, 89, 89), size=10.0)
    for j, w in enumerate([280] + [90] * len(M) + [380, 120]): tb.col_width(j, w)
    for j, h in enumerate(H): tb.write(0, j, h, style=head)
    for i, f in enumerate(fs, start=1):
        tb.write(i, 0, f['name'], style=lab)
        for k, m in enumerate(M, start=1):
            tb.write(i, k, 'yes' if f.get(m['id']) else 'no', style=val); tb.set_cell_formatting(i, k, 'popup', popup_values=['yes', 'no'], allow_none=False)
        tb.write(i, len(M) + 1, f.get('note', ''), style=note); tb.write(i, len(M) + 2, f['id'], style=key); tb.row_height(i, 30)
    SH, SS = sheads(t), t.get('settings', [])
    if SS:
        doc.add_sheet('Settings', 'Settings', num_rows=len(SS) + 1, num_cols=len(SH)); st = doc.sheets[-1].tables[0]; st.num_header_rows = 1; st.num_header_cols = 1
        for j, w in enumerate([280] + [90] * len(M) + [60, 90, 380, 120]): st.col_width(j, w)
        for j, h in enumerate(SH): st.write(0, j, h, style=head)
        for i, r in enumerate(SS, start=1):
            st.write(i, 0, r['name'], style=lab)
            for k, m in enumerate(M, start=1): st.write(i, k, r.get(m['id']), style=val)
            st.write(i, len(M) + 1, r['unit'], style=note); st.write(i, len(M) + 2, '%s to %s' % (fmt(r['min']), fmt(r['max'])), style=note)
            st.write(i, len(M) + 3, r.get('note', ''), style=note); st.write(i, len(M) + 4, r['id'], style=key); st.row_height(i, 40)
    doc.add_sheet('How to fill', 'How to fill', num_rows=len(TIPS) + 1, num_cols=3)
    g = doc.sheets[-1].tables[0]; g.num_header_rows = 1; g.num_header_cols = 1; g.col_width(0, 150); g.col_width(1, 460); g.col_width(2, 110)
    for j, h in enumerate(['Item', 'What it means', 'How to fill']): g.write(0, j, h, style=head)
    for i, row in enumerate(TIPS, start=1): g.write(i, 0, row[0], style=lab); g.write(i, 1, row[1], style=note); g.write(i, 2, row[2], style=note); g.row_height(i, 40)
    doc.save(dest)


def read_sheets(src):
    if src.lower().endswith('.numbers'):
        from numbers_parser import Document
        return {sh.name: [tuple(r) for r in sh.tables[0].rows(values_only=True)] for sh in Document(src).sheets}
    from openpyxl import load_workbook
    return {ws.title: list(ws.iter_rows(values_only=True)) for ws in load_workbook(src, data_only=True).worksheets}


def read_settings(sheets, t):              # {key: {mode: value}} from the Settings sheet, or None when the file has none
    data = next((r for n, r in sheets.items() if n.lower().startswith('setting')), None)
    if not data: return None
    hd = [str(h or '').strip().lower() for h in data[0]]
    col = {m['id']: hd.index(m['name'].lower()) for m in t['modes'] if m['name'].lower() in hd}
    kc = next((j for j, h in enumerate(hd) if h.startswith('key')), None)
    if kc is None: return None
    return {str(r[kc]).strip(): {m: (r[j] if j < len(r) else None) for m, j in col.items()} for r in data[1:] if kc < len(r) and r[kc]}


def read_rows(src, t, sheets=None):
    sheets = sheets or read_sheets(src)
    data = next((r for n, r in sheets.items() if n.lower().startswith('mode')), None) or next(iter(sheets.values()))
    hd = [str(h or '').strip().lower() for h in data[0]]
    col = {m['id']: hd.index(m['name'].lower()) for m in t['modes'] if m['name'].lower() in hd}
    kc = next((j for j, h in enumerate(hd) if h.startswith('key')), None)
    if kc is None or not col: raise SystemExit('Needs a column per mode (%s) and a Key column.' % ', '.join(m['name'].title() for m in t['modes']))
    return [{'key': r[kc] if kc < len(r) else None, **{m: (r[j] if j < len(r) else None) for m, j in col.items()}} for r in data[1:]]


def num(v):
    if isinstance(v, bool) or v is None: return None
    if isinstance(v, (int, float)): v = round(float(v), 6); return int(v) if v.is_integer() else v
    try: return num(float(str(v).strip().replace(',', '.')))
    except ValueError: return None


def do_import(src, dry):
    s, i, j, t = read_table(); F = {f['id']: f for f in t['features']}; errors, notes, ch = [], [], []; sheets = read_sheets(src); sch = []
    S = {r['id']: r for r in t.get('settings', [])}; got = read_settings(sheets, t)
    for k, vals in (got or {}).items():
        if k not in S: notes.append('Settings: %s is not a setting the game has; left out' % k); continue
        r = S[k]
        for m, raw in vals.items():
            v = num(raw)
            if v is None: errors.append('Settings / %s / %s: "%s" is not a number' % (r['name'], m, raw)); continue
            if v < r['min'] or v > r['max']: errors.append('Settings / %s / %s: %s is outside %s to %s' % (r['name'], m, fmt(v), fmt(r['min']), fmt(r['max']))); continue
            if v != r.get(m): sch.append((k, m, v))
    for n, r in enumerate(read_rows(src, t, sheets), start=2):
        k = str(r['key'] or '').strip()
        if not k: continue
        if k not in F: notes.append('Row %d: %s is not a feature the game has yet - build it first; left out' % (n, k)); continue
        for m in (x['id'] for x in t['modes']):
            if m not in r: continue
            v = r[m]; v = v if isinstance(v, bool) else str(v or '').strip().lower()
            if v not in (True, False, 'yes', 'no', 'true', 'false'): errors.append('Row %d (%s) / %s: "%s" is not yes or no' % (n, F[k]['name'], m, r[m])); continue
            b = v is True or v in ('yes', 'true')
            if b != bool(F[k].get(m)): ch.append((k, m, b))
    for x in notes: print('note: ' + x)
    if errors: print('Not imported - fix these first:'); [print('  ' + e) for e in errors]; sys.exit(1)
    n = len(ch) + len(sch); print('%d change%s' % (n, '' if n == 1 else 's') + (':' if n else ''))
    for k, m, b in ch: print('  %s / %s: %s' % (F[k]['name'], m, 'yes' if b else 'no'))
    for k, m, v in sch: print('  Settings / %s / %s: %s -> %s' % (S[k]['name'], m, fmt(S[k][m]), fmt(v)))
    if dry or not n: return
    lines = s[i:j].split('\n')
    for k, m, v in sch:
        hit = [x for x, l in enumerate(lines) if '"id": %s,' % json.dumps(k) in l and '"min"' in l]
        lines[hit[0]] = re.sub(r'("%s": )(-?[0-9.]+)' % m, lambda q: q.group(1) + fmt(v), lines[hit[0]], count=1)
    for k, m, b in ch:
        hit = [x for x, l in enumerate(lines) if '"id": %s,' % json.dumps(k) in l]
        lines[hit[0]] = re.sub(r'("%s": )(true|false)' % m, lambda q: q.group(1) + ('true' if b else 'false'), lines[hit[0]], count=1)
    out = s[:i] + '\n'.join(lines) + s[j:]; json.loads(out[i:out.index(B)]); open(JS, 'w', encoding='utf-8').write(out); print('written', os.path.relpath(JS))


if __name__ == '__main__':
    if len(sys.argv) < 3 or sys.argv[1] not in ('export', 'import'): print(__doc__); sys.exit(2)
    if sys.argv[1] == 'export': (export_numbers if sys.argv[2].lower().endswith('.numbers') else export_xlsx)(sys.argv[2], read_table()[3]); print('saved', sys.argv[2])
    else: do_import(sys.argv[2], '--dry-run' in sys.argv[3:])
