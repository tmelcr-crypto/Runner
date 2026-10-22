"""Character table <-> spreadsheet (Apple Numbers or Excel): every character of the missions (js/01y-character-data.js).

  python3 tools/characters_sheet.py export characters.numbers            write the table as a Numbers file (or .xlsx)
  python3 tools/characters_sheet.py import characters.numbers [--dry-run] read an edited file back, check it, rewrite the table

Sheets: Characters (one row per character: the name and the letter on the maps, the bubble colour, the clothes their figure wears in
the city, who they are and how they look, the recipe their pixel-art portrait is drawn from, their mumble), Faces (the moods that pick
the angry and the smiling face; any other mood: calm) and How to fill. Rows are found by the ID column: a new row with a new ID is a
new character; a row taken out is gone (not if a mission still has them). After an import the portraits are drawn again
(tools/portraits.py). Needs openpyxl for .xlsx and numbers-parser for .numbers."""
import json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__)); JS = os.path.join(HERE, '..', 'js', '01y-character-data.js')
MISSIONS = os.path.join(HERE, '..', 'js', '01x-mission-data.js')
A, B = '/*CHARACTERS-JSON*/', '/*END-CHARACTERS-JSON*/'
sys.path.insert(0, HERE)
import portraits as PT                       # the recipe's words the portraits know
WAVES = ['sine', 'triangle', 'square', 'sawtooth']
COLS = [  # (header, key, kind, list)   kind: text, upper, letter, colour, colour?, pick, extras, num(lo, hi)
    ('Name on screen', 'name', 'upper', None), ('Letter', 'letter', 'letter', None), ('Bubble colour', 'bubble', 'colour', None),
    ('Shirt in the city', 'shirt', 'colour', None), ('Trousers in the city', 'pants', 'colour', None),
    ('Who they are', 'role', 'text', None), ('How they look', 'look', 'text', None),
    ('Skin', 'skin', 'colour', None), ('Hair', 'hair', 'colour', None), ('Hair style', 'hairStyle', 'pick', PT.HAIR),
    ('Facial hair', 'facial', 'pick', PT.FACIAL), ('Facial hair colour', 'facialColour', 'colour?', None), ('Eyes', 'eyes', 'colour', None),
    ('Hat', 'hat', 'pick', PT.HATS), ('Hat colour', 'hatColour', 'colour?', None), ('Hat colour 2', 'hatColour2', 'colour?', None),
    ('Glasses', 'glasses', 'pick', PT.GLASSES), ('Glasses colour', 'glassesColour', 'colour?', None),
    ('Top', 'top', 'pick', PT.TOPS), ('Top colour', 'topColour', 'colour', None), ('Top colour 2', 'topColour2', 'colour?', None),
    ('Extras', 'extras', 'extras', PT.EXTRAS), ('Age', 'age', 'pick', PT.AGES), ('Build', 'build', 'pick', PT.BUILDS),
    ('Smiling eyes', 'smileEyes', 'pick', PT.SMILE_EYES), ('Calm mouth', 'calmMouth', 'pick', PT.CALM_MOUTH),
    ('Mumble wave', 'm.wave', 'pick', WAVES), ('Mumble low (Hz)', 'm.low', 'num', (40, 900)), ('Mumble high (Hz)', 'm.high', 'num', (40, 1200)),
    ('Syllables a second', 'm.rate', 'num', (1, 30)), ('Pause (s)', 'm.gap', 'num', (0, 1))]
HEAD = [c[0] for c in COLS] + ['ID']
TIPS = [('Yellow cells', 'What you can change. Hair style, Facial hair, Hat, Glasses, Top, Age, Build, Smiling eyes, Calm mouth and Mumble wave are lists to pick from.', 'tap a cell and type'),
        ('Colours', 'A colour is #rrggbb, like #ff9a3d. The facial hair, hat, glasses and the top\'s second colour can be left empty when not used.', ''),
        ('The portrait', 'Pixel art, 48 x 75, built from the owner\'s reference picture (a 3/4 view, transparent around the figure), drawn from the Skin ... Calm mouth columns by tools/portraits.py after the import; three faces: calm, angry, smiling. Extras: any of ' + ', '.join(PT.EXTRAS) + ', separated by commas.', ''),
        ('Smiling eyes', 'happy: the eyes arch when they smile; open: a smile that never reaches the eyes (Nico).', ''),
        ('Calm mouth', 'line, smirk (one corner up) or frown (both corners down) - the calm face.', ''),
        ('Faces', 'On the Faces sheet: the moods (as written on each line in the missions) that show the angry face, and the ones that show the smiling face, separated by commas. Any other mood: the calm face.', ''),
        ('Mumble', 'The sound of their voice while their words appear: the wave, the lowest and highest pitch, how many syllables a second, and the pause between words.', ''),
        ('New rows', 'A new row with a new ID (lowercase letters) is a new character. A row taken out is gone - unless a mission still has them.', ''),
        ('Sending it back', 'Upload the file in the chat as it is (.numbers or .xlsx). Every value is checked before anything is written.', '')]


def read_table():
    s = open(JS, encoding='utf-8').read(); i, j = s.index(A) + len(A), s.index(B); return s, i, j, json.loads(s[i:j])


def dump(t):
    f = t['faces']
    out = ['{', '"faces": {"angry": %s,' % json.dumps(f['angry']), '          "smile": %s},' % json.dumps(f['smile']), '"characters": [']
    for n, c in enumerate(t['characters']):
        out.append('  ' + json.dumps(c, ensure_ascii=False) + (',' if n < len(t['characters']) - 1 else ''))
    return '\n'.join(out + [']', '}'])


def get(c, key):
    if key.startswith('m.'): return c['mumble'][key[2:]]
    v = c.get(key, ''); return ', '.join(v) if isinstance(v, list) else v


def rows_of(t): return [[get(c, k) for _, k, _, _ in COLS] + [c['id']] for c in t['characters']]


def export_xlsx(dest, t):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    F = 'Arial'; f_in = PatternFill('solid', fgColor='FFF7CC'); f_head = PatternFill('solid', fgColor='2B2442'); f_key = PatternFill('solid', fgColor='EDEDED')
    thin = Side(style='thin', color='C9C9C9'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook(); ws = wb.active; ws.title = 'Characters'
    for j, h in enumerate(HEAD, start=1):
        c = ws.cell(row=1, column=j, value=h); c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = f_head; c.border = box; c.alignment = Alignment(horizontal='center', wrap_text=True)
    for j, (_, _, kind, lst) in enumerate(COLS, start=1):
        if kind == 'pick':
            dv = DataValidation(type='list', formula1='"%s"' % ','.join(lst), allow_blank=False, showErrorMessage=True, error='Pick from the list'); ws.add_data_validation(dv)
            dv.add('%s2:%s%d' % (col(j), col(j), len(t['characters']) + 1))
    for i, row in enumerate(rows_of(t), start=2):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v); c.border = box; last = j == len(HEAD)
            c.font = Font(name=F, size=10, bold=j == 1, color='595959' if last else '0000FF'); c.fill = f_key if last else f_in
            c.alignment = Alignment(vertical='center', wrap_text=len(str(v)) > 30)
        ws.row_dimensions[i].height = 48
    for j, h in enumerate(HEAD, start=1): ws.column_dimensions[col(j)].width = 40 if h in ('Who they are', 'How they look') else 14
    ws.freeze_panes = 'B2'
    fs = wb.create_sheet('Faces'); fs.append(['Face', 'Moods (separated by commas)']); fs.append(['angry', ', '.join(t['faces']['angry'])]); fs.append(['smile', ', '.join(t['faces']['smile'])])
    fs.column_dimensions['A'].width = 12; fs.column_dimensions['B'].width = 100
    g = wb.create_sheet('How to fill'); g.column_dimensions['A'].width = 18; g.column_dimensions['B'].width = 90; g.column_dimensions['C'].width = 18
    for i, row in enumerate([('Item', 'What it means', 'How to fill')] + TIPS, start=1):
        for j, v in enumerate(row, start=1): c = g.cell(row=i, column=j, value=v); c.font = Font(name=F, size=10, bold=i == 1 or j == 1); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = box
    wb.save(dest)


def col(j):
    s = ''
    while j: j, r = divmod(j - 1, 26); s = chr(65 + r) + s
    return s


def export_numbers(dest, t):
    from numbers_parser import Alignment, Document, RGB
    from settings_sheet import exact_numbers; exact_numbers()
    rows = rows_of(t)
    doc = Document(sheet_name='Characters', table_name='Characters', num_header_rows=1, num_header_cols=1, num_rows=len(rows) + 1, num_cols=len(HEAD))
    S = {}
    def style(name, **kw):
        if name not in S: S[name] = doc.add_style(name=name, font_name='Helvetica Neue', font_size=kw.pop('size', 11.0), alignment=Alignment(kw.pop('h', 'left'), 'middle'), **kw)
        return S[name]
    head = style('BR head', bg_color=RGB(43, 36, 66), font_color=RGB(255, 255, 255), bold=True, h='center'); val = style('BR value', bg_color=RGB(255, 247, 204), font_color=RGB(0, 0, 255))
    lab = style('BR label', bg_color=RGB(228, 224, 242), bold=True); key = style('BR key', bg_color=RGB(237, 237, 237), font_color=RGB(89, 89, 89), size=10.0)
    tb = doc.sheets[0].tables[0]
    for j, h in enumerate(HEAD): tb.write(0, j, h, style=head); tb.col_width(j, 240 if h in ('Who they are', 'How they look') else 100)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            tb.write(i, j, v, style=lab if j == 0 else key if j == len(HEAD) - 1 else val)
            if j < len(COLS) and COLS[j][2] == 'pick': tb.set_cell_formatting(i, j, 'popup', popup_values=COLS[j][3], allow_none=False)
        tb.row_height(i, 48)
    doc.add_sheet('Faces', 'Faces', num_rows=3, num_cols=2); fs = doc.sheets[-1].tables[0]
    for i, r in enumerate([['Face', 'Moods (separated by commas)'], ['angry', ', '.join(t['faces']['angry'])], ['smile', ', '.join(t['faces']['smile'])]]):
        for j, v in enumerate(r): fs.write(i, j, v, style=head if i == 0 else val if j == 1 else lab)
    fs.col_width(1, 600)
    doc.add_sheet('How to fill', 'How to fill', num_rows=len(TIPS) + 1, num_cols=3); g = doc.sheets[-1].tables[0]
    for i, r in enumerate([('Item', 'What it means', 'How to fill')] + TIPS):
        for j, v in enumerate(r): g.write(i, j, v)
    g.col_width(0, 130); g.col_width(1, 520); g.col_width(2, 120)
    doc.save(dest)


def read_sheets(src):
    if src.lower().endswith('.numbers'):
        from numbers_parser import Document
        return {sh.name: [tuple(r) for r in sh.tables[0].rows(values_only=True)] for sh in Document(src).sheets}
    from openpyxl import load_workbook
    return {ws.title: list(ws.iter_rows(values_only=True)) for ws in load_workbook(src, data_only=True).worksheets}


def table(sheets, title):
    data = next((r for n, r in sheets.items() if n.lower().startswith(title.lower())), None)
    if not data: return None, []
    hd = {str(h or '').strip().lower(): j for j, h in enumerate(data[0])}
    return hd, [r for r in data[1:] if any(c not in (None, '') for c in r)]


def text(v): return '' if v is None else str(v).strip()
def num(v):
    if isinstance(v, bool) or v is None or v == '': return None
    if isinstance(v, (int, float)): v = round(float(v), 6); return int(v) if v.is_integer() else v
    try: return num(float(str(v).strip().replace(',', '.')))
    except ValueError: return None
def colour(v):
    s = text(v).lstrip('#').lower(); return '#' + s if re.fullmatch(r'[0-9a-f]{6}', s) else None


def used_ids():                               # the characters the missions name
    s = open(MISSIONS, encoding='utf-8').read(); m = re.search(r'/\*MISSIONS-JSON\*/(.*?)/\*END-MISSIONS-JSON\*/', s, re.S)
    return {c for q in json.loads(m.group(1))['missions'] for c in q['chars']}


def do_import(src, dry):
    s, i, j, t = read_table(); sheets = read_sheets(src); errors, changes = [], []
    new = {'faces': {k: list(v) for k, v in t['faces'].items()}, 'characters': []}
    hd, rows = table(sheets, 'Characters')
    if hd is None: print('Not imported - the Characters sheet is missing'); sys.exit(1)
    old = {c['id']: c for c in t['characters']}
    for n, row in enumerate(rows, start=2):
        g = lambda h: row[hd[h.lower()]] if h.lower() in hd and hd[h.lower()] < len(row) else None
        cid = text(g('ID')).lower(); where = 'Characters row %d (%s)' % (n, cid or 'no ID')
        if not re.fullmatch(r'[a-z][a-z0-9_]*', cid): errors.append('%s: ID must be lowercase letters, digits or _' % where); continue
        if any(c['id'] == cid for c in new['characters']): errors.append('%s: ID used twice' % where); continue
        c = {'id': cid, 'mumble': {}}
        for h, key, kind, lst in COLS:
            raw = g(h); v = None
            if kind == 'text': v = text(raw)
            elif kind == 'upper':
                v = text(raw).upper()
                if not v: errors.append('%s: %s is needed' % (where, h))
            elif kind == 'letter':
                v = text(raw).upper()
                if not re.fullmatch(r'[A-Z0-9]', v): errors.append('%s: Letter must be one letter or digit, not "%s"' % (where, raw))
            elif kind in ('colour', 'colour?'):
                if kind == 'colour?' and not text(raw): v = ''
                else:
                    v = colour(raw)
                    if v is None: errors.append('%s: %s "%s" is not a colour like #ff9a3d' % (where, h, raw)); v = ''
            elif kind == 'pick':
                v = text(raw).lower()
                if v not in lst: errors.append('%s: %s "%s" is not one of %s' % (where, h, raw, ', '.join(lst)))
            elif kind == 'extras':
                v = [x.strip().lower() for x in text(raw).split(',') if x.strip()]
                bad = [x for x in v if x not in lst]
                if bad: errors.append('%s: Extras %s - only %s' % (where, ', '.join(bad), ', '.join(lst)))
            elif kind == 'num':
                v = num(raw); lo, hi = lst
                if v is None or v < lo or v > hi: errors.append('%s: %s must be a number from %g to %g, not "%s"' % (where, h, lo, hi, raw)); v = lo
            if key.startswith('m.'): c['mumble'][key[2:]] = v
            else: c[key] = v
        if c['mumble'].get('high', 0) < c['mumble'].get('low', 0): errors.append('%s: Mumble high must not be below Mumble low' % where)
        c = {k: c[k] for k in ['id', 'name', 'letter', 'bubble', 'shirt', 'pants', 'mumble', 'role', 'look'] + [x[1] for x in COLS if not x[1].startswith('m.') and x[1] not in ('name', 'letter', 'bubble', 'shirt', 'pants', 'role', 'look')]}
        new['characters'].append(c)
        if cid not in old: changes.append('new character %s (%s)' % (cid, c['name']))
        else:
            diff = [k for k in c if c[k] != old[cid].get(k)]
            if diff: changes.append('%s: %s' % (cid, ', '.join(diff) + ' changed'))
    gone = [k for k in old if k not in {c['id'] for c in new['characters']}]
    use = used_ids()
    for k in gone:
        if k in use: errors.append('Characters: %s is taken out but a mission still has them' % k)
        else: changes.append('character %s taken out' % k)
    hd, rows = table(sheets, 'Faces')
    if hd is not None:
        for row in rows:
            face, moods = text(row[0]).lower(), [w.strip().lower() for w in text(row[1] if len(row) > 1 else '').split(',') if w.strip()]
            if face not in ('angry', 'smile'): errors.append('Faces: "%s" is not angry or smile' % row[0]); continue
            if moods != new['faces'][face]: changes.append('Faces / %s: the moods changed' % face); new['faces'][face] = moods
        both = set(new['faces']['angry']) & set(new['faces']['smile'])
        if both: errors.append('Faces: %s cannot be both angry and smiling' % ', '.join(sorted(both)))
    if errors: print('Not imported - fix these first:'); [print('  ' + e) for e in errors]; sys.exit(1)
    print('%d change%s' % (len(changes), '' if len(changes) == 1 else 's') + (':' if changes else '')); [print('  ' + c) for c in changes]
    if dry or not changes: return
    open(JS, 'w', encoding='utf-8').write(s[:i] + dump(new) + s[j:]); print('written', os.path.relpath(JS))
    PT.draw_all()


if __name__ == '__main__':
    if len(sys.argv) < 3 or sys.argv[1] not in ('export', 'import'): print(__doc__); sys.exit(2)
    if sys.argv[1] == 'export': (export_numbers if sys.argv[2].lower().endswith('.numbers') else export_xlsx)(sys.argv[2], read_table()[3]); print('saved', sys.argv[2])
    else: do_import(sys.argv[2], '--dry-run' in sys.argv[3:])
