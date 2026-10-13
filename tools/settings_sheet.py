"""Settings tables <-> spreadsheet (Apple Numbers or Excel): the police (js/01b-police-data.js), day, night & weather (js/01d-sky-data.js),
the airport (js/01k-airport-data.js), the places - airport gates, military base, port, lunapark, space center, spray garages, the Colony's valet, Iron Key, Skyline One (js/01l-places-data.js) - and
the streets - traffic lights, rules of the road, walk signals, sirens, heat reducers (js/01m-streets-data.js) - and the vehicle jobs -
taxi, paramedic, firefighter, vigilante, their dispatch points, levels, pay and rewards (js/01o-jobs-data.js) - and the races - phone
booths, street races, the NASCAR race, the drag strip (js/01p-races-data.js; the fixed routes' junctions stay in the file) - and the stats -
what counts toward 100% and its reward (js/01t-stats-data.js) - and the casino strip - the ROYAL FLAMINGO's doors, slot machines, dice,
guards and cashier's cage, THE PULSE nightclub, the strip at night (js/01u-casino-data.js) - and the stunt jumps - slow motion, how a
ramp throws a car, a good landing, the bonus, each jump's speed and distance (js/01v-stunt-data.js).

  python3 tools/settings_sheet.py police export police.numbers            write the police table as a Numbers file (or .xlsx)
  python3 tools/settings_sheet.py sky export sky.numbers                  ... the day, night & weather table
  python3 tools/settings_sheet.py airport export airport.numbers          ... the airport and its planes
  python3 tools/settings_sheet.py places export places.numbers            ... the gates, the base, the port, the lunapark, the space center, the spray garages, the valet
  python3 tools/settings_sheet.py streets export streets.numbers          ... traffic lights, drivers, pedestrians, sirens, heat reducers
  python3 tools/settings_sheet.py jobs export jobs.numbers                ... the taxi, paramedic, firefighter and vigilante jobs
  python3 tools/settings_sheet.py races export races.numbers              ... phone booths, street races, the NASCAR race, the drag strip
  python3 tools/settings_sheet.py stats export stats.numbers              ... what counts toward 100%, the reward at 100%
  python3 tools/settings_sheet.py casino export casino.numbers            ... the casino, its games, guards and cage, the nightclub, the strip
  python3 tools/settings_sheet.py stunts export stunts.numbers            ... the stunt jumps: slow motion, lift, the landing, the bonus, each jump
  python3 tools/settings_sheet.py interiors export interiors.numbers      ... hangars, warehouses, the storage yard: hours, the guard, loot, rare vehicles
  python3 tools/settings_sheet.py police import police.numbers [--dry-run] read an edited file (.numbers or .xlsx) back, check it, rewrite the table

Both are made for Numbers on an iPhone or iPad: plain tables with one header row, no merged cells, no comments. Yellow cells are the
values to edit, YES / NO cells are pop-up menus. The Excel workbook also has grey Check cells (simple formulas); a Numbers file cannot
be written with formulas, so there the checks happen on import. Import matches rows by the ID column, so rows may be moved around,
but IDs must stay as they are. Needs openpyxl for .xlsx and numbers-parser for .numbers (pip install openpyxl numbers-parser)."""
import json, os, re, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
STARS = ['1 star', '2 stars', '3 stars', '4 stars', '5 stars']
TABLES = {'police': ('01b-police-data.js', 'POLICE', {'levels': 'Wanted levels', 'settings': 'Settings', 'crimes': 'Crimes'}),
          'sky': ('01d-sky-data.js', 'SKY', {'settings': 'Settings'}),
          'airport': ('01k-airport-data.js', 'AIRPORT', {'settings': 'Settings'}),
          'places': ('01l-places-data.js', 'PLACES', {'settings': 'Settings'}),
          'streets': ('01m-streets-data.js', 'STREETS', {'settings': 'Settings'}),
          'jobs': ('01o-jobs-data.js', 'JOBS', {'settings': 'Settings'}),
          'races': ('01p-races-data.js', 'RACES', {'settings': 'Settings'}),
          'stats': ('01t-stats-data.js', 'STATS', {'settings': 'Settings'}),
          'casino': ('01u-casino-data.js', 'CASINO', {'settings': 'Settings'}),
          'stunts': ('01v-stunt-data.js', 'STUNT', {'settings': 'Settings'}),
          'interiors': ('01w-interior-data.js', 'INTERIORS', {'settings': 'Settings'})}
JS = A = B = SHEETS = NAME = None
def use(name):                               # pick the table the commands work on
    global JS, A, B, SHEETS, NAME, TIPS
    f, mark, SHEETS = TABLES[name]; NAME = name
    JS = os.path.join(HERE, '..', 'js', f); A, B = '/*%s-JSON*/' % mark, '/*END-%s-JSON*/' % mark
    TIPS = {'police': TIPS_POLICE, 'sky': TIPS_SKY, 'airport': TIPS_AIRPORT, 'places': TIPS_PLACES, 'streets': TIPS_STREETS, 'jobs': TIPS_JOBS, 'races': TIPS_RACES, 'stats': TIPS_STATS, 'casino': TIPS_CASINO, 'stunts': TIPS_STUNTS, 'interiors': TIPS_INTERIORS}[name]


def read_table():
    s = open(JS, encoding='utf-8').read()
    i, j = s.index(A) + len(A), s.index(B)
    return s, i, j, json.loads(s[i:j])


def dump(t):                                # one row per line, the way the file is laid out
    out = ['{']
    for n, (k, rows) in enumerate(t.items()):
        out.append('"%s": [' % k)
        out += ['  ' + json.dumps(r, ensure_ascii=False) + (',' if m < len(rows) - 1 else '') for m, r in enumerate(rows)]
        out.append(']' + (',' if n < len(t) - 1 else ''))
    return '\n'.join(out + ['}'])


# ---------- export ----------
F = 'Arial'
f_in = PatternFill('solid', fgColor='FFF7CC'); f_calc = PatternFill('solid', fgColor='EDEDED'); f_head = PatternFill('solid', fgColor='2B2442')
f_lab = PatternFill('solid', fgColor='E4E0F2'); f_grp = PatternFill('solid', fgColor='F3F1FA')
thin = Side(style='thin', color='C9C9C9'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
yn = lambda v: 'YES' if v else 'NO'
allowed = lambda r: 'YES or NO' if r['unit'] == 'yes/no' else '%s to %s' % (fmt(r['min']), fmt(r['max']))
fmt = lambda v: ('%d' % v) if float(v).is_integer() else ('%g' % v)


def cell(ws, row, col, v, kind, wrap=False, bold=False, num='General'):
    c = ws.cell(row=row, column=col, value=v); c.border = box; c.number_format = num
    c.alignment = Alignment(wrap_text=wrap, vertical='center', horizontal='center' if kind in ('in', 'calc') and not wrap else 'left')
    if kind == 'head': c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = f_head
    elif kind == 'in': c.font = Font(name=F, size=10, color='0000FF'); c.fill = f_in
    elif kind == 'calc': c.font = Font(name=F, size=10, color='595959'); c.fill = f_calc
    elif kind == 'lab': c.font = Font(name=F, size=10, bold=True); c.fill = f_lab
    elif kind == 'id': c.font = Font(name=F, size=9, color='8A8A8A')
    else: c.font = Font(name=F, size=10, bold=bold)
    return c


def validate(ws, r, ref):
    if r['unit'] == 'yes/no':
        v = DataValidation(type='list', formula1='"YES,NO"', allow_blank=False, showErrorMessage=True, error='Pick YES or NO.')
    else:
        whole = r['unit'] in ('count', 'stars', 'level', 'floors')
        v = DataValidation(type='whole' if whole else 'decimal', operator='between', formula1=fmt(r['min']), formula2=fmt(r['max']), allow_blank=False,
                           showErrorMessage=True, error='Enter a %snumber from %s to %s.' % ('whole ' if whole else '', fmt(r['min']), fmt(r['max'])))
    ws.add_data_validation(v); v.add(ref)


def header(ws, names, widths):
    for j, (n, w) in enumerate(zip(names, widths), start=1): cell(ws, 1, j, n, 'head'); ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[1].height = 20


def export(path):
    t = read_table()[3]; wb = Workbook(); ws = wb.active
    if 'levels' in SHEETS: export_levels(ws, t)
    else: wb.remove(ws)
    flat_sheets(wb, t)
    tips_sheet(wb)
    wb.save(path); print('saved', path)


def export_levels(ws, t):
    # Wanted levels: one row per setting, one column per level
    ws.title = SHEETS['levels']
    header(ws, ['Setting'] + STARS + ['Unit', 'Allowed', 'What it does', 'ID'], [26, 9, 9, 9, 9, 9, 9, 13, 46, 11])
    rows = {}
    for i, r in enumerate(t['levels'], start=2):
        rows[r['id']] = i; cell(ws, i, 1, r['name'], 'lab', wrap=True)
        for k, v in enumerate(r['v']): cell(ws, i, 2 + k, yn(v) if r['unit'] == 'yes/no' else v, 'in')
        validate(ws, r, 'B%d:F%d' % (i, i))
        cell(ws, i, 7, r['unit'], ''); cell(ws, i, 8, allowed(r), ''); cell(ws, i, 9, r['note'], '', wrap=True); cell(ws, i, 10, r['id'], 'id')
        ws.row_dimensions[i].height = 44
    crime_row = {r['id']: i for i, r in enumerate(t['crimes'], start=2)}
    i = len(t['levels']) + 2; h = rows['heat']; cj = crime_row['carjack']
    cell(ws, i, 1, 'Calc: seen carjacks to get there', 'calc', wrap=True)
    for k in range(5):
        X = L(2 + k); cell(ws, i, 2 + k, '=IF(Crimes!$B$%d>0,ROUNDUP(%s%d/Crimes!$B$%d,0),"")' % (cj, X, h, cj), 'calc')
    cell(ws, i, 7, 'carjacks', 'calc'); cell(ws, i, 8, '', 'calc'); cell(ws, i, 9, 'How many carjackings a cop sees before you get this level.', 'calc', wrap=True); cell(ws, i, 10, '', 'calc')
    i += 1; cell(ws, i, 1, 'Calc: check', 'calc', wrap=True)
    for k in range(5):
        X, P = L(2 + k), L(1 + k)
        rise = '%s%d>0' % (X, h) if k == 0 else '%s%d>%s%d' % (X, h, P, h)
        cell(ws, i, 2 + k, '=IF(COUNTBLANK(%s2:%s%d)>0,"Fill every cell",IF(%s,"OK","Heat must be more than the level before"))' % (X, X, len(t['levels']) + 1, rise), 'calc', wrap=True)
    for j in (7, 8, 10): cell(ws, i, j, '', 'calc')
    cell(ws, i, 9, 'OK when the column is complete and needs more heat than the level before.', 'calc', wrap=True)
    ws.row_dimensions[i - 1].height = 30; ws.row_dimensions[i].height = 44
    ws.freeze_panes = 'B2'


def flat_sheets(wb, t):
    # Settings and Crimes: one row per value
    def flat(key, title, names, widths):
        ws = wb.create_sheet(title); header(ws, names, widths)
        for i, r in enumerate(t[key], start=2):
            cell(ws, i, 1, r['name'], 'lab', wrap=True)
            cell(ws, i, 2, yn(r['v']) if r['unit'] == 'yes/no' else r['v'], 'in'); validate(ws, r, 'B%d' % i)
            if r['unit'] == 'yes/no': chk = '=IF(OR(B{0}="YES",B{0}="NO"),"OK","Type YES or NO")'.format(i)
            else: chk = '=IF(AND(ISNUMBER(B{0}),B{0}>={1},B{0}<={2}),"OK","Must be {1} to {2}")'.format(i, fmt(r['min']), fmt(r['max']))
            cell(ws, i, 3, chk, 'calc'); cell(ws, i, 4, r['unit'], ''); cell(ws, i, 5, allowed(r), ''); cell(ws, i, 6, r['note'], '', wrap=True)
            j = 7
            if 'group' in r: c = cell(ws, i, j, r['group'], ''); c.fill = f_grp; j += 1
            cell(ws, i, j, r['id'], 'id'); ws.row_dimensions[i].height = 44 if len(r['note']) > 60 else 30
        ws.freeze_panes = 'B2'
    flat('settings', SHEETS['settings'], ['Setting', 'Value', 'Check', 'Unit', 'Allowed', 'What it does', 'Group', 'ID'], [27, 9, 10, 9, 13, 46, 15, 13])
    if 'crimes' in SHEETS: flat('crimes', SHEETS['crimes'], ['Crime', 'Heat', 'Check', 'Unit', 'Allowed', 'Note', 'ID'], [30, 8, 10, 8, 11, 46, 12])


def tips_sheet(wb):

    # How to fill
    g = wb.create_sheet('How to fill'); header(g, ['Item', 'What it means'], [24, 80])
    tips = [('Yellow cells', 'The values you change. Tap a cell and type. YES / NO cells offer a list.', f_in),
            ('Grey cells', 'Calculated - leave them. Check says OK, or what is wrong in that row or column.', f_calc)] + TIPS + [
            ('Sending it back', 'Share or export the file as Excel (.xlsx) and upload it in the chat. Every yellow value is read back into the game.', None)]
    for i, (a, b_, fill) in enumerate(tips, start=2):
        c = cell(g, i, 1, a, '', bold=True); cell(g, i, 2, b_, '', wrap=True)
        if fill: c.fill = fill
        g.row_dimensions[i].height = 44
    g.freeze_panes = 'A2'


TIPS_SKY = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
            ('Units', 'Times of day as hours 0 to 24 (17.5 = 17:30), durations in real minutes or seconds, % as 0 to 100.', None),
            ('Clock', 'The clock runs while you play: a whole day lasts the set number of real minutes. Dawn and dusk are the hour and a half around sunrise and sunset.', None),
            ('Weather', 'After each spell of weather the next one is picked by the five chances. A change blends in over the set seconds.', None),
            ('Rain', 'Wet roads give less grip (your car and traffic alike) and traffic slows down; at night the roads shine with the lights.', None),
            ('Police', 'At night and in rain or fog the police see less far than their normal sight range (police table).', None),
            ('Look', 'By day the ground turns from the night\'s purple to grey streets, green grass and sand, and buildings throw dark shadows; at night the city keeps only the moonlight set here, so lamps, windows, neon and headlights stand out.', None)]
TIPS_AIRPORT = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
            ('Units', 'Speeds in km/h, distances in metres (a car is 4.5 m long, the runway about 400 m), times in seconds, angles in degrees.', None),
            ('Traffic', 'One plane moves at a time: it lands and taxis to a free gate, or it is pushed back, taxis out and takes off. The next one starts the set seconds after the last one started (or as soon as it is done).', None),
            ('Gates', 'There are five gates. Below the lower count the next movement is a landing, above the higher one a take-off; in between they take turns.', None),
            ('Damage', 'Gunfire wears a plane down like a car. A blast close enough blows it up at once. A moving plane kills whoever it runs into and wrecks cars.', None),
            ('Fence', 'The fence round the airfield stops people and cars. Only vehicles at least this heavy, going at least this fast, break through.', None)]
TIPS_PLACES = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
            ('Units', 'Times in seconds or minutes, distances in metres, speeds in km/h, % as 0 to 100, stars as a wanted level 0 to 5.', None),
            ('Airport gates', 'Stop at a barrier and the guard lifts it; drive through without stopping and the boom breaks and you are wanted.', None),
            ('Military base', 'Inside the fence a warning counts down; still inside when it ends, the alarm gives the wanted level and the soldiers fire. They stay inside the base. What you take comes back after you have left.', None),
            ('Port', 'One ship moves at a time. Below the lower count the next movement is an arrival, above the higher one a departure.', None),
            ('Lunapark and space center', 'Visitors stroll inside the lunapark while you are there. The space center gate stays shut until the story opens it.', None),
            ('Spray garages', 'Drive in, the door shuts, pick a colour: repainted and repaired for the price (money in euros, also in the economy sheet). Not seen going in, the police lose you up to the stars set.', None),
            ('Colony', 'The valet at the hotel door greets you when you come close on foot or slowly in a car - once, until you have been away and come back.', None),
            ('City', 'Towers: in Palm Heights, Sunstrip and Seaview, this share of the bigger buildings behind the street fronts rise as towers, each between the two heights (a floor is 14 units of the game). Picked once with the city.', None),
            ('Factory', 'The cement works in Gravel Flats: the height of its tallest chimney (the other two are a little lower) and how much smoke rises from each.', None),
            ('Iron Key', 'The prison island: the free ferry shuttles between its pier and the pier on the Palm Heights shore, waiting at each end. Inside the walls a RESTRICTED AREA warning counts down; then the alarm gives stars and the tower guards fire. Busted, you are let out of its gate.', None),
            ('Construction site', 'Skyline One downtown: open (gate, workers, machines) between the two hours, locked with the night guard about the rest of the time. The ramps throw a car up by this share of its climbing speed; a car falls this far without damage; on foot a drop this high starts to hurt, and so much for each metre more. Taken machines and deck pickups come back after the time given, while you are away.', None)]
TIPS_JOBS = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
            ('Units', 'Times in seconds, distances in metres by road (a street block is about 100 m), speeds in km/h, money in euros, % as 0 to 100; a level and points are whole numbers.', None),
            ('Starting a job', 'Drive the right vehicle into the marker at its dispatch point and stop: taxi at a taxi rank, ambulance at a hospital, fire engine at a fire station, police car at a police station. The START card shows.', None),
            ('Levels', 'Level 1 asks for one fare, patient, fire or criminal car (Level 1 asks for), each level after for more. Each one adds time to the clock; a level done pays its bonus and the next begins.', None),
            ('The clock', 'Time is given as if you drove at the speed set for the job, over the road distance, plus the extra seconds. Too low a speed makes the job easy, too high makes it impossible.', None),
            ('Rewards', 'At the reward level, once per saved game: taxi - nitro in taxis; paramedic - more health; firefighter - fireproof; vigilante - more body armor.', None)]
TIPS_RACES = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
            ('Units', 'Times in seconds, distances in metres (a street block is about 100 m), money in euros, % as 0 to 100; counts are whole numbers.', None),
            ('Phone booths', 'Now and then one of the booths around you rings and shows on the maps. A booth by a fixed race offers that race (double ring) or a random one (fast trill).', None),
            ('Street races', 'Accept on foot at the booth: the fee is paid, then reach the start in time in an allowed car and stop - the countdown starts. Rivals do not wait for you.', None),
            ('Fixed races', 'Five routes, each with its fee and prizes (its own group). Their streets are set in the file (routes); ask Claude to change a route.', None),
            ('NASCAR race', 'At the small speedway on the Palm Heights beach: its race booth stands in the paddock by the shore street.', None),
            ('Drag strips', 'The Palm strip beside the speedway (201 m) and the Sandbar strip on the east beach (600 m), each with its booth: stop by it in a car - a race against a rival (fee and prize), or a test run alone.', None)]
TIPS_STATS = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
            ('Units', 'YES / NO switches a group in or out of the count; a level and counts are whole numbers; health and armor in points; time in minutes.', None),
            ('100%', 'Each thing counts once: a safehouse, a rampage, a kind of hidden weapon or item, a race, a job, a car delivered, a spray garage. The share done is the things done out of all of them, rounded down - 100% only when every one is done.', None),
            ('Modes', 'A group whose feature is switched off in the mode being played (modes table) is left out of the count, and so is a group switched off here.', None),
            ('The reward', 'At 100%: a one-of-a-kind gold car (vehicles table: gold) in your dearest garage with a free space, and health and body armor up to the points set. Lost, a new gold car waits there after the minutes set.', None)]
TIPS_CASINO = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
            ('Units', 'Money in euros; x bet: what three of a kind pays, times what the pull cost; counts and stars are whole numbers; times in seconds or minutes; hours 0 to 23.', None),
            ('Doors', 'With the dress code on, the bouncers let you in only when you wear at least one smart piece (the clothes table: Smart). Never while you are wanted; the club only while it is open.', None),
            ('Slot machines', 'Three reels; each holds as many of each picture as set. The rarer a picture, the more seldom three of it line up. Three flamingos win the jackpot, a fixed sum. With the numbers as they come, a machine pays back about 92% of what goes in.', None),
            ('Dice', 'Craps on the pass line: 7 or 11 on the first roll wins what you bet, 2, 3 or 12 loses it; any other total is the point, and you roll until it comes again (win) or a 7 does (lose). The house wins slightly more often (about 51 to 49).', None),
            ('Trouble', 'Shooting in the casino or hurting anyone in it brings the stars set and the guards open fire. Robbing the cashier\'s cage (a gun in your hand at its window) pays between the two sums, brings more stars, and the cage is empty for the time set.', None),
            ('THE PULSE', 'Open between the two hours (it may run past midnight: open at 22, close at 6). The music is a radio station, by its number in the radio list.', None)]
TIPS_STUNTS = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
            ('Units', 'Percent of the usual; m/s² how fast a car falls; degrees how far it may have spun round; money in euros; km/h and metres for each jump.', None),
            ('The jumps', 'Twenty: fourteen street ramps and five over the water (numbered as on the map) and Skyline One\'s launch ramp. A jump counts when you leave the lip at least as fast as its speed, fly at least its distance, land on your wheels not spun round further than set and without losing more of the car than set.', None),
            ('Lift and gravity', 'How hard a stunt ramp throws a car up and how fast it then falls; change them and each jump\'s speed needs setting again (a car flies about 29 m at 110 km/h, 44 m at 140 with the values as they are).', None),
            ('The bonus', 'Paid once per jump, the first time: the metres flown and the metres of height times their rates, between the least and the most.', None)]
TIPS_INTERIORS = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
            ('Units', 'Hours on the game clock (0 to 24); game h = hours on the game clock (24 = a game day); metres; km/h; seconds; money in euros; % as 0 to 100.', None),
            ('The places', 'Ten buildings: the airport\'s three hangars, three port warehouses, three more in Dockside and Gravel Flats, and the storage yard\'s eight lock-ups. They are always the same buildings.', None),
            ('Day and night', 'Between the two hours the doors are up and workers stand inside; the rest of the day the doors are down and a guard walks inside with a torch. The lock-ups have neither: force a lock any time.', None),
            ('Rare vehicles', 'The CLASSIC \'62, the RACE BIKE, the PROTOTYPE and the ARMOURED VAN (their numbers are in the vehicles table). Yours once you drive off; lost or left, a new one is back after the time set.', None)]
TIPS_STREETS = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
            ('Units', 'Times in seconds or minutes, distances in metres (a car is 4.5 m long, a street about 11 m wide), speeds in km/h, % as 0 to 100, stars as a wanted level 1 to 5.', None),
            ('Traffic lights', 'Each junction runs green, amber, all red for one group of approaches after another; opposite approaches share a group. Junctions are not in step with each other.', None),
            ('Drivers', 'Cars stop at the line on red (and on amber when they can), give way to oncoming cars when turning across them, and stop for people on the crossing of the street they turn into. Running a red light yourself is not a crime.', None),
            ('Pedestrians', 'People cross at the zebra crossings. The walk signal is green while the traffic across the crossing has red.', None),
            ('Lights and sirens', 'Ambulances and fire engines on a call and police cars after you go through red lights slowly. Cars ahead of them move to the right and slow down.', None),
            ('Heat reducers', 'Each takes one star off. One lies on the road while you are wanted enough; others wait in fixed back alleys. All of them are taken from inside a car too.', None),
            ('City sound', 'All made by the game, no recordings. Each % is the loudness of one layer (100 = as tuned, 0 = off); City sound volume scales them all. Sounds of cars and people come from where they are and fade with distance.', None)]
TIPS_POLICE = [('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
        ('Units', 'Distances in metres (a car is 4.5 m long, a street about 11 m wide), speeds in km/h or m/s, times in seconds, % as 0 to 100.', None),
        ('Heat and stars', 'Each crime the police see or hear adds heat (Crimes sheet). Enough heat gives a wanted level (Wanted levels: Heat needed). Heat below 1 star fades away; stars only go when you lose the police or pay a fine.', None),
        ('Seeing', 'A cop sees you within Sight range, inside the Field of view, with nothing in between. Whoever sees you tells the others by radio.', None),
        ('Searching', 'When nobody sees you they search around the spot where you were last seen (the red circle on the map, the stars blink). Out of sight for the level\'s time and they give up; inside the circle the time runs slower.', None),
        ('Stop order', 'Levels with Stop order and fine = YES: a cop near you orders you to stop. Stand still (on foot or in your car) for the fine; ignore it and you get the next level.', None),
        ('Arrest', 'Other levels: cars pull up beside you (or ram you where Cars ram you = YES), the crew gets out and a cop who holds you long enough arrests you. Drive away and they run back to their car.', None),
        ('Shooting', 'Only on levels with Shoot on sight = YES, or Shoot back if attacked = YES after you fired, rammed a police car or hurt a cop. They shout a warning first and never shoot with a passer-by in the way.', None),
        ('Busted / wasted', 'With Carry on after busted or wasted = YES you start again at the nearest police station or hospital and lose a share of your score (and your weapons, if set).', None),
        ('Roadblocks', 'From the levels with Roadblock police cars above 0: police cars across the street ahead of you while you drive, cops behind them with the level\'s guns, an APC in the middle where set. Ram through or turn off.', None),
        ('Armed people', 'Passers-by who carry a gun (weapon table: Carried by people) and turn on you shoot with the police\'s hit chance, which falls the farther you are, and the police\'s damage times their gun\'s damage against the pistol\'s. A revolver fires slower and harder, an SMG in short bursts.', None),
        ('Tank and APC', 'At the top wanted level the tank fires rockets at you; a police APC fires its turret gun in bursts when it sees you. The APC belt is also yours when you drive one (vehicle table: Weapon mg).', None)]


def exact_numbers():
    """numbers-parser stores a number as decimal128 through a float division, so 53 is saved as 52.99999999999999.
    Store the shortest decimal form of the value instead: 53 stays 53, 37.5 stays 37.5."""
    from decimal import Decimal
    import numbers_parser.cell as nc
    def pack(value):
        sign, digits, exp = (Decimal(value) if isinstance(value, int) else Decimal(repr(float(value)))).as_tuple()
        m, e, buf = int(''.join(map(str, digits)) or 0), exp + nc.DECIMAL128_BIAS, bytearray(16)
        buf[15] |= e >> 7; buf[14] |= (e & 0x7F) << 1
        for k in range(14): buf[k] = m & 0xFF; m >>= 8
        buf[14] |= m & 1
        if sign: buf[15] |= 0x80
        return buf
    nc._pack_decimal128 = pack


def export_numbers(path):                   # the same tables as a Numbers document: values, colours, pop-up menus (no formulas)
    from numbers_parser import Alignment, Document, RGB
    exact_numbers()
    t = read_table()[3]; doc = None
    def table(name, rows, cols, widths):
        nonlocal doc
        if doc is None: doc = Document(sheet_name=name, table_name=name, num_header_rows=1, num_header_cols=1, num_rows=rows, num_cols=cols)
        else: doc.add_sheet(name, name, num_rows=rows, num_cols=cols)
        tb = doc.sheets[-1].tables[0]; tb.num_header_rows = 1; tb.num_header_cols = 1
        for k, w in enumerate(widths): tb.col_width(k, w)
        return tb
    def style(name, **kw): return doc.add_style(name=name, font_name='Helvetica Neue', font_size=kw.pop('size', 11.0), alignment=Alignment(kw.pop('h', 'left'), 'middle'), **kw)
    S = {}
    def mk_styles():
      S.update({'head': style('BR head', bg_color=RGB(43, 36, 66), font_color=RGB(255, 255, 255), bold=True), 'lab': style('BR label', bg_color=RGB(228, 224, 242), bold=True),
         'in': style('BR value', bg_color=RGB(255, 247, 204), font_color=RGB(0, 0, 255), h='center'), 'txt': style('BR text'),
         'id': style('BR id', font_color=RGB(138, 138, 138), size=9.0), 'grp': style('BR group', bg_color=RGB(243, 241, 250))})
    def put(tb, r, c, v, k): tb.write(r, c, v, style=S[k])
    def value(tb, r, c, row, v):
        if row['unit'] == 'yes/no': put(tb, r, c, yn(v), 'in'); tb.set_cell_formatting(r, c, 'popup', popup_values=['YES', 'NO'], allow_none=False)
        else: put(tb, r, c, v, 'in')
    def head(tb, names): [put(tb, 0, k, n, 'head') for k, n in enumerate(names)]
    if 'levels' in SHEETS:
     tb = table(SHEETS['levels'], len(t['levels']) + 1, 10, [170, 58, 58, 58, 58, 58, 62, 80, 330, 80]); mk_styles()
     head(tb, ['Setting'] + STARS + ['Unit', 'Allowed', 'What it does', 'ID'])
     for i, r in enumerate(t['levels'], start=1):
        put(tb, i, 0, r['name'], 'lab')
        for k, v in enumerate(r['v']): value(tb, i, 1 + k, r, v)
        put(tb, i, 6, r['unit'], 'txt'); put(tb, i, 7, allowed(r), 'txt'); put(tb, i, 8, r['note'], 'txt'); put(tb, i, 9, r['id'], 'id'); tb.row_height(i, 48)
    for key, names, widths in (('settings', ['Setting', 'Value', 'Unit', 'Allowed', 'What it does', 'Group', 'ID'], [180, 62, 62, 80, 330, 110, 90]),
                               ('crimes', ['Crime', 'Heat', 'Unit', 'Allowed', 'Note', 'ID'], [210, 58, 58, 74, 330, 80])):
        if key not in SHEETS: continue
        tb = table(SHEETS[key], len(t[key]) + 1, len(names), widths)
        if not S: mk_styles()
        head(tb, names)
        for i, r in enumerate(t[key], start=1):
            put(tb, i, 0, r['name'], 'lab'); value(tb, i, 1, r, r['v']); put(tb, i, 2, r['unit'], 'txt'); put(tb, i, 3, allowed(r), 'txt'); put(tb, i, 4, r['note'], 'txt')
            if 'group' in r: put(tb, i, 5, r['group'], 'grp')
            put(tb, i, len(names) - 1, r['id'], 'id'); tb.row_height(i, 48 if len(r['note']) > 60 else 32)
    tips = [('Yellow cells', 'The values you change. Tap a cell and type; YES / NO cells are pop-up menus. The Allowed column shows what each one takes.'),
            ('Checks', 'This Numbers file has no formulas: every value is checked when the file comes back, and anything out of range is listed for you to fix.')] + \
           [(a, b_) for a, b_, _ in TIPS] + [('Sending it back', 'Upload the .numbers file in the chat as it is (or export it as Excel). Every yellow value is read back into the game.')]
    tb = table('How to fill', len(tips) + 1, 2, [150, 520]); head(tb, ['Item', 'What it means'])
    for i, (a, b_) in enumerate(tips, start=1): put(tb, i, 0, a, 'lab'); put(tb, i, 1, b_, 'txt'); tb.row_height(i, 48)
    doc.save(path); print('saved', path)


# ---------- import ----------
def num(v):
    if isinstance(v, bool): return None
    if isinstance(v, (int, float)): v = round(float(v), 6); return int(v) if v.is_integer() else v   # Numbers can give 50.00000000000001
    if isinstance(v, str):
        try: return num(float(v.strip().replace(',', '.')))
        except ValueError: return None
    return None


def yesno(v):
    if isinstance(v, bool): return v
    if isinstance(v, (int, float)) and v in (0, 1): return bool(v)
    if isinstance(v, str) and v.strip().upper() in ('YES', 'NO', 'TRUE', 'FALSE', 'Y', 'N'): return v.strip().upper() in ('YES', 'TRUE', 'Y')
    return None


def read_sheets(path):                      # {sheet name: rows of values}, from a Numbers file or an Excel workbook
    if path.lower().endswith('.numbers'):
        from numbers_parser import Document
        return {sh.name: [tuple(r) for tb in sh.tables for r in tb.rows(values_only=True)] for sh in Document(path).sheets}
    return {ws.title: list(ws.iter_rows(values_only=True)) for ws in load_workbook(path, data_only=True).worksheets}


def find_sheet(sheets, title):              # Numbers may rename "Settings" to "Settings - Table 1" on export to Excel
    for name, rows in sheets.items():
        if name == title or name.startswith(title + ' ') or name.startswith(title + '-'): return rows
    return None


def sheet_rows(rows_in):
    """header -> column map and {id: row values} for a sheet; the header is the first row that has an ID cell"""
    hdr, out = None, {}
    for row in rows_in:
        cells = [str(c).strip() if c is not None else '' for c in row]
        if hdr is None:
            if 'ID' in cells: hdr = {name.lower(): k for k, name in enumerate(cells) if name}
            continue
        k = hdr['id']
        if k < len(row) and row[k] not in (None, ''): out[str(row[k]).strip()] = row
    return hdr, out


def do_import(path, dry):
    s, i, j, t = read_table(); wb = read_sheets(path); errors, changes = [], []
    def take(r, raw, where):
        if r['unit'] == 'yes/no':
            v = yesno(raw)
            if v is None: errors.append('%s: "%s" is not YES or NO' % (where, raw))
            return v
        v = num(raw)
        if v is None: errors.append('%s: "%s" is not a number' % (where, raw)); return None
        if v < r['min'] or v > r['max']: errors.append('%s: %s is outside %s to %s' % (where, fmt(v), fmt(r['min']), fmt(r['max']))); return None
        if r['unit'] in ('count', 'stars', 'level', 'floors') and not float(v).is_integer(): errors.append('%s: %s must be a whole number' % (where, fmt(v))); return None
        return v
    show = lambda r, v: yn(v) if r['unit'] == 'yes/no' else fmt(v)
    for key, title in SHEETS.items():
        ws = find_sheet(wb, title)
        if ws is None: errors.append('sheet "%s" is missing' % title); continue
        hdr, rows = sheet_rows(ws)
        if hdr is None: errors.append('%s: no header row with an ID column' % title); continue
        for r in t[key]:
            if r['id'] not in rows: errors.append('%s: the row with ID %s is missing' % (title, r['id'])); continue
            row = rows[r['id']]
            if key == 'levels':
                cols = [hdr.get(n) for n in STARS]
                if None in cols: errors.append('%s: columns %s are needed' % (title, ', '.join(STARS))); break
                new = [take(r, row[c] if c < len(row) else None, '%s / %s / %s' % (title, r['name'], n)) for c, n in zip(cols, STARS)]
                if None in new: continue
                for n, a, b_ in zip(STARS, r['v'], new):
                    if a != b_: changes.append('%s / %s / %s: %s -> %s' % (title, r['name'], n, show(r, a), show(r, b_)))
                r['v'] = new
            else:
                c = hdr.get('value' if key == 'settings' else 'heat')
                if c is None: errors.append('%s: no %s column' % (title, 'Value' if key == 'settings' else 'Heat')); break
                v = take(r, row[c] if c < len(row) else None, '%s / %s' % (title, r['name']))
                if v is None: continue
                if v != r['v']: changes.append('%s / %s: %s -> %s' % (title, r['name'], show(r, r['v']), show(r, v)))
                r['v'] = v
    st = {r['id']: r['v'] for r in t['settings']}
    if NAME == 'police':
        heat = next(r['v'] for r in t['levels'] if r['id'] == 'heat')
        if any(b_ <= a for a, b_ in zip(heat, heat[1:])): errors.append('Wanted levels / Heat needed must rise from level to level: %s' % heat)
        if st['gapMax'] < st['gapMin']: errors.append('Settings: Time between shots to (%s) is less than from (%s)' % (fmt(st['gapMax']), fmt(st['gapMin'])))
        if st['rbAheadMax'] < st['rbAheadMin']: errors.append('Settings: Roadblocks ... to (%s) is less than from (%s)' % (fmt(st['rbAheadMax']), fmt(st['rbAheadMin'])))
        if st['pplGapMax'] < st['pplGapMin']: errors.append('Settings: Passer-by: time between shots to (%s) is less than from (%s)' % (fmt(st['pplGapMax']), fmt(st['pplGapMin'])))
        if st['pplBurstMax'] < st['pplBurstMin']: errors.append('Settings: Passer-by SMG burst to (%s) is less than from (%s)' % (fmt(st['pplBurstMax']), fmt(st['pplBurstMin'])))
    elif NAME == 'places':
        if st['shipsMax'] < st['shipsMin']: errors.append('Settings: Ships at the berths at most (%s) is less than at least (%s)' % (fmt(st['shipsMax']), fmt(st['shipsMin'])))
    elif NAME in ('streets', 'stats'): pass
    elif NAME == 'stunts':
        if st['bonusMax'] < st['bonusMin']: errors.append('Settings: The bonus - at most (%s) is less than at least (%s)' % (fmt(st['bonusMax']), fmt(st['bonusMin'])))
    elif NAME == 'interiors':
        if st['cashMax'] < st['cashMin']: errors.append('Settings: Cash in a stash - at most (%s) is less than at least (%s)' % (fmt(st['cashMax']), fmt(st['cashMin'])))
        if st['openTo'] <= st['openFrom']: errors.append('Settings: the doors shut (%s) must be later than they open (%s)' % (fmt(st['openTo']), fmt(st['openFrom'])))
    elif NAME == 'casino':
        if st['diceMax'] < st['diceMin']: errors.append('Settings: Dice - the biggest bet (%s) is less than the smallest (%s)' % (fmt(st['diceMax']), fmt(st['diceMin'])))
        if st['cageMax'] < st['cageMin']: errors.append("Settings: Cashier's cage - at most (%s) is less than at least (%s)" % (fmt(st['cageMax']), fmt(st['cageMin'])))
        if st['clubOpen'] == st['clubClose']: errors.append('Settings: THE PULSE - it cannot open and close at the same hour')
    elif NAME == 'jobs':
        name = {r['id']: r['name'] for r in t['settings']}
        for a, b_, what in (('taxiNear', 'taxiFar', 'Taxi: a call'), ('tripNear', 'tripFar', 'Taxi: a trip'), ('medNear', 'medFar', 'Paramedic: patients'),
                            ('fireNear', 'fireFar', 'Firefighter: a burning car'), ('vigNear', 'vigFar', 'Vigilante: a criminal car')):
            if st[b_] <= st[a]: errors.append('Settings: %s - "and at most" (%s) must be more than "at least" (%s)' % (what, fmt(st[b_]), fmt(st[a])))
    elif NAME == 'races':
        if st['randomFar'] <= st['randomNear']: errors.append('Settings: a random race "and at most" (%s) must be more than "at least" (%s)' % (fmt(st['randomFar']), fmt(st['randomNear'])))
        for g in ('random', 'seaview', 'palm', 'bridge', 'docks', 'coral', 'nascar'):
            p1, p2, p3 = st[g + 'Prize1'], st[g + 'Prize2'], st[g + 'Prize3']
            if not p1 >= p2 >= p3: errors.append('Settings: %s - the prizes must not grow from 1st to 3rd place (%s, %s, %s)' % (g, fmt(p1), fmt(p2), fmt(p3)))
    elif NAME == 'airport':
        if st['gatesMax'] < st['gatesMin']: errors.append('Settings: Planes at the gates at most (%s) is less than at least (%s)' % (fmt(st['gatesMax']), fmt(st['gatesMin'])))
        if st['takeRun'] > 380: errors.append('Settings: the take-off run must fit the runway (380 m at most)')
    else:
        if st['wxMax'] < st['wxMin']: errors.append('Settings: Weather lasts at most (%s) is less than at least (%s)' % (fmt(st['wxMax']), fmt(st['wxMin'])))
        if st['sunset'] <= st['sunrise'] + 3: errors.append('Settings: Sunset must come at least 3 hours after sunrise')
        if not any(st[k] > 0 for k in ('pClear', 'pCloudy', 'pRain', 'pStorm', 'pFog')): errors.append('Settings: at least one weather needs a chance above 0')
    if errors:
        print('Not imported - fix these first:'); [print('  ' + e) for e in errors]; sys.exit(1)
    print('%d change%s' % (len(changes), '' if len(changes) == 1 else 's') + (':' if changes else '')); [print('  ' + c) for c in changes]
    if dry or not changes: return
    open(JS, 'w', encoding='utf-8').write(s[:i] + dump(t) + s[j:]); print('written', os.path.relpath(JS))


def main(argv):
    if len(argv) < 3 or argv[0] not in TABLES or argv[1] not in ('export', 'import'): print(__doc__); sys.exit(2)
    use(argv[0])
    if argv[1] == 'export': (export_numbers if argv[2].lower().endswith('.numbers') else export)(argv[2])
    else: do_import(argv[2], '--dry-run' in argv[3:])


if __name__ == '__main__': main(sys.argv[1:])
