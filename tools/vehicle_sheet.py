"""Vehicle table <-> spreadsheet (Apple Numbers or Excel).

  python3 tools/vehicle_sheet.py export vehicles.numbers            write js/01c-vehicle-data.js as a Numbers file (or .xlsx)
  python3 tools/vehicle_sheet.py import vehicles.numbers [--dry-run] read an edited file (.numbers or .xlsx) back, check it, rewrite the table

One column per vehicle, one row per setting, the vehicle IDs in the top row. Rows are found by their label in the first column, so
rows may be moved; a row that is missing keeps the values already in the game. A vehicle whose column is gone is removed.
The Excel workbook has grey Calc rows (simple formulas); a Numbers file is written without formulas and checked on import.
Needs openpyxl for .xlsx and numbers-parser for .numbers."""
import json, os, re, sys

JS = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'js', '01c-vehicle-data.js')
A, B = '/*VEHICLE-JSON*/', '/*END-VEHICLE-JSON*/'
SHIFTERS = ['auto', 'sport', 'tall', 'throttle', 'bar']
BODIES = ['sedan', 'sports', 'truck', 'police', 'estate', 'pickup', 'limo', 'bike', 'bus', 'trash', 'ambulance', 'fire', 'apc', 'tank', 'new']
KERB_W, STALL_W = 2.33, 2.5                  # widest vehicle that fits the kerb parking lane / a parking lot stall (m)
# key, label, kind (text / choice / num / int / yesno / colour), choices or (min, max), what it means, unit / how to fill
ROWS = [
    ('id', 'ID', 'text', None, 'Internal key of the vehicle. Never shown in the game.', 'lowercase, no spaces, unique'),
    ('name', 'Name on screen', 'text', None, 'Name on the HUD ("PRESS E TO ENTER SEDAN").', 'short, CAPITALS look best'),
    ('body', 'Body (3D model)', 'choice', BODIES, 'Which 3D model it uses. new: a model that does not exist yet - describe it in Notes and Claude builds it.', 'pick from the list'),
    ('role', 'Role: civilian / police', 'choice', ['civilian', 'police'], 'police: light bar, patrols in traffic, joins chases (Police chase weight), and stealing it raises your wanted level more.', 'pick civilian or police'),
    ('len', 'Length (m)', 'num', (1, 20), 'Length of the body; also its collision size. Long vehicles turn wide.', 'metres'),
    ('wid', 'Width (m)', 'num', (0.5, 4), 'Width of the body; also its collision size. Up to 2.33 m parks at the kerb, up to 2.5 m in parking lots.', 'metres'),
    ('top', 'Top speed (km/h)', 'num', (5, 500), 'Highest speed on the flat. Traffic cruises at 40-55 km/h; police chase at 70-170 km/h by wanted level, never above this.', 'km/h'),
    ('accTo', 'Accelerates 0 to (km/h)', 'num', (5, 500), 'Speed used for the acceleration time below. At or above top speed it is read as the time to reach top speed.', 'km/h, usually 100'),
    ('accS', '...in seconds', 'num', (0.5, 120), 'Time from standstill to that speed, as the real vehicle does it. When you drive it, arcade handling makes it about 1.7x quicker.', 'seconds'),
    ('brake', 'Braking (m/s²)', 'num', (1, 40), 'Deceleration with the brakes full on (9.8 = 1 g). When you drive it, arcade handling brakes 1.5x harder.', 'm/s²'),
    ('steer', 'Steering (rad/s)', 'num', (0.2, 8), 'How fast it turns at full lock at town speed. At high speed grip limits the turn anyway.', 'radians per second, about 1-4'),
    ('grip', 'Tyre grip', 'num', (0.5, 15), 'How quickly a sideways slide stops. Low = drifty, high = on rails.', 'number, about 2-8'),
    ('hp', 'Health', 'num', (1, 100000), 'Damage it takes before it catches fire and explodes. A pistol hit does about 15, a rifle 88, a hard crash 10-40, a rocket up to 110.', 'points'),
    ('kg', 'Weight (kg)', 'num', (50, 100000), 'Mass in crashes and blasts: the heavier vehicle pushes the lighter one and takes less damage.', 'kilograms'),
    ('armored', 'Armored: yes / no', 'yesno', None, 'yes: a sniper round does not go through it.', 'yes or no'),
    ('traffic', 'Traffic weight', 'num', (0, 1000), 'How often it drives in traffic, relative to the others (0 = never). This sets the whole traffic mix, police included.', 'number, 0 or more'),
    ('parked', 'Parked weight', 'num', (0, 1000), 'How often it is parked at the kerb or in a parking lot (0 = never).', 'number, 0 or more'),
    ('chase', 'Police chase weight', 'num', (0, 1000), 'How often the police send it after you, among the police vehicles allowed at your wanted level. Only for role police.', 'number, 0 or more'),
    ('chaseFrom', 'Chase from wanted level', 'int', (1, 5), 'The lowest wanted level at which the police send it or let it join a chase.', '1 to 5 stars'),
    ('hidden', 'Hidden on map (count)', 'int', (0, 5), 'How many stand parked at secret spots, always in the same places, for you to find.', 'number, usually 0'),
    ('job', 'Job: none / trash / ambulance / fire', 'choice', ['none', 'trash', 'ambulance', 'fire'], 'trash: stops at the bins along its way. ambulance: comes for the dead. fire: comes to explosions and burning wrecks and puts them out.', 'pick from the list'),
    ('weapon', 'Weapon: none / rockets', 'choice', ['none', 'rockets'], 'rockets: when you drive it, FIRE launches rockets straight ahead (no ammo needed).', 'pick none or rockets'),
    ('shifter', 'Gear selector: auto / sport / tall / throttle / bar', 'choice', SHIFTERS, 'The gear selector on a touch screen (no letters; up is reverse, down is drive): auto - a T-handle lever, sport - a short round knob, tall - a long lever with a big knob, throttle - a military throttle handle, bar - a handlebar switch.', 'pick from the list'),
] + [('c%d' % k, 'Colour %d (#hex)' % k, 'colour', None, 'Paint colour; each vehicle picks one of its colours at random, and its neon underglow takes the same colour.' if k == 1 else 'Another paint colour (optional).',
      '#RRGGBB' + ('; at least one' if k == 1 else '')) for k in range(1, 7)] + [
    ('notes', 'Notes / look', 'text', None, 'How it looks or behaves, especially for a new body.', 'free text'),
]
ROW = {r[0]: r for r in ROWS}
# how a row label in a returned file is recognised (lowercase start of the label), most specific first
MATCH = [('id', 'id'), ('name', 'name on screen'), ('body', 'body'), ('role', 'role'), ('len', 'length'), ('wid', 'width'), ('top', 'top speed'),
         ('accTo', 'accelerates'), ('accS', '...in seconds'), ('accS', 'in seconds'), ('brake', 'braking'), ('steer', 'steering'), ('grip', 'tyre grip'),
         ('hp', 'health'), ('kg', 'weight'), ('armored', 'armored'), ('traffic', 'traffic weight'), ('parked', 'parked weight'), ('chase', 'police chase'),
         ('chaseFrom', 'chase from'), ('hidden', 'hidden'), ('job', 'job'), ('weapon', 'weapon'), ('shifter', 'gear selector')] + [('c%d' % k, 'colour %d' % k) for k in range(1, 7)] + [('notes', 'notes')]
DEFAULT = {'body': 'new', 'role': 'civilian', 'armored': False, 'traffic': 0, 'parked': 0, 'chase': 0, 'chaseFrom': 1, 'hidden': 0, 'job': '', 'weapon': '', 'shifter': 'auto', 'colors': [], 'notes': ''}
ORDER = ['id', 'name', 'body', 'role', 'len', 'wid', 'top', 'accTo', 'accS', 'brake', 'steer', 'grip', 'hp', 'kg', 'armored', 'traffic', 'parked', 'chase', 'chaseFrom',
         'hidden', 'job', 'weapon', 'shifter', 'colors', 'notes']


def read_table():
    s = open(JS, encoding='utf-8').read(); i, j = s.index(A) + len(A), s.index(B)
    return s, i, j, json.loads(s[i:j])


def dump(t):
    rows = t['vehicles']
    return '{\n"vehicles": [\n' + ',\n'.join('  ' + json.dumps(v, ensure_ascii=False) for v in rows) + '\n]\n}'


fmt = lambda v: ('%d' % v) if float(v).is_integer() else ('%g' % v)


def cell_value(v, key):                     # a vehicle field as it is shown in the sheet
    if key == 'armored': return 'yes' if v.get('armored') else 'no'
    if key in ('job', 'weapon'): return v.get(key) or 'none'
    if key[0] == 'c' and key[1:].isdigit(): k = int(key[1:]) - 1; return v['colors'][k] if k < len(v['colors']) else None
    return v.get(key)


# ---------- export ----------
def export_xlsx(path, t):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.datavalidation import DataValidation
    F = 'Arial'; f_in = PatternFill('solid', fgColor='FFF7CC'); f_calc = PatternFill('solid', fgColor='EDEDED'); f_head = PatternFill('solid', fgColor='2B2442')
    f_lab = PatternFill('solid', fgColor='E4E0F2'); f_labc = PatternFill('solid', fgColor='D9D9D9'); thin = Side(style='thin', color='C9C9C9'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    vs = t['vehicles']; NV = max(20, len(vs) + 6); C0, C1 = 2, 1 + NV; R = {r[0]: i for i, r in enumerate(ROWS, start=1)}
    wb = Workbook(); ws = wb.active; ws.title = 'Vehicles'; ws.column_dimensions['A'].width = 27
    for j in range(C0, C1 + 1): ws.column_dimensions[L(j)].width = 12
    for i, (key, lab, kind, *_r) in enumerate(ROWS, start=1):
        a = ws.cell(row=i, column=1, value=lab); a.font = Font(name=F, size=10, bold=True, color='FFFFFF' if key == 'id' else '000000')
        a.fill = f_head if key == 'id' else f_lab; a.alignment = Alignment(wrap_text=True, vertical='center'); a.border = box
        for n in range(NV):
            val = cell_value(vs[n], key) if n < len(vs) else None
            c = ws.cell(row=i, column=C0 + n, value=val); c.border = box; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=(key == 'notes'))
            if key == 'id': c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = f_head
            else: c.font = Font(name=F, size=10, color='0000FF'); c.fill = f_in
            if kind == 'colour' and val:
                hx = val.lstrip('#'); lum = 0.299 * int(hx[0:2], 16) + 0.587 * int(hx[2:4], 16) + 0.114 * int(hx[4:6], 16)
                c.fill = PatternFill('solid', fgColor=hx.upper()); c.font = Font(name=F, size=10, color='000000' if lum > 140 else 'FFFFFF')
        ws.row_dimensions[i].height = 58 if key == 'notes' else (30 if len(lab) > 26 else 18)
        span = '%s%d:%s%d' % (L(C0), i, L(C1), i)
        if kind in ('choice', 'yesno'):
            dv = DataValidation(type='list', formula1='"%s"' % ','.join(ROWS[i - 1][3] if kind == 'choice' else ['yes', 'no']), allow_blank=True, showErrorMessage=True, error='Pick from the list.')
            ws.add_data_validation(dv); dv.add(span)
        elif kind in ('num', 'int'):
            lo, hi = ROWS[i - 1][3]; dv = DataValidation(type='whole' if kind == 'int' else 'decimal', operator='between', formula1=fmt(lo), formula2=fmt(hi), allow_blank=True,
                                                       showErrorMessage=True, error='Enter a number from %s to %s.' % (fmt(lo), fmt(hi))); ws.add_data_validation(dv); dv.add(span)
    calc = [('Calc: Braking (g)', '0.00'), ('Calc: Stop from 100 km/h (m)', '0'), ('Calc: Traffic share', '0%'), ('Calc: Parked share', '0%'), ('Calc: Parks', '@'), ('Calc: Check', '@')]
    for k, (lab, nf) in enumerate(calc):
        i = len(ROWS) + 1 + k; a = ws.cell(row=i, column=1, value=lab); a.font = Font(name=F, size=10, bold=True, color='595959'); a.fill = f_labc; a.border = box
        for n in range(NV):
            X = L(C0 + n); r = lambda key: '%s%d' % (X, R[key]); sp = lambda key: '$%s%d:$%s%d' % (L(C0), R[key], L(C1), R[key])
            f = {'Calc: Braking (g)': '=IF({b}="","",{b}/9.81)'.format(b=r('brake')),
                 'Calc: Stop from 100 km/h (m)': '=IF(OR({b}="",{b}<=0),"",(100/3.6)^2/(2*{b}))'.format(b=r('brake')),
                 'Calc: Traffic share': '=IF(OR({i}="",{w}="",SUM({s})=0),"",{w}/SUM({s}))'.format(i=r('id'), w=r('traffic'), s=sp('traffic')),
                 'Calc: Parked share': '=IF(OR({i}="",{w}="",SUM({s})=0),"",{w}/SUM({s}))'.format(i=r('id'), w=r('parked'), s=sp('parked')),
                 'Calc: Parks': '=IF(OR({i}="",{p}="",{p}=0),"",IF({w}<={k},"kerb and lots",IF({w}<={s},"lots only","too wide")))'.format(i=r('id'), p=r('parked'), w=r('wid'), k=KERB_W, s=STALL_W),
                 'Calc: Check': ('=IF({i}="","",IF(COUNTIF($B${ri}:${C1}${ri},{i})>1,"Duplicate ID",IF(COUNTBLANK({n}:{kg})>0,"Fill every cell from Name to Weight",'
                                 'IF(AND({ch}<>"",{ch}>0,{role}<>"police"),"Chase weight needs role police",IF(AND({body}="new",{notes}=""),"Describe the new body in Notes",'
                                 'IF({c1}="","Add at least Colour 1",IF(OR(LEN({c1})<>7,LEFT({c1},1)<>"#"),"Colour 1 must look like #ff2bd6","OK")))))))').format(
                     i=r('id'), ri=R['id'], C1=L(C1), n=r('name'), kg=r('kg'), ch=r('chase'), role=r('role'), body=r('body'), notes=r('notes'), c1=r('c1'))}[lab]
            c = ws.cell(row=i, column=C0 + n, value=f); c.number_format = nf; c.font = Font(name=F, size=10); c.fill = f_calc; c.border = box
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=lab == 'Calc: Check')
        ws.row_dimensions[i].height = 44 if lab == 'Calc: Check' else 18
    ws.freeze_panes = 'B2'
    g = wb.create_sheet('How to fill'); g.column_dimensions['A'].width = 24; g.column_dimensions['B'].width = 60; g.column_dimensions['C'].width = 28
    for j, h in enumerate(['Item', 'What it means', 'Unit / how to fill'], start=1):
        c = g.cell(row=1, column=j, value=h); c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = f_head; c.border = box
    rows = TIPS + [('Grey rows', 'Calculated - leave them. Check says OK when a column is complete, otherwise what to fix.', '')] + \
        [(('Colour 1 ... 6' if lab.startswith('Colour 1') else lab), mean, unit) for key, lab, kind, ch, mean, unit in ROWS if not (kind == 'colour' and key != 'c1')]
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            c = g.cell(row=i, column=j, value=v); c.font = Font(name=F, size=10, bold=(j == 1)); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = box
            if j == 1: c.fill = f_lab
    g.freeze_panes = 'A2'; wb.save(path)


TIPS = [('Yellow cells', 'What you type in. Each column is one vehicle. Body, Role, Armored, Job and Weapon are lists to pick from.', 'tap a cell and type'),
        ('New vehicle', 'Use the next empty column: type its ID in the dark top row, then fill the yellow cells below it.', ''),
        ('Remove a vehicle', 'Delete its column (or clear its ID).', ''),
        ('Your own vehicle', 'When you drive it, it handles arcade style: about 1.7x quicker acceleration, 1.5x harder brakes, more grip unless the handbrake is pulled. Fill in real-world values.', ''),
        ('Police', 'How the police behave (when they come, how many, shooting) is set in the police table. Here: which vehicles they use and from which wanted level.', ''),
        ('Sending it back', 'Upload the file in the chat as it is (.numbers or .xlsx). Every column with an ID becomes a vehicle in the game.', '')]


def export_numbers(path, t):
    from numbers_parser import Alignment, Document, RGB
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); from settings_sheet import exact_numbers; exact_numbers()
    vs = t['vehicles']; NV = max(20, len(vs) + 6)
    doc = Document(sheet_name='Vehicles', table_name='Vehicles', num_header_rows=1, num_header_cols=1, num_rows=len(ROWS), num_cols=NV + 1); tb = doc.sheets[0].tables[0]
    S = {}
    def style(name, **kw):
        if name not in S: S[name] = doc.add_style(name=name, font_name='Helvetica Neue', font_size=kw.pop('size', 11.0), alignment=Alignment(kw.pop('h', 'left'), 'middle'), **kw)
        return S[name]
    head = style('BR head', bg_color=RGB(43, 36, 66), font_color=RGB(255, 255, 255), bold=True, h='center'); lab_s = style('BR label', bg_color=RGB(228, 224, 242), bold=True)
    inp = style('BR value', bg_color=RGB(255, 247, 204), font_color=RGB(0, 0, 255), h='center'); txt = style('BR text')
    tb.col_width(0, 200)
    for j in range(1, NV + 1): tb.col_width(j, 92)
    for i, (key, lab, kind, ch, *_r) in enumerate(ROWS):
        tb.write(i, 0, lab, style=head if key == 'id' else lab_s)
        for n in range(NV):
            val = cell_value(vs[n], key) if n < len(vs) else None
            if key == 'id':
                if val is not None: tb.write(i, 1 + n, val, style=head)
                else: tb.set_cell_style(i, 1 + n, head)
                continue
            if kind == 'colour' and val:
                hx = val.lstrip('#'); r, g, b = int(hx[0:2], 16), int(hx[2:4], 16), int(hx[4:6], 16); light = 0.299 * r + 0.587 * g + 0.114 * b > 140
                tb.write(i, 1 + n, val, style=style('BR colour ' + hx, bg_color=RGB(r, g, b), font_color=RGB(0, 0, 0) if light else RGB(255, 255, 255), h='center'))
            elif val is None or val == '':
                if kind in ('choice', 'yesno'): tb.write(i, 1 + n, '', style=inp)          # a pop-up menu needs a (blank) text cell under it
                else: tb.set_cell_style(i, 1 + n, inp)
            else: tb.write(i, 1 + n, val, style=inp)
            if kind in ('choice', 'yesno'): tb.set_cell_formatting(i, 1 + n, 'popup', popup_values=ch if kind == 'choice' else ['yes', 'no'], allow_none=True)
        tb.row_height(i, 64 if key == 'notes' else (40 if len(lab) > 26 else 26))
    rows = TIPS + [('Checks', 'This Numbers file has no formulas: every column is checked when the file comes back, and anything to fix is listed for you.', '')] + \
        [(('Colour 1 ... 6' if lab.startswith('Colour 1') else lab), mean, unit) for key, lab, kind, ch, mean, unit in ROWS if not (kind == 'colour' and key != 'c1')]
    doc.add_sheet('How to fill', 'How to fill', num_rows=len(rows) + 1, num_cols=3)
    g = doc.sheets[-1].tables[0]; g.num_header_rows = 1; g.num_header_cols = 1; g.col_width(0, 190); g.col_width(1, 430); g.col_width(2, 170)
    for j, h in enumerate(['Item', 'What it means', 'Unit / how to fill']): g.write(0, j, h, style=head)
    for i, row in enumerate(rows, start=1):
        g.write(i, 0, row[0], style=lab_s); g.write(i, 1, row[1], style=txt); g.write(i, 2, row[2] or '', style=txt); g.row_height(i, 48)
    doc.save(path)


# ---------- import ----------
def read_columns(path):
    """the Vehicles table as {field: value} per vehicle column, plus the fields the file has rows for"""
    if path.lower().endswith('.numbers'):
        from numbers_parser import Document
        sheets = {sh.name: [tuple(r) for r in sh.tables[0].rows(values_only=True)] for sh in Document(path).sheets}
    else:
        from openpyxl import load_workbook
        sheets = {ws.title: list(ws.iter_rows(values_only=True)) for ws in load_workbook(path, data_only=True).worksheets}
    rows = next((r for n, r in sheets.items() if n.lower().startswith('vehicles')), None) or next(iter(sheets.values()))
    field_row = {}
    for k, row in enumerate(rows):
        lab = str(row[0] or '').strip().lower()
        for key, start in MATCH:
            if lab == start if key == 'id' else lab.startswith(start):
                field_row.setdefault(key, k); break
    if 'id' not in field_row: raise SystemExit('No row labelled ID in the first column.')
    cols = []
    for j in range(1, len(rows[field_row['id']])):
        vid = rows[field_row['id']][j]
        if vid in (None, '') or str(vid).strip() == '': continue
        cols.append({key: (rows[k][j] if j < len(rows[k]) else None) for key, k in field_row.items()})
    return cols, set(field_row)


def norm_id(v): return re.sub(r'[^a-z0-9]+', '_', str(v).strip().lower()).strip('_')


def num(v):
    if isinstance(v, bool) or v is None: return None
    if isinstance(v, str):
        try: v = float(v.strip().replace(',', '.'))
        except ValueError: return None
    v = round(float(v), 6); return int(v) if v.is_integer() else v


def build(cols, present, old):
    """vehicles from the sheet columns; rows the file does not have keep their values. -> vehicles, errors, warnings"""
    olds = {v['id']: v for v in old}; out, errors, warns, seen = [], [], [], set()
    for col in cols:
        vid = norm_id(col['id']); where = 'Column %s' % vid
        if vid in seen: errors.append('%s: the ID is used twice' % where); continue
        seen.add(vid); v = dict(DEFAULT, **olds.get(vid, {})); v['id'] = vid; v['colors'] = list(v.get('colors', []))
        for key in present:
            if key == 'id': continue
            raw = col.get(key); kind = ROW[key][2]; lab = ROW[key][1]
            if isinstance(raw, str): raw = raw.strip()
            if kind == 'colour':
                k = int(key[1:]) - 1
                while len(v['colors']) <= k: v['colors'].append(None)
                if raw in (None, ''): v['colors'][k] = None
                elif not re.fullmatch(r'#?[0-9a-fA-F]{6}', str(raw)): errors.append('%s / %s: "%s" is not a colour like #ff2bd6' % (where, lab, raw))
                else: v['colors'][k] = '#' + str(raw).lstrip('#').lower()
            elif kind == 'text': v[key] = '' if raw is None else str(raw)
            elif kind == 'yesno':
                if raw in (None, ''): v[key] = False
                elif str(raw).lower() in ('yes', 'no', 'true', 'false'): v[key] = str(raw).lower() in ('yes', 'true')
                elif isinstance(raw, bool): v[key] = raw
                else: errors.append('%s / %s: "%s" is not yes or no' % (where, lab, raw))
            elif kind == 'choice':
                s = str(raw or '').lower()
                if key in ('job', 'weapon') and s in ('', 'none'): v[key] = ''
                elif s in ROW[key][3]: v[key] = s
                else: errors.append('%s / %s: "%s" is not one of %s' % (where, lab, raw, ', '.join(ROW[key][3])))
            else:
                n = num(raw); lo, hi = ROW[key][3]
                if n is None:
                    if key in ('traffic', 'parked', 'chase', 'hidden'): v[key] = 0
                    elif key == 'chaseFrom': v[key] = 1
                    else: errors.append('%s / %s: missing or not a number' % (where, lab))
                elif n < lo or n > hi: errors.append('%s / %s: %s is outside %s to %s' % (where, lab, fmt(n), fmt(lo), fmt(hi)))
                elif kind == 'int' and not float(n).is_integer(): errors.append('%s / %s: %s must be a whole number' % (where, lab, fmt(n)))
                else: v[key] = n
        v['colors'] = [c for c in v['colors'] if c]
        for key in ('name', 'len', 'wid', 'top', 'accTo', 'accS', 'brake', 'steer', 'grip', 'hp', 'kg'):
            if v.get(key) in (None, ''): errors.append('%s: %s is missing' % (where, ROW[key][1]))
        if not v['colors']: errors.append('%s: add at least Colour 1' % where)
        if v['chase'] and v['role'] != 'police': errors.append('%s: Police chase weight needs role police' % where)
        if v['body'] == 'new': warns.append('%s: body "new" - shown as a sedan until its model is built' % where)
        if isinstance(v.get('accTo'), (int, float)) and isinstance(v.get('top'), (int, float)) and v['accTo'] >= v['top']:
            warns.append('%s: accelerates to %s km/h but tops out at %s - read as the time to reach top speed' % (where, fmt(v['accTo']), fmt(v['top'])))
        if v['parked'] and isinstance(v.get('wid'), (int, float)) and v['wid'] > KERB_W:
            warns.append('%s: %s m wide - parks %s' % (where, fmt(v['wid']), 'only in parking lots' if v['wid'] <= STALL_W else 'nowhere (too wide even for a parking lot)'))
        out.append({k: v.get(k, DEFAULT.get(k)) for k in ORDER})
    if not any(v['role'] == 'civilian' and v['traffic'] > 0 for v in out): errors.append('At least one civilian vehicle needs a Traffic weight')
    if not any(v['role'] == 'police' and v['chase'] > 0 and v['chaseFrom'] <= 1 for v in out): errors.append('At least one police vehicle needs a Police chase weight from wanted level 1')
    if not any(v['parked'] > 0 and v['wid'] <= KERB_W for v in out if isinstance(v.get('wid'), (int, float))): errors.append('At least one vehicle that fits the kerb needs a Parked weight')
    gone = [k for k in olds if k not in seen]
    if gone: warns.append('Removed (their columns are gone): ' + ', '.join(gone))
    return out, errors, warns


def changes(old, new):
    o = {v['id']: v for v in old}; out = []
    for v in new:
        if v['id'] not in o: out.append('new vehicle %s (%s)' % (v['id'], v['name'])); continue
        for k in ORDER:
            if v.get(k) != o[v['id']].get(k): out.append('%s / %s: %s -> %s' % (v['id'], k, o[v['id']].get(k), v.get(k)))
    return out


def write(s, i, j, vehicles):
    open(JS, 'w', encoding='utf-8').write(s[:i] + dump({'vehicles': vehicles}) + s[j:]); print('written', os.path.relpath(JS))


def do_import(path, dry):
    s, i, j, t = read_table(); cols, present = read_columns(path)
    vs, errors, warns = build(cols, present, t['vehicles'])
    for w in warns: print('note: ' + w)
    if errors: print('Not imported - fix these first:'); [print('  ' + e) for e in errors]; sys.exit(1)
    ch = changes(t['vehicles'], vs); print('%d change%s' % (len(ch), '' if len(ch) == 1 else 's') + (':' if ch else '')); [print('  ' + c) for c in ch]
    if not dry and (ch or [v['id'] for v in vs] != [v['id'] for v in t['vehicles']]): write(s, i, j, vs)


if __name__ == '__main__':
    if len(sys.argv) < 3 or sys.argv[1] not in ('export', 'import'): print(__doc__); sys.exit(2)
    if sys.argv[1] == 'export':
        t = read_table()[3]; (export_numbers if sys.argv[2].lower().endswith('.numbers') else export_xlsx)(sys.argv[2], t); print('saved', sys.argv[2])
    else: do_import(sys.argv[2], '--dry-run' in sys.argv[3:])
