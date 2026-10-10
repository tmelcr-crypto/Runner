"""Clothes table <-> spreadsheet (Apple Numbers or Excel): the ten clothes shops and every piece they sell (js/01n-clothes-data.js).

  python3 tools/clothes_sheet.py export clothes.numbers            write the table as a Numbers file (or .xlsx)
  python3 tools/clothes_sheet.py import clothes.numbers [--dry-run] read an edited file back, check it, rewrite the table

Sheets: Settings (new clothes lose the police), Shops (name, colour, what it sells), Items (one row per piece: name, shop, slot, style,
two colours, price). Rows are found by the ID column; a new row with a new ID is a new shop or a new piece, a row taken out is gone.
Items with an empty Shop are what you wear in a new game (one top, bottoms and shoes) and are never sold. Prices can also be edited in
tools/economy_sheet.py. Needs openpyxl for .xlsx and numbers-parser for .numbers."""
import json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__)); JS = os.path.join(HERE, '..', 'js', '01n-clothes-data.js')
A, B = '/*CLOTHES-JSON*/', '/*END-CLOTHES-JSON*/'
STYLES = {'hat': ['cap', 'beanie', 'bucket', 'fedora', 'cowboy', 'hardhat', 'beret', 'bandana', 'headband'],
          'glasses': ['shades', 'aviators', 'round', 'visor'],
          'top': ['tee', 'tank', 'hawaii', 'shirt', 'check', 'hoodie', 'jacket', 'suit', 'leather', 'hivis', 'jersey', 'camo'],
          'bottoms': ['jeans', 'trousers', 'shorts', 'track', 'cargo'],
          'shoes': ['sneakers', 'boots', 'sandals', 'loafers', 'cowboy']}
SLOTS = list(STYLES)
ALL_STYLES = sorted({s for v in STYLES.values() for s in v})
TIPS = [('Yellow cells', 'What you can change. Slot, Style and Shop are lists to pick from (Style must fit the slot - see Styles below).', 'tap a cell and type'),
        ('Items', 'One row per piece: its name on screen, the shop that sells it (its ID from the Shops sheet), the slot it is worn in, the style it is drawn with, its colours and its price in euros.', ''),
        ('Colours', 'Colour 1 is the main colour; Colour 2 a second one - the print on a T-shirt, a jacket\'s shirt or tie, stripes, a hat\'s band or brim, a glasses frame, a shoe\'s sole. Empty: none.', '#RRGGBB'),
        ('Styles', '; '.join('%s: %s' % (k, ', '.join(v)) for k, v in STYLES.items()), ''),
        ('Starting clothes', 'Rows with an empty Shop are what you wear in a new game (one top, bottoms and shoes); they are never sold and have no price.', ''),
        ('New rows', 'A new row with a new ID is a new piece (or on Shops, a new shop - the game finds it a building). Take a row out and it is gone.', ''),
        ('Settings', 'Changing at least one piece at a clothes shop while no cop saw you go in makes the police lose you, up to the wanted level set; above it, the level drops to the second setting.', ''),
        ('Sending it back', 'Upload the file in the chat as it is (.numbers or .xlsx). Every value is checked before anything is written.', '')]
SH_HEAD = ['Shop', 'Colour (#hex)', 'What it sells', 'ID']
IT_HEAD = ['Name on screen', 'Shop', 'Slot', 'Style', 'Colour 1', 'Colour 2', 'Price (€)', 'Notes', 'ID']
ST_HEAD = ['Setting', 'Value', 'Unit', 'Allowed', 'What it does', 'ID']
fmt = lambda v: ('%d' % v) if float(v).is_integer() else ('%g' % v)


def read_table():
    s = open(JS, encoding='utf-8').read(); i, j = s.index(A) + len(A), s.index(B); return s, i, j, json.loads(s[i:j])


def dump(t):
    out = ['{']
    for n, (k, rows) in enumerate(t.items()):
        out.append('"%s": [' % k)
        out += ['  ' + json.dumps(r, ensure_ascii=False) + (',' if m < len(rows) - 1 else '') for m, r in enumerate(rows)]
        out.append(']' + (',' if n < len(t) - 1 else ''))
    return '\n'.join(out + ['}'])


def sheet_data(t):
    st = [[r['name'], r['v'], r['unit'], '%s to %s' % (fmt(r['min']), fmt(r['max'])), r.get('note', ''), r['id']] for r in t['settings']]
    sh = [[s['name'], s['color'], s.get('theme', ''), s['id']] for s in t['shops']]
    it = [[i['name'], i['shop'], i['slot'], i['style'], i['c1'], i.get('c2', ''), i.get('price'), i.get('notes', ''), i['id']] for i in t['items']]
    return st, sh, it


def export_xlsx(dest, t):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    F = 'Arial'; f_in = PatternFill('solid', fgColor='FFF7CC'); f_head = PatternFill('solid', fgColor='2B2442'); f_lab = PatternFill('solid', fgColor='E4E0F2'); f_key = PatternFill('solid', fgColor='EDEDED')
    thin = Side(style='thin', color='C9C9C9'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    st, sh, it = sheet_data(t); wb = Workbook(); first = True
    def sheet(title, head, rows, widths, inputs, lists):
        nonlocal first
        ws = wb.active if first else wb.create_sheet(title); ws.title = title; first = False
        for j, h in enumerate(head, start=1):
            c = ws.cell(row=1, column=j, value=h); c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = f_head; c.border = box; c.alignment = Alignment(horizontal='center')
        dvs = {}
        for j, vals in lists.items():
            dv = DataValidation(type='list', formula1='"%s"' % ','.join(vals), allow_blank=True, showErrorMessage=True, error='Pick from the list'); ws.add_data_validation(dv); dvs[j] = dv
        for i, row in enumerate(rows, start=2):
            for j, v in enumerate(row, start=1):
                c = ws.cell(row=i, column=j, value=v); c.border = box; inp = j in inputs
                c.font = Font(name=F, size=10, bold=j == 1, color='0000FF' if inp else '595959' if j == len(head) else '000000')
                c.fill = f_in if inp else f_key if j == len(head) else f_lab if j == 1 else PatternFill()
                c.alignment = Alignment(vertical='center', wrap_text=len(str(v or '')) > 40)
                if j in dvs: dvs[j].add(c.coordinate)
        for k, w in enumerate(widths): ws.column_dimensions[chr(65 + k)].width = w
        ws.freeze_panes = 'B2'
    sheet('Items', IT_HEAD, it, [26, 10, 10, 11, 11, 11, 10, 30, 18], {1, 2, 3, 4, 5, 6, 7, 8}, {2: [s['id'] for s in t['shops']], 3: SLOTS, 4: ALL_STYLES})
    sheet('Shops', SH_HEAD, sh, [22, 14, 60, 12], {1, 2, 3}, {})
    sheet('Settings', ST_HEAD, st, [40, 9, 8, 10, 60, 12], {2}, {})
    g = wb.create_sheet('How to fill'); g.column_dimensions['A'].width = 20; g.column_dimensions['B'].width = 90; g.column_dimensions['C'].width = 18
    for i, row in enumerate([('Item', 'What it means', 'How to fill')] + TIPS, start=1):
        for j, v in enumerate(row, start=1): c = g.cell(row=i, column=j, value=v); c.font = Font(name=F, size=10, bold=i == 1 or j == 1); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = box
    wb.save(dest)


def export_numbers(dest, t):
    from numbers_parser import Alignment, Document, RGB
    sys.path.insert(0, HERE); from settings_sheet import exact_numbers; exact_numbers()
    st, sh, it = sheet_data(t); doc = None; S = {}
    def style(name, **kw):
        if name not in S: S[name] = doc.add_style(name=name, font_name='Helvetica Neue', font_size=kw.pop('size', 11.0), alignment=Alignment(kw.pop('h', 'left'), 'middle'), **kw)
        return S[name]
    def sheet(title, head, rows, widths, inputs, lists):
        nonlocal doc
        if doc is None: doc = Document(sheet_name=title, table_name=title, num_header_rows=1, num_header_cols=1, num_rows=len(rows) + 1, num_cols=len(head))
        else: doc.add_sheet(title, title, num_rows=len(rows) + 1, num_cols=len(head))
        tb = doc.sheets[-1].tables[0]; tb.num_header_rows = 1; tb.num_header_cols = 1
        head_s = style('BR head', bg_color=RGB(43, 36, 66), font_color=RGB(255, 255, 255), bold=True, h='center'); lab = style('BR label', bg_color=RGB(228, 224, 242), bold=True)
        val = style('BR value', bg_color=RGB(255, 247, 204), font_color=RGB(0, 0, 255)); note = style('BR note', size=10.0); key = style('BR key', bg_color=RGB(237, 237, 237), font_color=RGB(89, 89, 89), size=10.0)
        for j, w in enumerate(widths): tb.col_width(j, w)
        for j, h in enumerate(head): tb.write(0, j, h, style=head_s)
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                k = j + 1; s = val if k in inputs and k != 1 else lab if k == 1 else key if k == len(head) else note
                tb.write(i, j, '' if v is None else v, style=s)
                if k in lists: tb.set_cell_formatting(i, j, 'popup', popup_values=lists[k], allow_none=True)
            tb.row_height(i, 30)
    sheet('Items', IT_HEAD, it, [190, 80, 80, 90, 90, 90, 80, 200, 130], {1, 2, 3, 4, 5, 6, 7, 8}, {2: [s['id'] for s in t['shops']], 3: SLOTS, 4: ALL_STYLES})
    sheet('Shops', SH_HEAD, sh, [160, 100, 420, 90], {1, 2, 3}, {})
    sheet('Settings', ST_HEAD, st, [280, 70, 60, 80, 380, 90], {2}, {})
    sheet('How to fill', ['Item', 'What it means', 'How to fill'], [list(r) for r in TIPS], [150, 520, 120], set(), {})
    doc.save(dest)


def read_sheets(src):
    if src.lower().endswith('.numbers'):
        from numbers_parser import Document
        return {sh.name: [tuple(r) for r in sh.tables[0].rows(values_only=True)] for sh in Document(src).sheets}
    from openpyxl import load_workbook
    return {ws.title: list(ws.iter_rows(values_only=True)) for ws in load_workbook(src, data_only=True).worksheets}


def table(sheets, title):                     # (header -> column, rows) of a sheet; Numbers may rename "Items" to "Items - Table 1"
    data = next((r for n, r in sheets.items() if n.lower().startswith(title.lower())), None)
    if not data: return None, []
    hd = {str(h or '').strip().lower(): j for j, h in enumerate(data[0])}
    return hd, [r for r in data[1:] if any(c not in (None, '') for c in r)]


def num(v):
    if isinstance(v, bool) or v is None or v == '': return None
    if isinstance(v, (int, float)): v = round(float(v), 6); return int(v) if v.is_integer() else v
    try: return num(float(str(v).strip().replace(',', '').replace('€', '').replace('$', '')))
    except ValueError: return None


def text(v): return '' if v is None else str(v).strip()


def colour(v):
    s = text(v).lstrip('#').lower()
    return '#' + s if re.fullmatch(r'[0-9a-f]{6}', s) else None


def do_import(src, dry):
    s, i, j, t = read_table(); sheets = read_sheets(src); errors, changes = [], []
    new = {'settings': [dict(r) for r in t['settings']], 'shops': [], 'items': []}
    hd, rows = table(sheets, 'Settings')
    if hd is not None:
        S = {r['id']: r for r in new['settings']}; vc, kc = hd.get('value'), hd.get('id')
        for row in rows:
            k = text(row[kc]) if kc is not None and kc < len(row) else ''
            if k not in S: continue
            r = S[k]; v = num(row[vc] if vc < len(row) else None)
            if v is None or v < r['min'] or v > r['max'] or (r['unit'] in ('stars', 'count') and not float(v).is_integer()):
                errors.append('Settings / %s: "%s" must be a whole number from %s to %s' % (r['name'], row[vc] if vc < len(row) else '', fmt(r['min']), fmt(r['max']))); continue
            if v != r['v']: changes.append('Settings / %s: %s -> %s' % (r['name'], fmt(r['v']), fmt(v))); r['v'] = v
    hd, rows = table(sheets, 'Shops')
    if hd is None: errors.append('the Shops sheet is missing'); rows = []
    old = {x['id']: x for x in t['shops']}
    for n, row in enumerate(rows, start=2):
        g = lambda name: row[hd[name]] if name in hd and hd[name] < len(row) else None
        sid, name, col = text(g('id')).lower(), text(g('shop')).upper(), colour(g('colour (#hex)'))
        where = 'Shops row %d' % n
        if not re.fullmatch(r'[a-z][a-z0-9_]*', sid): errors.append('%s: ID "%s" must be lowercase letters, digits or _' % (where, sid)); continue
        if any(x['id'] == sid for x in new['shops']): errors.append('%s: ID %s is used twice' % (where, sid)); continue
        if not name: errors.append('%s: the shop needs a name' % where)
        if not col: errors.append('%s: "%s" is not a colour like #2bf3ff' % (where, g('colour (#hex)')))
        x = {'id': sid, 'name': name, 'color': col or '', 'theme': text(g('what it sells'))}; new['shops'].append(x)
        if sid not in old: changes.append('new shop %s (%s)' % (sid, name))
        elif x != old[sid]: changes.append('shop %s changed' % sid)
    for sid in old:
        if sid not in {x['id'] for x in new['shops']}: changes.append('shop %s taken out' % sid)
    shop_ids = {x['id'] for x in new['shops']}
    hd, rows = table(sheets, 'Items')
    if hd is None: errors.append('the Items sheet is missing'); rows = []
    old = {x['id']: x for x in t['items']}
    for n, row in enumerate(rows, start=2):
        g = lambda name: row[hd[name]] if name in hd and hd[name] < len(row) else None
        iid = text(g('id')).lower(); where = 'Items row %d (%s)' % (n, iid or 'no ID')
        if not re.fullmatch(r'[a-z][a-z0-9_]*', iid): errors.append('%s: ID must be lowercase letters, digits or _' % where); continue
        if any(x['id'] == iid for x in new['items']): errors.append('%s: ID used twice' % where); continue
        name, shop, slot, sty = text(g('name on screen')).upper(), text(g('shop')).lower(), text(g('slot')).lower(), text(g('style')).lower()
        c1, c2raw = colour(g('colour 1')), text(g('colour 2')); c2 = colour(c2raw) if c2raw else ''
        pr = num(g('price (€)'))
        if not name: errors.append('%s: a name is needed' % where)
        if shop and shop not in shop_ids: errors.append('%s: shop "%s" is not on the Shops sheet' % (where, shop))
        if slot not in STYLES: errors.append('%s: slot "%s" is not one of %s' % (where, slot, ', '.join(SLOTS))); continue
        if sty not in STYLES[slot]: errors.append('%s: style "%s" does not fit a %s (%s)' % (where, sty, slot, ', '.join(STYLES[slot])))
        if not c1: errors.append('%s: Colour 1 "%s" is not a colour like #ffd23f' % (where, g('colour 1')))
        if c2 is None: errors.append('%s: Colour 2 "%s" is not a colour (or leave it empty)' % (where, c2raw)); c2 = ''
        if shop:
            if pr is None or pr < 0 or pr > 100000 or not float(pr).is_integer(): errors.append('%s: price must be a whole number of euros, 0 to 100000' % where); pr = 0
        else: pr = None
        x = {'id': iid, 'name': name, 'shop': shop, 'slot': slot, 'style': sty, 'c1': c1 or '', 'c2': c2, 'price': pr, 'notes': text(g('notes'))}; new['items'].append(x)
        if iid not in old: changes.append('new piece %s (%s, %s)' % (iid, name, shop or 'starting clothes'))
        elif x != old[iid]: changes.append('piece %s: ' % iid + ', '.join('%s %s -> %s' % (k, old[iid].get(k), x[k]) for k in x if x[k] != old[iid].get(k)))
    for iid in old:
        if iid not in {x['id'] for x in new['items']}: changes.append('piece %s taken out' % iid)
    for sl in ('top', 'bottoms', 'shoes'):
        if not any(x['slot'] == sl and not x['shop'] for x in new['items']): errors.append('Items: one %s with an empty Shop is needed (what you wear in a new game)' % sl)
    for sid in shop_ids:
        if not any(x['shop'] == sid for x in new['items']): errors.append('Shops: %s sells nothing - give it some items' % sid)
    if errors: print('Not imported - fix these first:'); [print('  ' + e) for e in errors]; sys.exit(1)
    print('%d change%s' % (len(changes), '' if len(changes) == 1 else 's') + (':' if changes else '')); [print('  ' + c) for c in changes]
    if dry or not changes: return
    open(JS, 'w', encoding='utf-8').write(s[:i] + dump(new) + s[j:]); print('written', os.path.relpath(JS))


if __name__ == '__main__':
    if len(sys.argv) < 3 or sys.argv[1] not in ('export', 'import'): print(__doc__); sys.exit(2)
    if sys.argv[1] == 'export': (export_numbers if sys.argv[2].lower().endswith('.numbers') else export_xlsx)(sys.argv[2], read_table()[3]); print('saved', sys.argv[2])
    else: do_import(sys.argv[2], '--dry-run' in sys.argv[3:])
