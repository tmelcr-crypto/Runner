"""Police settings <-> spreadsheet (iOS Numbers or Excel).

  python3 tools/police_sheet.py export police.xlsx            write the police table from js/01b-police-data.js to a workbook
  python3 tools/police_sheet.py import police.xlsx [--dry-run] read an edited workbook back, check it and rewrite the table

The workbook is made for Numbers on an iPhone or iPad: plain tables with one header row, no merged cells, no comments, only simple
functions. Yellow cells are the values to edit; grey cells are calculated. Import matches rows by the ID column, so rows may be
moved around, but IDs must stay as they are. Needs openpyxl (pip install openpyxl)."""
import json, os, re, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

JS = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'js', '01b-police-data.js')
A, B = '/*POLICE-JSON*/', '/*END-POLICE-JSON*/'
STARS = ['1 star', '2 stars', '3 stars', '4 stars', '5 stars']
SHEETS = {'levels': 'Wanted levels', 'settings': 'Settings', 'crimes': 'Crimes'}


def read_table():
    s = open(JS, encoding='utf-8').read()
    i, j = s.index(A) + len(A), s.index(B)
    return s, i, j, json.loads(s[i:j])


def dump(t):                                # one row per line, the way the file is laid out
    out = ['{']
    for n, (k, rows) in enumerate(t.items()):
        out.append('"%s": [' % k)
        out += ['  ' + json.dumps(r, ensure_ascii=False) + (',' if m < len(rows) - 1 else '') for m, r in enumerate(rows)]
        out.append(']' + (',' if n < len(t) - 1 else ''))
    return '\n'.join(out + ['}'])


# ---------- export ----------
F = 'Arial'
f_in = PatternFill('solid', fgColor='FFF7CC'); f_calc = PatternFill('solid', fgColor='EDEDED'); f_head = PatternFill('solid', fgColor='2B2442')
f_lab = PatternFill('solid', fgColor='E4E0F2'); f_grp = PatternFill('solid', fgColor='F3F1FA')
thin = Side(style='thin', color='C9C9C9'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
yn = lambda v: 'YES' if v else 'NO'
allowed = lambda r: 'YES or NO' if r['unit'] == 'yes/no' else '%s to %s' % (fmt(r['min']), fmt(r['max']))
fmt = lambda v: ('%d' % v) if float(v).is_integer() else ('%g' % v)


def cell(ws, row, col, v, kind, wrap=False, bold=False, num='General'):
    c = ws.cell(row=row, column=col, value=v); c.border = box; c.number_format = num
    c.alignment = Alignment(wrap_text=wrap, vertical='center', horizontal='center' if kind in ('in', 'calc') and not wrap else 'left')
    if kind == 'head': c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = f_head
    elif kind == 'in': c.font = Font(name=F, size=10, color='0000FF'); c.fill = f_in
    elif kind == 'calc': c.font = Font(name=F, size=10, color='595959'); c.fill = f_calc
    elif kind == 'lab': c.font = Font(name=F, size=10, bold=True); c.fill = f_lab
    elif kind == 'id': c.font = Font(name=F, size=9, color='8A8A8A')
    else: c.font = Font(name=F, size=10, bold=bold)
    return c


def validate(ws, r, ref):
    if r['unit'] == 'yes/no':
        v = DataValidation(type='list', formula1='"YES,NO"', allow_blank=False, showErrorMessage=True, error='Pick YES or NO.')
    else:
        whole = r['unit'] == 'count'
        v = DataValidation(type='whole' if whole else 'decimal', operator='between', formula1=fmt(r['min']), formula2=fmt(r['max']), allow_blank=False,
                           showErrorMessage=True, error='Enter a %snumber from %s to %s.' % ('whole ' if whole else '', fmt(r['min']), fmt(r['max'])))
    ws.add_data_validation(v); v.add(ref)


def header(ws, names, widths):
    for j, (n, w) in enumerate(zip(names, widths), start=1): cell(ws, 1, j, n, 'head'); ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[1].height = 20


def export(path):
    t = read_table()[3]; wb = Workbook()
    # Wanted levels: one row per setting, one column per level
    ws = wb.active; ws.title = SHEETS['levels']
    header(ws, ['Setting'] + STARS + ['Unit', 'Allowed', 'What it does', 'ID'], [26, 9, 9, 9, 9, 9, 9, 13, 46, 11])
    rows = {}
    for i, r in enumerate(t['levels'], start=2):
        rows[r['id']] = i; cell(ws, i, 1, r['name'], 'lab', wrap=True)
        for k, v in enumerate(r['v']): cell(ws, i, 2 + k, yn(v) if r['unit'] == 'yes/no' else v, 'in')
        validate(ws, r, 'B%d:F%d' % (i, i))
        cell(ws, i, 7, r['unit'], ''); cell(ws, i, 8, allowed(r), ''); cell(ws, i, 9, r['note'], '', wrap=True); cell(ws, i, 10, r['id'], 'id')
        ws.row_dimensions[i].height = 44
    crime_row = {r['id']: i for i, r in enumerate(t['crimes'], start=2)}
    i = len(t['levels']) + 2; h = rows['heat']; cj = crime_row['carjack']
    cell(ws, i, 1, 'Calc: seen carjacks to get there', 'calc', wrap=True)
    for k in range(5):
        X = L(2 + k); cell(ws, i, 2 + k, '=IF(Crimes!$B$%d>0,ROUNDUP(%s%d/Crimes!$B$%d,0),"")' % (cj, X, h, cj), 'calc')
    cell(ws, i, 7, 'carjacks', 'calc'); cell(ws, i, 8, '', 'calc'); cell(ws, i, 9, 'How many carjackings a cop sees before you get this level.', 'calc', wrap=True); cell(ws, i, 10, '', 'calc')
    i += 1; cell(ws, i, 1, 'Calc: check', 'calc', wrap=True)
    for k in range(5):
        X, P = L(2 + k), L(1 + k)
        rise = '%s%d>0' % (X, h) if k == 0 else '%s%d>%s%d' % (X, h, P, h)
        cell(ws, i, 2 + k, '=IF(COUNTBLANK(%s2:%s%d)>0,"Fill every cell",IF(%s,"OK","Heat must be more than the level before"))' % (X, X, len(t['levels']) + 1, rise), 'calc', wrap=True)
    for j in (7, 8, 10): cell(ws, i, j, '', 'calc')
    cell(ws, i, 9, 'OK when the column is complete and needs more heat than the level before.', 'calc', wrap=True)
    ws.row_dimensions[i - 1].height = 30; ws.row_dimensions[i].height = 44
    ws.freeze_panes = 'B2'

    # Settings and Crimes: one row per value
    def flat(key, title, names, widths):
        ws = wb.create_sheet(title); header(ws, names, widths)
        for i, r in enumerate(t[key], start=2):
            cell(ws, i, 1, r['name'], 'lab', wrap=True)
            cell(ws, i, 2, yn(r['v']) if r['unit'] == 'yes/no' else r['v'], 'in'); validate(ws, r, 'B%d' % i)
            if r['unit'] == 'yes/no': chk = '=IF(OR(B{0}="YES",B{0}="NO"),"OK","Type YES or NO")'.format(i)
            else: chk = '=IF(AND(ISNUMBER(B{0}),B{0}>={1},B{0}<={2}),"OK","Must be {1} to {2}")'.format(i, fmt(r['min']), fmt(r['max']))
            cell(ws, i, 3, chk, 'calc'); cell(ws, i, 4, r['unit'], ''); cell(ws, i, 5, allowed(r), ''); cell(ws, i, 6, r['note'], '', wrap=True)
            j = 7
            if 'group' in r: c = cell(ws, i, j, r['group'], ''); c.fill = f_grp; j += 1
            cell(ws, i, j, r['id'], 'id'); ws.row_dimensions[i].height = 44 if len(r['note']) > 60 else 30
        ws.freeze_panes = 'B2'
    flat('settings', SHEETS['settings'], ['Setting', 'Value', 'Check', 'Unit', 'Allowed', 'What it does', 'Group', 'ID'], [27, 9, 10, 9, 13, 46, 15, 13])
    flat('crimes', SHEETS['crimes'], ['Crime', 'Heat', 'Check', 'Unit', 'Allowed', 'Note', 'ID'], [30, 8, 10, 8, 11, 46, 12])

    # How to fill
    g = wb.create_sheet('How to fill'); header(g, ['Item', 'What it means'], [24, 80])
    tips = [
        ('Yellow cells', 'The values you change. Tap a cell and type. YES / NO cells offer a list.', f_in),
        ('Grey cells', 'Calculated - leave them. Check says OK, or what is wrong in that row or column.', f_calc),
        ('ID column', 'How the game finds each row. Do not change it; rows may be moved or sorted.', None),
        ('Units', 'Distances in metres (a car is 4.5 m long, a street about 11 m wide), speeds in km/h or m/s, times in seconds, % as 0 to 100.', None),
        ('Heat and stars', 'Each crime the police see or hear adds heat (Crimes sheet). Enough heat gives a wanted level (Wanted levels: Heat needed). Heat below 1 star fades away; stars only go when you lose the police or pay a fine.', None),
        ('Seeing', 'A cop sees you within Sight range, inside the Field of view, with nothing in between. Whoever sees you tells the others by radio.', None),
        ('Searching', 'When nobody sees you they search around the spot where you were last seen (the red circle on the map, the stars blink). Out of sight for the level\'s time and they give up; inside the circle the time runs slower.', None),
        ('Stop order', 'Levels with Stop order and fine = YES: a cop near you orders you to stop. Stand still (on foot or in your car) for the fine; ignore it and you get the next level.', None),
        ('Arrest', 'Other levels: cars pull up beside you (or ram you where Cars ram you = YES), the crew gets out and a cop who holds you long enough arrests you. Drive away and they run back to their car.', None),
        ('Shooting', 'Only on levels with Shoot on sight = YES, or Shoot back if attacked = YES after you fired, rammed a police car or hurt a cop. They shout a warning first and never shoot with a passer-by in the way.', None),
        ('Busted / wasted', 'With Carry on after busted or wasted = YES you start again at the nearest police station or hospital and lose a share of your score (and your weapons, if set).', None),
        ('Sending it back', 'Share or export the file as Excel (.xlsx) and upload it in the chat. Every yellow value is read back into the game.', None),
    ]
    for i, (a, b_, fill) in enumerate(tips, start=2):
        c = cell(g, i, 1, a, '', bold=True); cell(g, i, 2, b_, '', wrap=True)
        if fill: c.fill = fill
        g.row_dimensions[i].height = 44
    g.freeze_panes = 'A2'
    wb.save(path); print('saved', path)


# ---------- import ----------
def num(v):
    if isinstance(v, bool): return None
    if isinstance(v, (int, float)): return int(v) if float(v).is_integer() else round(float(v), 6)
    if isinstance(v, str):
        try: return num(float(v.strip().replace(',', '.')))
        except ValueError: return None
    return None


def yesno(v):
    if isinstance(v, bool): return v
    if isinstance(v, (int, float)) and v in (0, 1): return bool(v)
    if isinstance(v, str) and v.strip().upper() in ('YES', 'NO', 'TRUE', 'FALSE', 'Y', 'N'): return v.strip().upper() in ('YES', 'TRUE', 'Y')
    return None


def find_sheet(wb, title):                  # Numbers may rename "Settings" to "Settings - Table 1" on export
    for ws in wb.worksheets:
        if ws.title == title or ws.title.startswith(title + ' ') or ws.title.startswith(title + '-'): return ws
    return None


def sheet_rows(ws):
    """header -> column map and {id: row values} for a sheet; the header is the first row that has an ID cell"""
    hdr, out = None, {}
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else '' for c in row]
        if hdr is None:
            if 'ID' in cells: hdr = {name.lower(): k for k, name in enumerate(cells) if name}
            continue
        k = hdr['id']
        if k < len(row) and row[k] not in (None, ''): out[str(row[k]).strip()] = row
    return hdr, out


def do_import(path, dry):
    s, i, j, t = read_table(); wb = load_workbook(path, data_only=True); errors, changes = [], []
    def take(r, raw, where):
        if r['unit'] == 'yes/no':
            v = yesno(raw)
            if v is None: errors.append('%s: "%s" is not YES or NO' % (where, raw))
            return v
        v = num(raw)
        if v is None: errors.append('%s: "%s" is not a number' % (where, raw)); return None
        if v < r['min'] or v > r['max']: errors.append('%s: %s is outside %s to %s' % (where, fmt(v), fmt(r['min']), fmt(r['max']))); return None
        if r['unit'] == 'count' and not float(v).is_integer(): errors.append('%s: %s must be a whole number' % (where, fmt(v))); return None
        return v
    show = lambda r, v: yn(v) if r['unit'] == 'yes/no' else fmt(v)
    for key, title in SHEETS.items():
        ws = find_sheet(wb, title)
        if ws is None: errors.append('sheet "%s" is missing' % title); continue
        hdr, rows = sheet_rows(ws)
        if hdr is None: errors.append('%s: no header row with an ID column' % title); continue
        for r in t[key]:
            if r['id'] not in rows: errors.append('%s: the row with ID %s is missing' % (title, r['id'])); continue
            row = rows[r['id']]
            if key == 'levels':
                cols = [hdr.get(n) for n in STARS]
                if None in cols: errors.append('%s: columns %s are needed' % (title, ', '.join(STARS))); break
                new = [take(r, row[c] if c < len(row) else None, '%s / %s / %s' % (title, r['name'], n)) for c, n in zip(cols, STARS)]
                if None in new: continue
                for n, a, b_ in zip(STARS, r['v'], new):
                    if a != b_: changes.append('%s / %s / %s: %s -> %s' % (title, r['name'], n, show(r, a), show(r, b_)))
                r['v'] = new
            else:
                c = hdr.get('value' if key == 'settings' else 'heat')
                if c is None: errors.append('%s: no %s column' % (title, 'Value' if key == 'settings' else 'Heat')); break
                v = take(r, row[c] if c < len(row) else None, '%s / %s' % (title, r['name']))
                if v is None: continue
                if v != r['v']: changes.append('%s / %s: %s -> %s' % (title, r['name'], show(r, r['v']), show(r, v)))
                r['v'] = v
    heat = next(r['v'] for r in t['levels'] if r['id'] == 'heat')
    if any(b_ <= a for a, b_ in zip(heat, heat[1:])): errors.append('Wanted levels / Heat needed must rise from level to level: %s' % heat)
    st = {r['id']: r['v'] for r in t['settings']}
    if st['gapMax'] < st['gapMin']: errors.append('Settings: Time between shots to (%s) is less than from (%s)' % (fmt(st['gapMax']), fmt(st['gapMin'])))
    if errors:
        print('Not imported - fix these first:'); [print('  ' + e) for e in errors]; sys.exit(1)
    print('%d change%s' % (len(changes), '' if len(changes) == 1 else 's') + (':' if changes else '')); [print('  ' + c) for c in changes]
    if dry or not changes: return
    open(JS, 'w', encoding='utf-8').write(s[:i] + dump(t) + s[j:]); print('written', os.path.relpath(JS))


if __name__ == '__main__':
    if len(sys.argv) < 3 or sys.argv[1] not in ('export', 'import'): print(__doc__); sys.exit(2)
    if sys.argv[1] == 'export': export(sys.argv[2])
    else: do_import(sys.argv[2], '--dry-run' in sys.argv[3:])
